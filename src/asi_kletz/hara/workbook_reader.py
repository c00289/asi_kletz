"""Read a HARA workbook (the wheel loader's ``WL RA_rev2.1.xlsx`` and its kin) into the plain dataset
``HaraService.create_from_dataset`` loads (0.28.0).

The workbook's anatomy is the dissection's (Roy's Brain, work product playground/hara/dissection): six kinds
of sheet, one chain of ids, one lookup. The reader finds every sheet by name (case and spaces forgiven), reads
the cached values (``data_only``: what Excel last computed, which is the evidence of what the document
states), and answers with:

- ``hara``: the document parameters (Document Information, Cover, Scope); a person's name is never read, a
  role parameter is the Brain's business;
- ``stpa``: step 1 for ``StpaService.create_from_dataset``: the losses, the hazards H-n with their scenarios
  H-n.m and hazardous events H-n.m.k as dotted hazards, the system-level constraints SC-n (the safety goals),
  every row cited to its sheet and row;
- ``vehicles``, ``functions``, ``use_cases``, ``failure_modes``, ``assumptions``: the reference lists;
- ``ratings``: the rating table, RA(ISO 19014) widened by RA(for all vehicles) where the former has no letter;
- ``assessments``: the rows of the four context sheets of one **generation** (1: the sheets Document
  Information names, whose MPLr formula answers a rating; 2: the hidden HARA-n twins, whose MPLr formula
  answers a use-case name because the vehicle code was empty), the hazard and the scenario derived from the
  event code, the typed values kept beside;
- ``approvals``: the Y / N of every approver **role** (the header "Robot Chief: <name>" becomes the role slug
  ``robot_chief``; the name is not read);
- ``warnings``: what the workbook could not say.

Standard library and openpyxl only; nothing here touches the database.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
from collections.abc import Iterable, Sequence
from pathlib import Path
from typing import Any

from ..errors import CommandError, NotFound
from ..schema import slugify
from ..schema.definitions import CUMULATIVE_CODE, HARA_CONTEXTS
from ..services.hara_service import NOT_FOUND, norm_key, norm_mplr, split_event

GENERATIONS = (1, 2)
# the sheet each part comes from, by its normalised name (lower, no spaces)
SHEETS = {
    "document_information": "documentinformation",
    "cover": "cover",
    "scope": "scope",
    "vocabulary": "vocabulary",
    "function_list": "functionlist",
    "use_case_list": "usecaselist",
    "failure_modes": "failuremodes",
    "assumption_list": "assumptionlist",
    "loss": "loss",
    "hazard_list": "hazardlist",
    "scenarios_list": "scenarioslist",
    "hazardous_events": "hazardousevents",
    "ratings": "ra(iso19014)",
    "ratings_all": "ra(forallvehicles)",
    "constraints": "system-levelconstraint",
    "summary": "summary",
}
REQUIRED = ("hazard_list", "hazardous_events", "ratings")
_LETTER = re.compile(r"^[A-Z]$")
_REVISION = re.compile(r"^(RA|Rev\.?|R|Revision)\s*\d+(\.\d+)*$", re.IGNORECASE)
_ROLE_HEADER = re.compile(r"^\s*([^:]+?)\s*:\s*(.*)$")
SAFETY_GOAL = re.compile(r"Safety Goal\s*:\s*(.+)", re.IGNORECASE | re.DOTALL)
CUMULATIVE_HEADER = "cummulati"  # the workbook's spelling of the Foundation Vehicle Cumulative column
SPECIFIC_HEADER = "specific mplr"


def _norm_name(text: Any) -> str:
    return "".join(str(text or "").split()).casefold()


def _text(value: Any) -> str:
    """A cell as text: ``None`` empty, a datetime its date, whitespace collapsed at the ends only."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\xa0", " ").strip()


def _flat(value: Any) -> str:
    """A cell as one line of text."""
    return " ".join(_text(value).split())


def _codes(value: Any) -> list[str]:
    """``"L-1, L-2, L-3"`` as a list."""
    return [c.strip() for c in re.split(r"[,;\n]+", _text(value)) if c.strip()]


def sha256_of(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


class Sheet:
    """The rows of one worksheet as lists of cells, 1-based rows kept so a cite can name them."""

    def __init__(self, title: str, rows: Sequence[Sequence[Any]]) -> None:
        self.title = title
        self.rows = [list(r) for r in rows]

    def find_row(self, *needles: str, start: int = 0) -> int | None:
        """The index of the first row from ``start`` holding every needle (as a cell's flat text, casefolded)."""
        wanted = [n.casefold() for n in needles]
        for index in range(start, len(self.rows)):
            cells = [_flat(c).casefold() for c in self.rows[index]]
            if all(any(cell == n or cell.startswith(n) for cell in cells) for n in wanted):
                return index
        return None

    def columns(self, index: int) -> dict[str, int]:
        """Header text (flat, casefolded) to column index, for the row at ``index``."""
        out: dict[str, int] = {}
        for col, cell in enumerate(self.rows[index]):
            text = _flat(cell).casefold()
            if text and text not in out:
                out[text] = col
        return out

    def cell(self, index: int, col: int | None) -> Any:
        if col is None or index >= len(self.rows) or col >= len(self.rows[index]):
            return None
        return self.rows[index][col]

    def cite(self, workbook: str, index: int, col: int | None = None) -> str:
        from openpyxl.utils import get_column_letter

        where = f"{get_column_letter(col + 1)}{index + 1}" if col is not None else f"row {index + 1}"
        return f"{workbook} › {self.title}!{where}"


def load_sheets(path: Path) -> dict[str, Sheet]:
    """Every worksheet, cached values, by its normalised name."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # a missing file, not a workbook, a locked OneDrive copy
        raise CommandError(f"cannot read {path}: {exc}") from exc
    sheets = {}
    for ws in wb.worksheets:
        sheets[_norm_name(ws.title)] = Sheet(ws.title, list(ws.iter_rows(values_only=True)))
    wb.close()
    return sheets


def _column_of(headers: dict[str, int], *needles: str) -> int | None:
    """The column whose header starts with (or contains) one of ``needles``."""
    for needle in needles:
        n = needle.casefold()
        for text, col in headers.items():
            if text == n or text.startswith(n) or n in text:
                return col
    return None


# -- the document ------------------------------------------------------------------------------------
def read_document(sheets: dict[str, Sheet], warnings: list[str]) -> dict[str, Any]:
    head: dict[str, Any] = {}
    info = sheets.get(SHEETS["document_information"])
    if info is not None:
        labels: dict[str, Any] = {}
        description_at = None
        for index, row in enumerate(info.rows):
            cells = list(row)
            for col in range(len(cells) - 1):
                label = _flat(cells[col]).casefold()
                value = cells[col + 1]
                if label and label != "description" and value is not None and _flat(value):
                    labels.setdefault(label, value)
            if description_at is None and any(_flat(c).casefold() == "description" for c in cells):
                description_at = index  # the first such row: the system description follows it
        head["system"] = _flat(labels.get("system name"))
        head["company"] = _flat(labels.get("company"))
        head["department"] = _flat(labels.get("department"))
        head["supporting_document"] = _flat(labels.get("supporting document"))
        head["project"] = _flat(labels.get("project"))
        head["vehicle_name"] = _flat(labels.get("vehicle name"))
        date = labels.get("hazard analysis date") or labels.get("analysis date")
        if isinstance(date, dt.datetime | dt.date):
            head["analysis_date"] = _text(date)
        elif date is not None:
            head["analysis_date"] = _text(date)
        code = _flat(labels.get("vehicle code"))
        if code:
            head["vehicle_code"] = code
        # the standard sits beside an empty label (E7 "Risk Assessment- ISO 19014")
        for row in info.rows:
            for cell in row:
                text = _flat(cell)
                if "iso 19014" in text.casefold() and "assess" in text.casefold() and len(text) < 60:
                    head.setdefault("standard", text)
        if description_at is not None:
            for row in info.rows[description_at + 1 : description_at + 4]:
                for cell in row:
                    text = _text(cell)
                    if len(text) > 40:
                        head["description"] = text
                        break
                if head.get("description"):
                    break
        # "Responsible" is a person's name: a role parameter in the Brain, never a value here
    else:
        warnings.append("no Document Information sheet: the document parameters are empty")
    cover = sheets.get(SHEETS["cover"])
    if cover is not None:
        for row in cover.rows:
            for cell in row:
                text = _flat(cell)
                if not text:
                    continue
                if "hazard analysis" in text.casefold() and "risk assessment" in text.casefold():
                    head.setdefault("title", text)
                elif _REVISION.match(text):
                    head.setdefault("revision", text)
    scope = sheets.get(SHEETS["scope"])
    if scope is not None:
        for row in scope.rows:
            cells = list(row)
            for col, cell in enumerate(cells):
                if _flat(cell).casefold().rstrip(":") == "scope" and col + 1 < len(cells):
                    text = _text(cells[col + 1])
                    if text:
                        head["scope"] = text
    return head


def read_vehicles(sheets: dict[str, Sheet], warnings: list[str]) -> list[dict[str, Any]]:
    scope = sheets.get(SHEETS["scope"])
    if scope is None:
        warnings.append("no Scope sheet: no vehicles")
        return []
    header = scope.find_row("annex")
    if header is None:
        warnings.append("Scope: no Annex header; no vehicles")
        return []
    cols = scope.columns(header)
    letter_col = cols.get("annex")
    iso_col = _column_of(cols, "vehicles on iso", "iso 19014")
    asi_col = _column_of(cols, "asi vehicles", "asi vehicle")
    project_col = _column_of(cols, "asi projects", "project")
    out = []
    for index in range(header + 1, len(scope.rows)):
        letter = _flat(scope.cell(index, letter_col)).upper()
        if not _LETTER.match(letter):
            continue
        out.append(
            {
                "letter": letter,
                "iso_vehicle": _flat(scope.cell(index, iso_col)),
                "asi_vehicle": _flat(scope.cell(index, asi_col)),
                "projects": _flat(scope.cell(index, project_col)),
            }
        )
    return out


def read_vocabulary(sheets: dict[str, Sheet]) -> dict[str, dict[str, str]]:
    """``{type: {term (casefold): definition}}``."""
    vocab = sheets.get(SHEETS["vocabulary"])
    out: dict[str, dict[str, str]] = {}
    if vocab is None:
        return out
    header = vocab.find_row("type", "terms")
    if header is None:
        return out
    cols = vocab.columns(header)
    for index in range(header + 1, len(vocab.rows)):
        kind = _flat(vocab.cell(index, cols.get("type"))).casefold()
        term = _flat(vocab.cell(index, cols.get("terms")))
        definition = _text(vocab.cell(index, cols.get("definition")))
        if kind and term:
            out.setdefault(kind, {})[term.casefold()] = definition
    return out


def read_functions(
    sheets: dict[str, Sheet], vocabulary: dict[str, dict[str, str]], warnings: list[str]
) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """The functions, and the map of vehicle name (casefold) to annex letter the other sheets need."""
    sheet = sheets.get(SHEETS["function_list"])
    if sheet is None:
        warnings.append("no Function List sheet: no functions")
        return [], {}
    header = sheet.find_row("sr.#", "annex")
    if header is None:
        warnings.append("Function List: no Sr.# / Annex header; no functions")
        return [], {}
    letters: dict[int, str] = {}  # column -> letter
    for col, cell in enumerate(sheet.rows[header]):
        text = _flat(cell).upper()
        if _LETTER.match(text):
            letters[col] = text
    names_row = header + 1
    vehicle_letters: dict[str, str] = {}
    cumulative_col = specific_col = None
    for col, cell in enumerate(sheet.rows[names_row] if names_row < len(sheet.rows) else []):
        text = _flat(cell)
        low = text.casefold()
        if col in letters and text:
            vehicle_letters[low] = letters[col]
        elif CUMULATIVE_HEADER in low or "cumulative" in low:
            cumulative_col = col
        elif SPECIFIC_HEADER in low:
            specific_col = col
    sr_col = sheet.columns(header).get("sr.#", 1)
    name_col = sheet.columns(header).get("annex", 2)
    out = []
    kind = "base"
    definitions = vocabulary.get("function", {})
    for index in range(names_row, len(sheet.rows)):
        sr = _flat(sheet.cell(index, sr_col))
        name = _flat(sheet.cell(index, name_col))
        if not name:
            continue
        if not sr:
            low = name.casefold()
            if "work implement" in low:
                kind = "work_implement"
            elif "base function" in low:
                kind = "base"
            continue
        vehicles = [letters[c] for c in letters if _flat(sheet.cell(index, c)).casefold() == "x"]
        out.append(
            {
                "name": name,
                "kind": kind,
                "vehicles": vehicles,
                "cumulative_mplr": norm_mplr(sheet.cell(index, cumulative_col)),
                "specific_mplr": norm_mplr(sheet.cell(index, specific_col)),
                "definition": definitions.get(name.casefold(), ""),
            }
        )
    return out, vehicle_letters


def read_use_cases(
    sheets: dict[str, Sheet],
    vehicle_letters: dict[str, str],
    vocabulary: dict[str, dict[str, str]],
    warnings: list[str],
) -> list[dict[str, Any]]:
    sheet = sheets.get(SHEETS["use_case_list"])
    if sheet is None:
        warnings.append("no Use Case List sheet: no use cases")
        return []
    header = sheet.find_row("sr.#")
    vehicles_row = sheet.find_row("vehicles")
    if header is None:
        warnings.append("Use Case List: no Sr.# header; no use cases")
        return []
    cols = sheet.columns(header)
    sr_col = cols.get("sr.#", 1)
    name_col = _column_of(cols, "use cases", "use case") or sr_col + 1
    letter_of: dict[int, str] = {}
    if vehicles_row is not None:
        for col, cell in enumerate(sheet.rows[vehicles_row]):
            low = _flat(cell).casefold()
            if low in vehicle_letters:
                letter_of[col] = vehicle_letters[low]
        if not letter_of:
            warnings.append("Use Case List: the vehicle names match no Function List vehicle; no letters")
    definitions = vocabulary.get("use case", {})
    out = []
    for index in range(header + 1, len(sheet.rows)):
        sr = _flat(sheet.cell(index, sr_col))
        name = _flat(sheet.cell(index, name_col))
        if not name or not sr or not sr[0].isdigit():
            continue
        vehicles = [letter_of[c] for c in letter_of if _flat(sheet.cell(index, c)).casefold() == "x"]
        out.append(
            {
                "name": name,
                "base": name.endswith("*"),
                "vehicles": vehicles,
                "definition": definitions.get(name.casefold(), "") or definitions.get(name.rstrip("* ").casefold(), ""),
            }
        )
    return out


def read_failure_modes(sheets: dict[str, Sheet], warnings: list[str]) -> list[dict[str, Any]]:
    sheet = sheets.get(SHEETS["failure_modes"])
    if sheet is None:
        warnings.append("no Failure modes sheet: no failure modes")
        return []
    header = sheet.find_row("sr.#")
    if header is None:
        return []
    cols = sheet.columns(header)
    sr_col = cols.get("sr.#", 1)
    name_col = _column_of(cols, "failure mode") or sr_col + 1
    out = []
    for index in range(header + 1, len(sheet.rows)):
        name = _flat(sheet.cell(index, name_col))
        if name:
            out.append({"name": name})
    return out


def read_assumptions(sheets: dict[str, Sheet], warnings: list[str]) -> list[dict[str, Any]]:
    sheet = sheets.get(SHEETS["assumption_list"])
    if sheet is None:
        return []
    header = sheet.find_row("sr#")
    if header is None:
        header = sheet.find_row("assumption")
    if header is None:
        return []
    # two tables side by side: every "Sr#" column opens one; the row above names its source
    starts = [col for col, cell in enumerate(sheet.rows[header]) if _flat(cell).casefold() in ("sr#", "sr.#")]
    sources = []
    for start in starts:
        source = "ISO 19014"
        if header > 0:
            above = [_flat(c) for c in sheet.rows[header - 1][start : start + 3]]
            if any("asi" in a.casefold() and "iso" not in a.casefold() for a in above):
                source = "ASI"
        sources.append(source)
    out = []
    for start, source in zip(starts, sources, strict=True):
        vehicle = ""
        for index in range(header + 1, len(sheet.rows)):
            vehicle = _flat(sheet.cell(index, start + 1)) or vehicle
            text = _text(sheet.cell(index, start + 2))
            if text:
                out.append({"source": source, "vehicle": vehicle, "text": text})
    return out


# -- step 1 -------------------------------------------------------------------------------------------
def read_step1(sheets: dict[str, Sheet], workbook: str, warnings: list[str]) -> tuple[dict[str, Any], dict[str, str]]:
    """The STPA dataset (losses, hazards with scenarios and events, constraints) and the hazard to safety-goal
    map ``{"H-1": "SC-1"}``."""
    losses: list[dict[str, Any]] = []
    hazards: list[dict[str, Any]] = []
    constraints: list[dict[str, Any]] = []
    goals: dict[str, str] = {}
    cut: dict[str, int] = {}  # sheet -> ids the workbook's LEFT() cut short, repaired from the text
    sheet = sheets.get(SHEETS["loss"])
    if sheet is not None:
        header = sheet.find_row("id", "loss")
        if header is not None:
            cols = sheet.columns(header)
            for index in range(header + 1, len(sheet.rows)):
                code = _flat(sheet.cell(index, cols.get("id")))
                if not code.upper().startswith("L"):
                    continue
                losses.append(
                    {
                        "code": code,
                        "text": _flat(sheet.cell(index, cols.get("loss"))),
                        "note": _text(sheet.cell(index, cols.get("detail"))),
                        "cite": sheet.cite(workbook, index, cols.get("id")),
                    }
                )
    sheet = sheets.get(SHEETS["hazard_list"])
    if sheet is None:
        raise NotFound(f"{workbook}: no Hazard List sheet")
    header = sheet.find_row("id", "hazard")
    if header is None:
        raise NotFound(f"{workbook}: Hazard List has no ID / Hazard header")
    cols = sheet.columns(header)
    status_col = max(cols.values()) + 1 if cols else None
    for index in range(header + 1, len(sheet.rows)):
        code = _flat(sheet.cell(index, cols.get("id")))
        if not code.upper().startswith("H"):
            continue
        detail = _text(sheet.cell(index, cols.get("detail")))
        status = _flat(sheet.cell(index, status_col))
        note = detail
        if status:
            note = f"{detail}\n[{status}]" if detail else f"[{status}]"
        hazards.append(
            {
                "code": code,
                "text": _flat(sheet.cell(index, cols.get("hazard"))),
                "losses": _codes(sheet.cell(index, cols.get("relation"))),
                "note": note,
                "cite": sheet.cite(workbook, index, cols.get("id")),
            }
        )
    sheet = sheets.get(SHEETS["scenarios_list"])
    if sheet is not None:
        header = sheet.find_row("scenario id")
        if header is not None:
            cols = sheet.columns(header)
            na_col = cols.get("na")
            for index in range(header + 1, len(sheet.rows)):
                code = _flat(sheet.cell(index, cols.get("scenario id")))
                if not code.upper().startswith("H"):
                    continue
                text = _flat(sheet.cell(index, cols.get("scenario")))
                code = _full_code(code, text, sheet.title, cut)
                hazards.append(
                    {
                        "code": code,
                        "text": _strip_code(text, code),
                        "parent": _parent(code, sheet.cell(index, cols.get("relations")), sheet.title, warnings),
                        "retired": bool(_flat(sheet.cell(index, na_col))),
                        "cite": sheet.cite(workbook, index, cols.get("scenario id")),
                    }
                )
    sheet = sheets.get(SHEETS["hazardous_events"])
    if sheet is None:
        raise NotFound(f"{workbook}: no Hazardous Events sheet")
    header = sheet.find_row("hazardous event id")
    if header is None:
        raise NotFound(f"{workbook}: Hazardous Events has no Hazardous Event ID header")
    cols = sheet.columns(header)
    comment_col = _column_of(cols, "comments")
    for index in range(header + 1, len(sheet.rows)):
        code = _flat(sheet.cell(index, cols.get("hazardous event id")))
        if not code.upper().startswith("H"):
            continue
        comment = _text(sheet.cell(index, comment_col))
        code = _full_code(code, _flat(sheet.cell(index, cols.get("hazardous event"))), sheet.title, cut)
        hazards.append(
            {
                "code": code,
                "text": _strip_code(_flat(sheet.cell(index, cols.get("hazardous event"))), code),
                "parent": _parent(code, sheet.cell(index, cols.get("relation")), sheet.title, warnings),
                "note": f"approval comment: {comment}" if comment else "",
                "cite": sheet.cite(workbook, index, cols.get("hazardous event id")),
            }
        )
    sheet = sheets.get(SHEETS["constraints"])
    if sheet is not None:
        header = sheet.find_row("id", "system-level constraint")
        if header is not None:
            cols = sheet.columns(header)
            for index in range(header + 1, len(sheet.rows)):
                code = _flat(sheet.cell(index, cols.get("id")))
                if not code.upper().startswith("SC"):
                    continue
                name = _flat(sheet.cell(index, cols.get("system-level constraint")))
                detail = _text(sheet.cell(index, cols.get("detail")))
                goal = SAFETY_GOAL.search(detail)
                text = " ".join(goal.group(1).split()) if goal else (detail or name)
                relation = _codes(sheet.cell(index, cols.get("relation")))
                for hazard in relation:
                    goals.setdefault(hazard, code)
                note_lines = [name] if name else []
                note_lines.extend(f"hazard: {h}" for h in relation)
                constraints.append(
                    {
                        "code": code,
                        "text": text,
                        "note": "\n".join(note_lines),
                        "cite": sheet.cite(workbook, index, cols.get("id")),
                    }
                )
    else:
        warnings.append("no System-Level Constraint sheet: no safety goals")
    for title, n in cut.items():
        warnings.append(
            f"{title}: {n} id(s) were cut short by the workbook's LEFT() (H-2.3.10 read as H-2.3.1); the code in "
            "front of the text is used; the assessment rows still carry the short id"
        )
    if not losses:
        warnings.append("no losses read")
    return {"losses": losses, "hazards": hazards, "constraints": constraints}, goals


def read_approvals(sheets: dict[str, Sheet], warnings: list[str]) -> list[dict[str, Any]]:
    sheet = sheets.get(SHEETS["hazardous_events"])
    if sheet is None:
        return []
    header = sheet.find_row("hazardous event id")
    if header is None:
        return []
    cols = sheet.columns(header)
    roles: dict[int, str] = {}
    for col, cell in enumerate(sheet.rows[header]):
        text = _flat(cell)
        match = _ROLE_HEADER.match(text)
        if match and "comment" not in text.casefold():
            roles[col] = slugify(match.group(1))  # the role; the name after the colon is not read
    if not roles:
        warnings.append("Hazardous Events: no 'Role: Name' approval columns")
        return []
    out = []
    for index in range(header + 1, len(sheet.rows)):
        code = _flat(sheet.cell(index, cols.get("hazardous event id")))
        if not code.upper().startswith("H"):
            continue
        for col, role in roles.items():
            decision = _flat(sheet.cell(index, col)).upper()[:1]
            if decision in ("Y", "N"):
                out.append({"event": code, "role": role, "decision": decision})
    return out


# -- the rating table ----------------------------------------------------------------------------------
def read_ratings(
    sheets: dict[str, Sheet], vehicle_letters: dict[str, str], warnings: list[str]
) -> list[dict[str, Any]]:
    """RA(ISO 19014) first (its header row carries the annex letters), then RA(for all vehicles) (its
    vehicles named, mapped to letters through the Function List) for the letters the first lacks."""
    table: dict[tuple[str, ...], dict[str, Any]] = {}
    order: list[tuple[str, ...]] = []
    primary = sheets.get(SHEETS["ratings"])
    if primary is None:
        raise NotFound("no RA(ISO 19014) sheet: no rating table")
    for sheet, is_primary in ((primary, True), (sheets.get(SHEETS["ratings_all"]), False)):
        if sheet is None:
            continue
        header = sheet.find_row("use case", "failure type")
        if header is None:
            warnings.append(f"{sheet.title}: no Use Case / Failure type header; skipped")
            continue
        letters_row = header - 1 if header > 0 else None
        cols = sheet.columns(header)
        function_col = cols.get("base functions", cols.get("annex", 1))
        use_col = cols.get("use case")
        mode_col = cols.get("failure type")
        cumulative_col = _column_of(cols, CUMULATIVE_HEADER, "cumulative")
        comment_col = _column_of(cols, "comments", "comment")
        letter_of: dict[int, str] = {}
        for col, cell in enumerate(sheet.rows[header]):
            if col in (function_col, use_col, mode_col, cumulative_col, comment_col):
                continue
            letter = _flat(sheet.cell(letters_row, col)).upper() if letters_row is not None else ""
            if _LETTER.match(letter):
                letter_of[col] = letter
            else:
                name = _flat(cell).casefold()
                if name in vehicle_letters:
                    letter_of[col] = vehicle_letters[name]
        if not letter_of:
            warnings.append(f"{sheet.title}: no vehicle column could be named by an annex letter; skipped")
            continue
        function = ""
        skipped = 0
        for index in range(header + 1, len(sheet.rows)):
            function = _flat(sheet.cell(index, function_col)) or function
            use_case = _flat(sheet.cell(index, use_col))
            mode = _flat(sheet.cell(index, mode_col))
            if not (function and use_case and mode):
                if function and any(_flat(sheet.cell(index, c)) for c in letter_of):
                    skipped += 1
                continue
            key = norm_key(function, use_case, mode)
            entry = table.get(key)
            if entry is None:
                entry = {
                    "function": function,
                    "use_case": use_case,
                    "failure_mode": mode,
                    "ratings": {},
                    "cumulative": "",
                    "comment": "",
                }
                table[key] = entry
                order.append(key)
            for col, letter in letter_of.items():
                value = norm_mplr(sheet.cell(index, col))
                if value and (is_primary or letter not in entry["ratings"]):
                    entry["ratings"][letter] = value
                elif is_primary and letter not in entry["ratings"]:
                    entry["ratings"][letter] = value  # an empty cell, kept so the column is known
            if is_primary or not entry["cumulative"]:
                entry["cumulative"] = entry["cumulative"] or norm_mplr(sheet.cell(index, cumulative_col))
            comment = _text(sheet.cell(index, comment_col))
            if comment and not entry["comment"]:
                entry["comment"] = comment
        if skipped:
            warnings.append(
                f"{sheet.title}: {skipped} rated row(s) carry no use case or failure type (function-level "
                "ratings, not a key of the table); skipped"
            )
    return [table[k] for k in order]


# -- the assessments -----------------------------------------------------------------------------------
def context_sheets(sheets: dict[str, Sheet], generation: int) -> dict[str, Sheet]:
    """The four context sheets of a generation: ``HARA(BF-BUC)`` (1) or ``HARA-1 (BF-BUC)`` (2)."""
    out: dict[str, Sheet] = {}
    for name, sheet in sheets.items():
        if not name.startswith("hara"):
            continue
        is_twin = name.startswith("hara-")
        if (generation == 2) != is_twin:
            continue
        for context in HARA_CONTEXTS:
            if context.casefold() in name:
                out[context] = sheet
    return out


def read_assessments(
    sheets: dict[str, Sheet], generation: int, goals: dict[str, str], warnings: list[str]
) -> list[dict[str, Any]]:
    found = context_sheets(sheets, generation)
    if not found:
        warnings.append(f"no assessment sheet of generation {generation}")
        return []
    missing = [c for c in HARA_CONTEXTS if c not in found]
    if missing:
        warnings.append(f"generation {generation}: no sheet for {', '.join(missing)}")
    out = []
    for context in HARA_CONTEXTS:
        sheet = found.get(context)
        if sheet is None:
            continue
        header = sheet.find_row("hazardous events id")
        if header is None:
            warnings.append(f"{sheet.title}: no Hazardous Events ID header; skipped")
            continue
        cols = sheet.columns(header)
        event_col = cols.get("hazardous events id")
        function_col = _column_of(cols, "functions", "function")
        use_col = _column_of(cols, "use cases", "use case")
        mode_col = _column_of(cols, "failure mode")
        mplr_col = cols.get("mplr")
        cumulative_col = _column_of(cols, "cummulative", "cumulative")
        comment_col = _column_of(cols, "comments", "comment")
        gid_col = _column_of(cols, "sg global id")
        goal_col = _column_of(cols, "safety constraint")
        sg_col = _column_of(cols, "sg mplr")
        garbage = 0
        for index in range(header + 1, len(sheet.rows)):
            event = _flat(sheet.cell(index, event_col))
            if not event.upper().startswith("H"):
                continue
            hazard, scenario, event = split_event(event)
            raw = sheet.cell(index, mplr_col)
            stated = norm_mplr(raw)
            if not stated and _flat(raw) not in ("", "0", NOT_FOUND):
                garbage += 1
            elif _flat(raw) == NOT_FOUND:
                stated = NOT_FOUND
            goal_text = _flat(sheet.cell(index, goal_col))
            out.append(
                {
                    "context": context,
                    "hazard": hazard,
                    "scenario": scenario,
                    "event": event,
                    "function": _flat(sheet.cell(index, function_col)),
                    "use_case": _flat(sheet.cell(index, use_col)),
                    "failure_mode": _flat(sheet.cell(index, mode_col)),
                    "stated_mplr": stated,
                    "cumulative_mplr": norm_mplr(sheet.cell(index, cumulative_col)),
                    "comment": _text(sheet.cell(index, comment_col)),
                    "sg_global_id": _flat(sheet.cell(index, gid_col)),
                    "safety_goal": goals.get(hazard, "") or goal_text,
                    "sg_mplr": norm_mplr(sheet.cell(index, sg_col)),
                }
            )
        if garbage:
            warnings.append(
                f"{sheet.title}: {garbage} MPLr cell(s) hold text that is no rating (the second generation's "
                "formula answers a use-case name when the vehicle code is empty); kept empty"
            )
    return out


_CODE_IN_TEXT = re.compile(r"^(H-\d+(?:\.\d+)*)\s*[:\-]")


def _full_code(code: str, text: str, sheet: str, cut: dict[str, int]) -> str:
    """The code in front of the text when it is longer than the id cell's (``=LEFT(B2,7)`` cut ``H-2.3.10``
    to ``H-2.3.1``); the id cell's otherwise."""
    match = _CODE_IN_TEXT.match(text)
    if match and match.group(1) != code and match.group(1).startswith(code):
        cut[sheet] = cut.get(sheet, 0) + 1
        return match.group(1)
    return code


def _parent(code: str, relation: Any, sheet: str, warnings: list[str]) -> str:
    """The parent a dotted code implies (H-1.1 for H-1.1.1); the sheet's Relation cell is checked against it
    and a disagreement (a typo such as ``H-4.1.``) is a warning, never the parent."""
    hazard, scenario, _event = split_event(code)
    implied = scenario if code.count(".") >= 2 else hazard
    stated = _flat(relation)
    if stated and implied and stated != implied:
        warnings.append(f"{sheet}: {code} relates to {stated!r}; its code implies {implied}, which is used")
    return implied or stated


def _strip_code(text: str, code: str) -> str:
    """``"H-1.1: Direct person interactions"`` without its own code in front."""
    if text.startswith(code):
        rest = text[len(code) :].lstrip()
        if rest.startswith(":") or rest.startswith("-"):
            return rest[1:].strip()
        return rest or text
    return text


def guess_vehicle_code(head: dict[str, Any], vehicles: Iterable[dict[str, Any]]) -> str | None:
    """The annex letter whose ISO 19014 vehicle name contains the document's vehicle name (casefold), or
    ``None``; what fills an empty Document Information C8."""
    name = " ".join(str(head.get("vehicle_name") or "").split()).casefold()
    if not name:
        return None
    singular = name.rstrip("s")
    for v in vehicles:
        iso = " ".join(str(v.get("iso_vehicle") or "").split()).casefold()
        if name in iso or singular in iso:
            return str(v["letter"])
    return None


def read_hara_workbook(path: str | Path, *, generation: int = 1, vehicle_code: str | None = None) -> dict[str, Any]:
    """The dataset of a HARA workbook (see the module docstring). ``generation`` picks the assessment sheets;
    ``vehicle_code`` overrides Document Information's (and the guess from the vehicle name when that is
    empty, which it is in rev 2.1)."""
    file = Path(path)
    if not file.is_file():
        raise NotFound(f"workbook not found: {file}")
    if generation not in GENERATIONS:
        raise CommandError(f"generation must be 1 or 2, not {generation!r}")
    warnings: list[str] = []
    sheets = load_sheets(file)
    for key in REQUIRED:
        if SHEETS[key] not in sheets:
            raise NotFound(f"{file.name}: no sheet named like {SHEETS[key]!r} (it has {', '.join(sheets)})")
    head = read_document(sheets, warnings)
    vocabulary = read_vocabulary(sheets)
    vehicles = read_vehicles(sheets, warnings)
    functions, vehicle_letters = read_functions(sheets, vocabulary, warnings)
    for v in vehicles:  # the Scope sheet names the vehicles too: a second way to a letter
        vehicle_letters.setdefault(" ".join(str(v["iso_vehicle"]).split()).casefold(), v["letter"])
    use_cases = read_use_cases(sheets, vehicle_letters, vocabulary, warnings)
    failure_modes = read_failure_modes(sheets, warnings)
    assumptions = read_assumptions(sheets, warnings)
    stpa, goals = read_step1(sheets, file.name, warnings)
    approvals = read_approvals(sheets, warnings)
    ratings = read_ratings(sheets, vehicle_letters, warnings)
    assessments = read_assessments(sheets, generation, goals, warnings)
    code = vehicle_code or head.get("vehicle_code")
    if not code:
        code = guess_vehicle_code(head, vehicles)
        if code:
            warnings.append(
                f"Document Information has no vehicle code; {code} taken from the vehicle name "
                f"{head.get('vehicle_name')!r} (pass --vehicle to choose another)"
            )
        else:
            warnings.append("no vehicle code: pass --vehicle <annex letter> or set it on the HARA; no MPLr is computed")
    if code and code.casefold() != CUMULATIVE_CODE:
        code = code.strip().upper()
    hara = {
        "name": head.get("title") or f"HARA {file.stem}",
        "system": head.get("system"),
        "vehicle_code": code,
        "vehicle_name": head.get("vehicle_name"),
        "project": head.get("project"),
        "department": head.get("department"),
        "company": head.get("company"),
        "standard": head.get("standard"),
        "supporting_document": head.get("supporting_document"),
        "analysis_date": head.get("analysis_date"),
        "revision": head.get("revision"),
        "description": head.get("description"),
        "scope": head.get("scope"),
        "source": str(file),
        "source_sha256": sha256_of(file),
        "source_generation": generation,
        "status": "Draft",
    }
    stpa["analysis"] = {
        "system": head.get("system"),
        "source_kind": "workbook",
        "source_bundle": str(file),
        "source_body_sha256": hara["source_sha256"],
        "status": "Draft",
        "note": f"step 1 read from {file.name} by `hara new from_workbook`",
    }
    return {
        "hara": hara,
        "stpa": stpa,
        "vehicles": vehicles,
        "functions": functions,
        "use_cases": use_cases,
        "failure_modes": failure_modes,
        "assumptions": assumptions,
        "ratings": ratings,
        "assessments": assessments,
        "approvals": approvals,
        "warnings": warnings,
    }
