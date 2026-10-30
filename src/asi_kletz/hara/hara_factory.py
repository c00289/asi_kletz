"""HaraFactory: the builders of a HARA and its export (0.28.0).

A **builder** turns options into the plain dataset ``HaraService.create_from_dataset`` loads through the
services, so every HARA obeys the same rules however it was made: ``blank`` (the document and nothing else),
``from_workbook`` (**the primary builder**: a HARA workbook such as ``WL RA_rev2.1.xlsx``, read by
``workbook_reader.read_hara_workbook``, its step 1 into an STPA analysis of the same slug) and ``clone``.
The same ``Builder`` shape as the control-structure and STPA factories.

``export`` writes the HARA back in the workbook's own shape (``<slug>.xlsx``: Document Information, Scope,
Function List, Use Case List, Failure modes, Assumption List, Loss, Hazard List, Scenarios List, Hazardous
Events with the approvals by role, RA(ISO 19014), one HARA(<context>) sheet per context with the **computed**
MPLr beside the stated one, System-Level Constraint, Summary derived) and a Markdown report (``<slug>.md``:
the parameters, the counts, the Summary, the check). The data is the database's; the files are views.
"""

from __future__ import annotations

from pathlib import Path
from typing import TYPE_CHECKING, Any

from ..errors import CommandError, NotFound
from ..schema.definitions import CUMULATIVE_CODE, HARA_CONTEXTS
from ..services.hara_service import (
    DATASET_KEYS,
    HARA_ENTITY,
    PART_ENTITIES,
    equivalents,
    ratings_of,
    split_event,
)
from ..stpa.factory import Builder

if TYPE_CHECKING:
    from ..brain import Brain

FORMATS = ("xlsx", "md")
DEFAULT_FORMATS = ("xlsx", "md")
EXPORT_DIR = "hara"  # under Brain.out_dir when --out is not given
_BUILDERS: dict[str, Builder] = {}


def register_hara_builder(builder: Builder) -> Builder:
    if builder.name in _BUILDERS:
        raise ValueError(f"HARA builder {builder.name!r} already registered")
    _BUILDERS[builder.name] = builder
    return builder


def _require(options: dict[str, Any], key: str, builder: str) -> str:
    value = options.get(key)
    if value in (None, ""):
        raise CommandError(f"{builder} needs --{'from' if key == 'path' else key}")
    return str(value)


def _blank(brain: Brain, options: dict[str, Any]) -> dict[str, Any]:
    return {"hara": {"name": _require(options, "name", "blank")}}


def _from_workbook(brain: Brain, options: dict[str, Any]) -> dict[str, Any]:
    from .workbook_reader import read_hara_workbook

    generation = options.get("generation")
    try:
        generation = int(generation) if generation not in (None, "") else 1
    except (TypeError, ValueError):
        raise CommandError(f"--generation must be 1 or 2, not {generation!r}") from None
    data = read_hara_workbook(
        _require(options, "path", "from_workbook"), generation=generation, vehicle_code=options.get("vehicle")
    )
    if options.get("analysis"):  # step 1 already in Kletz: the HARA points at it, the workbook's is not loaded
        data.pop("stpa", None)
        data["hara"]["analysis"] = str(options["analysis"])
    return data


def _clone(brain: Brain, options: dict[str, Any]) -> dict[str, Any]:
    source = brain.hara.resolve(_require(options, "path", "clone"))
    data = brain.hara.dataset(source)
    head = dict(data["hara"])
    head.pop("slug", None)
    data["hara"] = {**head, "status": "Draft", "note": f"duplicate of {source['slug']}"}
    for a in data.get("assessments") or []:
        a.pop("mplr", None)  # calculated again on load
    return data


register_hara_builder(
    Builder("blank", "The document and nothing else, for a HARA built by hand.", _blank, ("name", "system"))
)
register_hara_builder(
    Builder(
        "from_workbook",
        "A HARA workbook (--from <xlsx>): its lists, ratings, assessments and step 1; --vehicle <letter> picks the "
        "MPLr column, --generation 1|2 the assessment sheets, --analysis <slug> an STPA analysis already in Kletz.",
        _from_workbook,
        ("path", "vehicle", "generation", "analysis", "name", "slug", "system"),
    )
)
register_hara_builder(
    Builder(
        "clone",
        "A deep copy of an existing HARA (--from its id or slug), the same STPA analysis.",
        _clone,
        ("path", "name", "slug"),
    )
)


class HaraFactory:
    def __init__(self, brain: Brain) -> None:
        self.brain = brain

    @staticmethod
    def builders() -> list[str]:
        return list(_BUILDERS)

    @staticmethod
    def builder(name: str) -> Builder:
        try:
            return _BUILDERS[str(name)]
        except KeyError:
            raise NotFound(f"unknown builder {name!r}; known: {', '.join(_BUILDERS)}") from None

    def catalog(self) -> list[dict[str, Any]]:
        return [b.to_dict(self.brain) for b in _BUILDERS.values()]

    def create(self, builder: str, actor: str = "python", **options: Any) -> tuple[dict[str, Any], list[str]]:
        """Build a HARA with ``builder``: the dataset it returns is loaded through the services. ``name``,
        ``slug``, ``system`` and ``vehicle`` given here replace the dataset's; the reader's warnings come
        first."""
        spec = self.builder(builder)
        reason = spec.available(self.brain)
        if reason:
            raise NotFound(f"{spec.name} is unavailable: {reason}")
        dataset = spec.build(self.brain, dict(options))
        head: dict[str, Any] = {k: options[k] for k in ("name", "slug", "system") if options.get(k)}
        if options.get("vehicle") and builder != "from_workbook":
            head["vehicle_code"] = options["vehicle"]
        row, warnings = self.brain.hara.create_from_dataset(dataset, actor, **head)
        return row, warnings

    # -- the files ------------------------------------------------------
    def export(
        self,
        ref: Any,
        out_dir: str | Path | None = None,
        formats: tuple[str, ...] | None = None,
        *,
        actor: str = "python",
    ) -> dict[str, Any]:
        """Write ``<out>/<slug>.xlsx`` (the workbook shape, computed MPLr) and ``<out>/<slug>.md`` (the report).
        Audited as ``hara.export``, published as ``hara.exported``."""
        formats = tuple(dict.fromkeys(formats or DEFAULT_FORMATS))
        unknown = [f for f in formats if f not in FORMATS]
        if unknown:
            raise CommandError(f"unknown format(s) {', '.join(unknown)}; choose from {', '.join(FORMATS)}")
        row = self.brain.hara.resolve(ref)
        out = Path(out_dir) if out_dir else Path(self.brain.out_dir) / EXPORT_DIR
        out.mkdir(parents=True, exist_ok=True)
        paths: dict[str, str] = {}
        if "xlsx" in formats:
            paths["xlsx"] = str(self.export_workbook(row, out / f"{row['slug']}.xlsx"))
        if "md" in formats:
            paths["md"] = str(self.export_report(row, out / f"{row['slug']}.md"))
        data = {
            "id": row["id"],
            "slug": row["slug"],
            "name": row["name"],
            "out": str(out),
            "formats": list(formats),
            "paths": paths,
            "files": list(paths.values()),
            "primary": paths.get("xlsx") or paths.get("md"),
            "issues": [],
        }
        self.brain.storage.audit(actor, "hara.export", HARA_ENTITY, row["id"], {"slug": row["slug"], "paths": paths})
        self.brain.events.publish("hara.exported", hara=row["slug"], paths=paths, actor=actor)
        return data

    def export_workbook(self, row: dict[str, Any], path: Path) -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        brain = self.brain
        slug = row["slug"]
        parts = brain.hara.parts(slug)
        hazards = brain.hara.hazards(row)
        constraints = brain.hara.constraints(row)
        summary = brain.hara.hazard_summary(row)
        bold = Font(bold=True)
        differ = PatternFill("solid", fgColor="FFF4C2C2")
        wb = Workbook()
        wb.remove(wb.active)

        def sheet(title: str, headers: list[str], rows: list[list[Any]], widths: dict[int, int] | None = None):
            ws = wb.create_sheet(title[:31])
            ws.append(headers)
            for cell in ws[1]:
                cell.font = bold
            for r in rows:
                ws.append(r)
            for col, width in (widths or {}).items():
                ws.column_dimensions[get_column_letter(col)].width = width
            ws.freeze_panes = "A2"
            return ws

        info = [
            ["System Name", row.get("system")],
            ["Company", row.get("company")],
            ["Department", row.get("department")],
            ["Hazard Analysis date", row.get("analysis_date")],
            ["Supporting document", row.get("supporting_document")],
            ["Project", row.get("project")],
            ["Standard", row.get("standard")],
            ["Vehicle code", row.get("vehicle_code")],
            ["Vehicle Name", row.get("vehicle_name")],
            ["Revision", row.get("revision")],
            ["STPA analysis", row.get("analysis")],
            ["Source", row.get("source")],
            ["Source SHA-256", row.get("source_sha256")],
            ["Generation read", row.get("source_generation")],
            ["Description", row.get("description")],
            ["Scope", row.get("scope")],
        ]
        sheet("Document Information", ["Document Information", "Value"], info, {1: 24, 2: 90})
        sheet(
            "Scope",
            ["Annex", "Vehicles on ISO 19014", "ASI Vehicles", "ASI Projects"],
            [[v["letter"], v.get("iso_vehicle"), v.get("asi_vehicle"), v.get("projects")] for v in parts["vehicles"]],
            {2: 44, 3: 30, 4: 30},
        )
        letters = [v["letter"] for v in parts["vehicles"]] or sorted(
            {letter for f in parts["functions"] for letter in (f.get("vehicles") or [])}
        )
        sheet(
            "Function List",
            ["Sr.#", "Function", "Kind", *letters, "Foundation Vehicle Cumulative", "Specific MPLr", "Definition"],
            [
                [
                    f.get("order"),
                    f["name"],
                    f.get("kind"),
                    *["x" if letter in (f.get("vehicles") or []) else None for letter in letters],
                    f.get("cumulative_mplr"),
                    f.get("specific_mplr"),
                    f.get("definition"),
                ]
                for f in parts["functions"]
            ],
            {2: 34},
        )
        sheet(
            "Use Case List",
            ["Sr.#", "Use Case", "Base", *letters, "Definition"],
            [
                [
                    u.get("order"),
                    u["name"],
                    "x" if u.get("base") else None,
                    *["x" if letter in (u.get("vehicles") or []) else None for letter in letters],
                    u.get("definition"),
                ]
                for u in parts["use_cases"]
            ],
            {2: 44},
        )
        sheet(
            "Failure modes",
            ["Sr.#", "Failure Modes"],
            [[m.get("order"), m["name"]] for m in parts["failure_modes"]],
            {2: 32},
        )
        sheet(
            "Assumption List",
            ["Sr#", "Source", "Vehicles", "Assumption"],
            [[a.get("order"), a.get("source"), a.get("vehicle"), a.get("text")] for a in parts["assumptions"]],
            {3: 32, 4: 90},
        )
        analysis = row.get("analysis")
        losses = brain.service("stpa_loss").list({"analysis": analysis}) if analysis else []
        sheet(
            "Loss",
            ["ID", "Loss", "Detail", "Relation"],
            [
                [
                    x["code"],
                    x.get("text"),
                    x.get("note"),
                    ", ".join(h for h, hz in hazards.items() if "." not in h and x["code"] in (hz.get("losses") or [])),
                ]
                for x in losses
            ],
            {2: 44, 3: 60, 4: 30},
        )
        top = [h for code, h in hazards.items() if "." not in code]
        sheet(
            "Hazard List",
            ["ID", "Hazard", "Detail", "Relation"],
            [[h["code"], h.get("text"), h.get("note"), ", ".join(h.get("losses") or [])] for h in top],
            {2: 60, 3: 60, 4: 30},
        )
        scenarios = [h for code, h in hazards.items() if code.count(".") == 1]
        sheet(
            "Scenarios List",
            ["Scenario ID", "Scenario", "Relations", "NA"],
            [
                [h["code"], f"{h['code']}: {h.get('text') or ''}", h.get("parent"), "x" if h.get("retired") else None]
                for h in scenarios
            ],
            {2: 70},
        )
        roles = sorted({a["role"] for a in parts["approvals"]})
        decisions: dict[tuple[str, str], str] = {
            (a["event"], a["role"]): a.get("decision") or "" for a in parts["approvals"]
        }
        events = [h for code, h in hazards.items() if code.count(".") >= 2]
        sheet(
            "Hazardous Events",
            ["Hazardous Event ID", "Hazardous Event", "Relation", *roles, "Comments"],
            [
                [
                    h["code"],
                    f"{h['code']}: {h.get('text') or ''}",
                    h.get("parent"),
                    *[decisions.get((h["code"], role), None) for role in roles],
                    h.get("note"),
                ]
                for h in events
            ],
            {2: 80},
        )
        rating_letters = sorted({letter for r in parts["ratings"] for letter in ratings_of(r.get("ratings"))})
        sheet(
            "RA(ISO 19014)",
            [
                "Sr.#",
                "Function",
                "Use Case",
                "Failure type",
                *rating_letters,
                "Foundation Vehicle Cumulative",
                "Comments",
            ],
            [
                [
                    r.get("order"),
                    r["function"],
                    r["use_case"],
                    r["failure_mode"],
                    *[ratings_of(r.get("ratings")).get(letter) or None for letter in rating_letters],
                    r.get("cumulative"),
                    r.get("comment"),
                ]
                for r in parts["ratings"]
            ],
            {2: 26, 3: 40, 4: 26},
        )
        code = row.get("vehicle_code") or ""
        mplr_header = f"MPLr ({'cumulative' if code == CUMULATIVE_CODE else code or 'no vehicle code'})"
        for context in HARA_CONTEXTS:
            rows_of = [a for a in parts["assessments"] if a.get("context") == context]
            ws = sheet(
                f"HARA({context})",
                [
                    "Hazard ID",
                    "Hazards",
                    "Scenarios ID",
                    "Scenarios",
                    "Hazardous Events ID",
                    "Hazardous Events",
                    "Functions",
                    "Use Cases",
                    "Failure Mode",
                    mplr_header,
                    "Stated MPLr",
                    "Cumulative Hazard MPLr",
                    "Comments",
                    "SG Global ID",
                    "Safety Constraint (Safety Goal)",
                    "SG MPLr",
                ],
                [],
                {2: 40, 4: 40, 6: 50, 7: 22, 8: 30, 9: 24},
            )
            for a in rows_of:
                hazard = a.get("hazard") or split_event(a.get("event"))[0]
                scenario = a.get("scenario") or split_event(a.get("event"))[1]
                goal = constraints.get(a.get("safety_goal") or "")
                ws.append(
                    [
                        hazard,
                        (hazards.get(hazard) or {}).get("text"),
                        scenario,
                        (hazards.get(scenario) or {}).get("text"),
                        a.get("event"),
                        (hazards.get(a.get("event")) or {}).get("text"),
                        a.get("function"),
                        a.get("use_case"),
                        a.get("failure_mode"),
                        a.get("mplr"),
                        a.get("stated_mplr"),
                        a.get("cumulative_mplr"),
                        a.get("comment"),
                        a.get("sg_global_id"),
                        f"{a.get('safety_goal')}: {goal.get('text')}" if goal else a.get("safety_goal"),
                        a.get("sg_mplr"),
                    ]
                )
                if a.get("stated_mplr") and a.get("mplr") != a.get("stated_mplr"):
                    ws.cell(row=ws.max_row, column=10).fill = differ
        sheet(
            "System-Level Constraint",
            ["ID", "System-Level Constraint", "Safety Goal", "Relation"],
            [
                [
                    c["code"],
                    (c.get("note") or "").splitlines()[0] if c.get("note") else "",
                    c.get("text"),
                    ", ".join(
                        line.split(":", 1)[1].strip()
                        for line in (c.get("note") or "").splitlines()
                        if line.startswith("hazard:")
                    ),
                ]
                for c in constraints.values()
            ],
            {2: 44, 3: 80},
        )
        ws = sheet(
            "Summary",
            [
                "Hazard",
                "Safety Goal",
                *HARA_CONTEXTS,
                "MPLr (ISO 19014)",
                "PLr (ISO 13849)",
                "AgPLr (ISO 25119)",
                "SIL (IEC 61508)",
                *[f"typed {c}" for c in HARA_CONTEXTS],
                "typed MPLr",
                "agrees",
            ],
            [],
            {1: 60, 2: 70},
        )
        for e in summary["hazards"]:
            ws.append(
                [
                    f"{e['hazard']}: {e['text']}",
                    f"{e['safety_goal']}: {e['safety_goal_text']}" if e["safety_goal"] else "",
                    *[e["contexts"][c] for c in HARA_CONTEXTS],
                    e["mplr"],
                    e["plr"],
                    e["agplr"],
                    e["sil"],
                    *[e["typed_cumulative"][c] for c in HARA_CONTEXTS],
                    e["typed_mplr"],
                    "" if e["agrees"] is None else ("yes" if e["agrees"] else "no"),
                ]
            )
            if e["agrees"] is False:
                ws.cell(row=ws.max_row, column=7).fill = differ
        wb.save(path)
        return path

    def export_report(self, row: dict[str, Any], path: Path) -> Path:
        brain = self.brain
        slug = row["slug"]
        counts = brain.hara.counts(slug)
        summary = brain.hara.hazard_summary(row)
        check = brain.hara.check(row)
        lines = [f"# {row['name']}", ""]
        lines.append(
            f"Vehicle code **{row.get('vehicle_code') or '(none)'}** ({row.get('vehicle_name') or '?'}); "
            f"STPA analysis `{row.get('analysis') or '-'}`; status {row.get('status')}."
        )
        if row.get("source"):
            lines.append(
                f"Read from `{Path(str(row['source'])).name}` (generation {row.get('source_generation')}, "
                f"SHA-256 `{(row.get('source_sha256') or '')[:16]}…`)."
            )
        lines += ["", "## Document", "", "| Parameter | Value |", "|---|---|"]
        for key in (
            "system",
            "company",
            "department",
            "project",
            "standard",
            "supporting_document",
            "analysis_date",
            "revision",
        ):
            lines.append(f"| {key.replace('_', ' ')} | {row.get(key) or ''} |")
        lines += ["", "## Counts", "", "| Part | Rows |", "|---|---|"]
        for key, n in counts.items():
            lines.append(f"| {key.replace('_', ' ')} | {n} |")
        lines += ["", "## Summary (computed)", ""]
        lines.append(
            "| Hazard | Safety goal | " + " | ".join(HARA_CONTEXTS) + " | MPLr | PLr | AgPLr | SIL | typed | agrees |"
        )
        lines.append("|---|---|" + "---|" * (len(HARA_CONTEXTS) + 6))
        for e in summary["hazards"]:
            agrees = "" if e["agrees"] is None else ("yes" if e["agrees"] else "**no**")
            lines.append(
                f"| {e['hazard']} | {e['safety_goal']} | "
                + " | ".join(e["contexts"][c] or "" for c in HARA_CONTEXTS)
                + f" | {e['mplr']} | {e['plr']} | {e['agplr']} | {e['sil']} | {e['typed_mplr']} | {agrees} |"
            )
        lines += ["", "## Check: computed against the workbook", ""]
        lines.append("| Context | Rows | Agree | Differ | No key | Unstated |")
        lines.append("|---|---|---|---|---|---|")
        for context, c in check["contexts"].items():
            lines.append(
                f"| {context} | {c['rows']} | {c['agree']} | {c['differ']} | {c['not_found']} | {c['unstated']} |"
            )
        for context, c in check["contexts"].items():
            if c["examples"]:
                lines += ["", f"Examples, {context}:", ""]
                for x in c["examples"]:
                    lines.append(
                        f"- {x['event']} · {x['function']} · {x['use_case']} · {x['failure_mode']}: "
                        f"computed {x['computed'] or '(empty)'}, stated {x['stated']}"
                    )
        if any(check["unknown"].values()):
            lines += ["", "Used but not listed:", ""]
            for bucket, names in check["unknown"].items():
                if names:
                    lines.append(f"- {bucket.replace('_', ' ')}: {', '.join(names)}")
        lines.append("")
        path.write_text("\n".join(lines), encoding="utf-8")
        return path


__all__ = [
    "DEFAULT_FORMATS",
    "FORMATS",
    "HaraFactory",
    "register_hara_builder",
    "DATASET_KEYS",
    "PART_ENTITIES",
    "equivalents",
]
