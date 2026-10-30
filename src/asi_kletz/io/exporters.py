"""Exporters: the workbook and the JSON dump are *views* regenerated from the database.

Since 0.9.0 both carry the hidden ``uid`` of every row: the JSON dump as any key, the workbook as the
last column of every sheet, headed ``uid`` and hidden, so the sheet reads as before and still round-trips.

0.20.0, the owner: "make sure that the excel export includes links to the work products wherever they may
be... on multiple urls just go to first". The Safety Plan sheet gets a derived **Work product** column (no
field; a clickable hyperlink to the row's first URL, the cell text the URL's host, the Hyperlink style) before
the **Work product URLs** column (the ``urls`` field, one URL per line, read back by the import). A row with
no Jira ticket prints ``no ticket attached`` in its Jira key cell, which the import reads as no key.

0.21.0, the owner: "turn the name of the safety deliverable into a clickable URL". The Safety Deliverables cell
of a row that has a URL is itself a hyperlink to the first one (the Hyperlink style), beside the Work product
column; the import reads its text as before.

0.22.0 (Phase A of the Confluence integration plan): the Work product cell's text is the first URL's kind and
host (``Confluence, asirobots.atlassian.net``; a placeholder row reads ``placeholder``), so a reviewer sees at a
glance which deliverables still point nowhere; the hyperlink is the first URL as before.

0.27.0, the owner: "bug when exporting safety plan excel". The sheet had linked every deliverable name and every
Work product cell to the placeholder in hyperlink blue, the gap 0.26.0 paints pinkish red in the Word export. The
sheet now marks the gaps the same way: a placeholder link stays a link but dark red ``#B71C1C`` over pinkish red
``#F4C2C2`` (the name and the Work product cell both), a deliverable with no URL is shaded pinkish red without a
link, ``no ticket attached`` is shaded too, a real link stays blue; the summary counts them under ``gaps``. And
the workbook carries the timestamp suffix (:mod:`asi_kletz.io.stamp`); the JSON dump keeps the exact name it is
given, a machine artefact that ``load`` and the scripts read back by that path.
"""

from __future__ import annotations

import abc
import datetime as dt
import json
from collections.abc import Callable
from pathlib import Path
from typing import TYPE_CHECKING, Any, ClassVar

from ..errors import NotFound
from ..schema import EntitySchema, FieldType
from ..schema.definitions import NO_TICKET, WORK_PRODUCT_HEADER
from ..schema.fields import json_text
from ..schema.urls import first_url, kind_label, url_kind, urls_lines
from ..version import SCHEMA_VERSION, __version__
from .stamp import next_sibling, stamped

if TYPE_CHECKING:
    from ..brain import Brain

DELIVERABLE = "deliverable"  # the Safety Plan row's name field: its cell links to the first URL (0.21.0)


def _timestamp() -> str:
    return dt.datetime.now(dt.UTC).replace(microsecond=0).isoformat()


class Exporter(abc.ABC):
    format: ClassVar[str] = ""
    suffix: ClassVar[str] = ""

    @abc.abstractmethod
    def export(self, brain: Brain, out: str | Path | None, entity: str | None = None) -> dict[str, Any]:
        """Write the view. Returns a summary with ``path``, ``entities`` (counts) and ``warnings``."""


class ExporterFactory:
    _exporters: dict[str, type[Exporter]] = {}

    @classmethod
    def register(cls, exporter: type[Exporter]) -> type[Exporter]:
        cls._exporters[exporter.format] = exporter
        return exporter

    @classmethod
    def create(cls, fmt: str) -> Exporter:
        try:
            return cls._exporters[fmt]()
        except KeyError:
            raise NotFound(f"unknown export format {fmt!r}; known: {', '.join(cls._exporters)}") from None

    @classmethod
    def formats(cls) -> list[str]:
        return list(cls._exporters)


def _schemas(brain: Brain, entity: str | None) -> list[EntitySchema]:
    return [brain.registry.get(entity)] if entity else list(brain.registry)


def sheet_columns(schema: EntitySchema) -> list[tuple[str, Any]]:
    """The sheet's columns in order, ``(header, FieldSpec or None)``: every field, and before the ``urls``
    field the derived ``Work product`` column (``None``), the hyperlink to the row's first URL (0.20.0)."""
    columns: list[tuple[str, Any]] = []
    for spec in schema.fields:
        if spec.picker == "urls":
            columns.append((WORK_PRODUCT_HEADER, None))
        columns.append((spec.header, spec))
    return columns


def _save_with_fallback(save: Callable[[Path], None], out: Path, stamp: bool = True) -> tuple[Path, str | None]:
    """Save ``out`` under its timestamped name (:func:`~asi_kletz.io.stamp.stamped`; the exact name when ``stamp``
    is off: the JSON dump is a machine artefact that scripts and tests read back by the path they gave); if that
    is locked (open in another program), save to its ``-2`` sibling instead."""
    out.parent.mkdir(parents=True, exist_ok=True)
    target = stamped(out) if stamp else out
    try:
        save(target)
        return target, None
    except PermissionError:
        alt = next_sibling(target)
        save(alt)
        return alt, f"{target.name} is locked (open in another program); wrote {alt.name} instead"


@ExporterFactory.register
class JsonExporter(Exporter):
    format = "json"
    suffix = ".json"

    def payload(self, brain: Brain, entity: str | None = None) -> dict[str, Any]:
        base = {"schema_version": SCHEMA_VERSION, "app_version": __version__, "exported_at": _timestamp()}
        if entity:
            schema = brain.registry.get(entity)
            return {**base, "entity": schema.name, "rows": brain.service(schema.name).list()}
        return {**base, "entities": {s.name: brain.service(s.name).list() for s in brain.registry}}

    def export(self, brain: Brain, out: str | Path | None, entity: str | None = None) -> dict[str, Any]:
        payload = self.payload(brain, entity)
        target = Path(out) if out else Path(f"{payload.get('entity', 'asi_kletz')}.json")

        def save(path: Path) -> None:
            path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

        path, warning = _save_with_fallback(save, target, stamp=False)
        if entity:
            counts = {payload["entity"]: len(payload["rows"])}
        else:
            counts = {name: len(rows) for name, rows in payload["entities"].items()}
        return {"format": self.format, "path": str(path), "entities": counts, "warnings": [warning] if warning else []}


@ExporterFactory.register
class ExcelExporter(Exporter):
    format = "excel"
    suffix = ".xlsx"
    HEADER_FILL = "FFD9E1F2"
    GAP_FILL = "FFF4C2C2"      # 0.27.0: pinkish red on a gap, asi_gutenberg's attention mark on the sheet
    GAP_FONT = "FFB71C1C"      # dark red for a link that still points at the placeholder

    def _gap(self, cell: Any) -> None:
        from openpyxl.styles import PatternFill

        cell.fill = PatternFill("solid", fgColor=self.GAP_FILL)

    def _link(self, cell: Any, url: str, gap: bool = False) -> None:
        """A hyperlink in Excel's blue; over the placeholder, dark red on pinkish red and still a link."""
        from openpyxl.styles import Font

        cell.hyperlink = url
        cell.style = "Hyperlink"
        if gap:
            cell.font = Font(color=self.GAP_FONT, underline="single")
            self._gap(cell)

    def export(self, brain: Brain, out: str | Path | None, entity: str | None = None) -> dict[str, Any]:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = Workbook()
        wb.remove(wb.active)
        counts: dict[str, int] = {}
        gaps = {"placeholder": 0, "no_url": 0, "no_ticket": 0}
        for schema in _schemas(brain, entity):
            ws = wb.create_sheet(schema.sheet[:31])
            rows = brain.service(schema.name).list()
            counts[schema.name] = len(rows)
            columns = sheet_columns(schema)
            widths: list[int] = []
            for col, (header, _spec) in enumerate(columns, start=1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor=self.HEADER_FILL)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                widths.append(len(header))
            has_urls = any(spec is None for _header, spec in columns)
            for r, row in enumerate(rows, start=2):
                link = first_url(row.get("urls")) if has_urls else None
                placeholder = bool(link) and url_kind(link, brain.url_hosts) == "placeholder"
                for col, (_header, spec) in enumerate(columns, start=1):
                    cell = ws.cell(row=r, column=col)
                    if spec is None:  # the derived Work product column: the link to the first URL
                        if not link:
                            self._gap(cell)  # 0.27.0: no URL at all is a gap, shaded, no link
                            gaps["no_url"] += 1
                            continue
                        cell.value = kind_label(link, brain.url_hosts)  # 0.22.0: the kind and the host
                        self._link(cell, link, gap=placeholder)
                        if placeholder:
                            gaps["placeholder"] += 1
                        widths[col - 1] = max(widths[col - 1], min(len(str(cell.value)), 60))
                        continue
                    value = row.get(spec.name)
                    if value is None:
                        if spec.name == "jira_key":  # 0.20.0, the owner: say so instead of an empty cell
                            cell.value = NO_TICKET
                            if has_urls:  # a gap on the Safety Plan only: a meeting or a question has no ticket by nature
                                self._gap(cell)
                                gaps["no_ticket"] += 1
                            widths[col - 1] = max(widths[col - 1], len(NO_TICKET))
                        continue
                    if spec.type is FieldType.BOOL:
                        cell.value = 1 if value else 0
                    elif spec.type is FieldType.DATE:
                        cell.value = dt.date.fromisoformat(value)
                        cell.number_format = "yyyy-mm-dd"
                    elif spec.type is FieldType.URL:
                        cell.value = value
                        cell.hyperlink = value
                        cell.style = "Hyperlink"
                    elif spec.picker == "urls":  # the work product URLs: one per line, read back by the import
                        cell.value = urls_lines(value)
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                    elif spec.type is FieldType.JSON:  # the people of a role: JSON text, read back by the import
                        cell.value = value if isinstance(value, str) else json_text(value)
                    else:
                        cell.value = value
                        if spec.name == DELIVERABLE and has_urls:
                            if link:  # 0.21.0: the deliverable's name links to the first URL, the Hyperlink style
                                self._link(cell, link, gap=placeholder)
                            else:     # 0.27.0: a deliverable that points nowhere is shaded, like the Word export
                                self._gap(cell)
                        if spec.type is FieldType.MULTILINE:
                            cell.alignment = Alignment(wrap_text=True, vertical="top")
                    widths[col - 1] = max(widths[col - 1], min(len(str(value)), 60))
            for col, (_header, spec) in enumerate(columns, start=1):
                if spec is None:
                    ws.column_dimensions[get_column_letter(col)].width = max(6, min(widths[col - 1] + 2, 62))
                    continue
                base = spec.width or widths[col - 1]
                if spec.type is FieldType.BOOL:
                    base = min(base, 14)
                ws.column_dimensions[get_column_letter(col)].width = max(6, min(base + 2, 62))
                if spec.hidden:  # the uid: present in the data, hidden on the sheet, so it reads as before
                    ws.column_dimensions[get_column_letter(col)].hidden = True
                if spec.lookup:
                    values = brain.lookups.get(spec.lookup, ())
                    formula = '"' + ",".join(values) + '"'
                    if values and len(formula) <= 255 and not any("," in v or '"' in v for v in values):
                        dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
                        ws.add_data_validation(dv)
                        letter = get_column_letter(col)
                        dv.add(f"{letter}2:{letter}{max(len(rows) + 200, 500)}")
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = ws.dimensions
        target = Path(out) if out else Path(brain.workbook_path)
        path, warning = _save_with_fallback(wb.save, target)
        return {"format": self.format, "path": str(path), "entities": counts, "gaps": gaps,
                "warnings": [warning] if warning else []}
