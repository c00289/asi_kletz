"""Importers: JSON dumps (round trip) and the Excel workbook (seeding, and edits made in Excel).

0.20.0: the Safety Plan sheet's ``Work product`` column is the export's own (a hyperlink to the first URL) and
is skipped silently; ``Work product URLs`` (one per line) and a Jira key cell reading ``no ticket attached``
are read back by the plan row's ``on_load`` (``WorkProductService.loaded``): the list parsed, the text no key.
"""

from __future__ import annotations

import abc
import json
from pathlib import Path
from typing import TYPE_CHECKING, Any, ClassVar

from ..errors import NotFound
from ..schema import EntitySchema
from ..schema.definitions import WORK_PRODUCT_HEADER

DERIVED_HEADERS = {WORK_PRODUCT_HEADER}  # the export's own columns (0.20.0): not fields, ignored without a word

if TYPE_CHECKING:
    from ..brain import Brain


class Importer(abc.ABC):
    format: ClassVar[str] = ""

    @abc.abstractmethod
    def import_(
        self,
        brain: Brain,
        path: str | Path,
        mode: str = "merge",
        entities: list[str] | None = None,
        actor: str = "import",
    ) -> dict[str, Any]: ...


class ImporterFactory:
    _importers: dict[str, type[Importer]] = {}

    @classmethod
    def register(cls, importer: type[Importer]) -> type[Importer]:
        cls._importers[importer.format] = importer
        return importer

    @classmethod
    def create(cls, fmt: str) -> Importer:
        try:
            return cls._importers[fmt]()
        except KeyError:
            raise NotFound(f"unknown import format {fmt!r}; known: {', '.join(cls._importers)}") from None

    @classmethod
    def formats(cls) -> list[str]:
        return list(cls._importers)


def _require_file(path: str | Path) -> Path:
    file = Path(path)
    if not file.is_file():
        raise NotFound(f"file not found: {file}")
    return file


def _collect(
    brain: Brain, summaries: list[dict[str, Any]], path: Path, fmt: str, extra_warnings: list[str] | None = None
) -> dict[str, Any]:
    """The import's result; every entity is loaded and the calculated fields recomputed by now, so a row that
    carried a calculated value the tool disagrees with is warned about here, once per entity."""
    for summary in summaries:
        warning = brain.service(summary["entity"]).calculated_warning(summary)
        if warning:
            summary["warnings"].append(warning)
    return {
        "format": fmt,
        "path": str(path),
        "entities": {
            s["entity"]: {k: s[k] for k in ("inserted", "updated", "deleted", "skipped")} for s in summaries
        },
        "warnings": [*(extra_warnings or []), *(f"{s['entity']}: {w}" for s in summaries for w in s["warnings"])],
        "errors": [f"{s['entity']}: {e}" for s in summaries for e in s["errors"]],
    }


@ImporterFactory.register
class JsonImporter(Importer):
    format = "json"

    def import_(self, brain, path, mode="merge", entities=None, actor="import"):
        file = _require_file(path)
        payload = json.loads(file.read_text(encoding="utf-8"))
        wanted = {brain.registry.get(e).name for e in entities} if entities else None
        if "entity" in payload:
            sections = {payload["entity"]: payload.get("rows", [])}
        elif "entities" in payload:
            sections = dict(payload["entities"])
        else:
            raise NotFound(f"{file.name} is not an ASI Kletz dump (no 'entity' or 'entities' key)")
        summaries, warnings = [], []
        with brain.schedule.deferred():  # the plan's calculated fields: once, after every entity is in
            for name, rows in sections.items():
                if name not in brain.registry:
                    warnings.append(f"skipped unknown entity {name!r}")
                    continue
                schema = brain.registry.get(name)
                if wanted and schema.name not in wanted:
                    continue
                summaries.append(brain.service(schema.name).load_rows(rows, mode, actor))
        return _collect(brain, summaries, file, self.format, warnings)


def normalise_header(value: Any) -> str:
    return " ".join(str(value if value is not None else "").replace("\xa0", " ").split()).casefold()


def map_columns(schema: EntitySchema, headers: list[Any]) -> tuple[dict[str, int], list[str]]:
    """Match fields to columns by header text, in order, so repeated headers map to successive fields."""
    columns = [(idx, normalise_header(h)) for idx, h in enumerate(headers) if h not in (None, "")]
    used: set[int] = set()
    mapping: dict[str, int] = {}
    for spec in schema.fields:
        target = normalise_header(spec.header)
        for idx, header in columns:
            if idx not in used and header == target:
                mapping[spec.name] = idx
                used.add(idx)
                break
    derived = {normalise_header(h) for h in DERIVED_HEADERS}
    unmapped = [str(headers[idx]).strip() for idx, h in columns if idx not in used and h not in derived]
    return mapping, unmapped


@ImporterFactory.register
class ExcelImporter(Importer):
    format = "excel"

    def import_(self, brain, path, mode="merge", entities=None, actor="import"):
        from openpyxl import load_workbook

        file = _require_file(path)
        wb = load_workbook(file, data_only=True, read_only=True)
        try:
            by_name = {normalise_header(ws.title): ws for ws in wb.worksheets}
            schemas = [brain.registry.get(e) for e in entities] if entities else list(brain.registry)
            summaries, warnings = [], []
            with brain.schedule.deferred():  # the plan's calculated fields: once, after every sheet is in
                for schema in schemas:
                    ws = by_name.get(normalise_header(schema.sheet))
                    if ws is None:
                        warnings.append(f"{schema.name}: sheet {schema.sheet!r} not found in workbook")
                        continue
                    rows_iter = ws.iter_rows(values_only=True)
                    headers = list(next(rows_iter, ()))
                    mapping, unmapped = map_columns(schema, headers)
                    if not mapping:
                        warnings.append(f"{schema.name}: no recognised headers on sheet {schema.sheet!r}")
                        continue
                    if unmapped:
                        warnings.append(f"{schema.name}: ignored column(s) {', '.join(unmapped)}")
                    rows: list[dict[str, Any]] = []
                    for raw in rows_iter:
                        record = {name: (raw[idx] if idx < len(raw) else None) for name, idx in mapping.items()}
                        if all(v is None or (isinstance(v, str) and not v.strip()) for v in record.values()):
                            continue
                        rows.append(record)
                    summaries.append(brain.service(schema.name).load_rows(rows, mode, actor))
        finally:
            wb.close()
        return _collect(brain, summaries, file, self.format, warnings)
