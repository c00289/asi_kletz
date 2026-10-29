from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from asi_kletz import Brain
from tests.conftest import ENTITIES, PLAN, SHEETS, WORKBOOK


def test_json_dump_and_load_modes(seeded, tmp_path):
    out = tmp_path / "dump.json"
    dumped = seeded.execute("dump", out=str(out))
    assert dumped.ok and dumped.data["entities"] == {
        PLAN: 3, "sprint": 0, "sprint_item": 0, "week": 0, "holiday": 0, "parameter": 0,
    }
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 10 and list(payload["entities"]) == ENTITIES

    uids = {r["id"]: r["uid"] for r in payload["entities"][PLAN]}
    assert len(uids) == 3 and all(uids.values())  # the dump carries the hidden identity like any field
    with Brain(tmp_path / "other.db") as other:
        loaded = other.execute("load", path=str(out), mode="replace")
        assert loaded.ok, loaded.errors
        assert other.execute("count", entity=PLAN).data["count"] == 3
        assert {r["id"]: r["uid"] for r in other.execute("list", entity=PLAN).data} == uids  # replace keeps them
        # merge updates the rows the uids name; so does append: a row that exists is never inserted twice
        assert other.execute("load", path=str(out), mode="merge").data["entities"][PLAN]["updated"] == 3
        assert other.execute("count", entity=PLAN).data["count"] == 3
        appended = other.execute("load", path=str(out), mode="append").data["entities"][PLAN]
        assert appended["updated"] == 3 and appended["inserted"] == 0
        assert other.execute("count", entity=PLAN).data["count"] == 3
        # a 0.8.x dump has no uids: append makes new rows with new ids, and stamps them
        for row in payload["entities"][PLAN]:
            del row["uid"]
        old = tmp_path / "dump-0.8.json"
        old.write_text(json.dumps(payload), encoding="utf-8")
        assert other.execute("load", path=str(old), mode="append").data["entities"][PLAN]["inserted"] == 3
        rows = other.execute("list", entity=PLAN, order_by="id").data
        assert [r["id"] for r in rows] == [1, 2, 3, 4, 5, 6] and len({r["uid"] for r in rows}) == 6
        assert all(r["uid"] for r in rows)

    single = tmp_path / "plan.json"
    assert seeded.execute("dump", entity=PLAN, out=str(single)).ok
    assert json.loads(single.read_text(encoding="utf-8"))["entity"] == PLAN


def test_merge_matches_by_uid_before_id(seeded, tmp_path):
    """A dump with every id shifted by 100 merges as three updates and no insert: the uid names the row."""
    out = tmp_path / "dump.json"
    assert seeded.execute("dump", entity=PLAN, out=str(out)).ok
    payload = json.loads(out.read_text(encoding="utf-8"))
    for row in payload["rows"]:
        row["id"] += 100
        row["work_product_location"] = f"moved {row['id']}"
    shifted = tmp_path / "shifted.json"
    shifted.write_text(json.dumps(payload), encoding="utf-8")
    result = seeded.execute("load", path=str(shifted), mode="merge")
    assert result.ok and result.data["entities"][PLAN] == {"inserted": 0, "updated": 3, "deleted": 0, "skipped": 0}
    rows = seeded.execute("list", entity=PLAN).data
    assert [r["id"] for r in rows] == [1, 2, 3]
    assert [r["work_product_location"] for r in rows] == ["moved 101", "moved 102", "moved 103"]
    assert [r["uid"] for r in rows] == [r["uid"] for r in payload["rows"]]
    # a sibling's cached rows take this tool's uids: merged by id, the file's uid replaces the stored one
    with Brain(tmp_path / "sibling.db") as sibling:
        for row in payload["rows"]:
            sibling.execute("create", entity=PLAN, data={"deliverable": row["deliverable"], "order": row["order"]})
        own = [r["uid"] for r in sibling.execute("list", entity=PLAN).data]
        for row in payload["rows"]:
            row["id"] -= 100
        shifted.write_text(json.dumps(payload), encoding="utf-8")
        merged = sibling.execute("load", path=str(shifted), mode="merge").data["entities"][PLAN]
        assert merged == {"inserted": 0, "updated": 3, "deleted": 0, "skipped": 0}
        theirs = [r["uid"] for r in sibling.execute("list", entity=PLAN).data]
        assert theirs == [r["uid"] for r in payload["rows"]] and not set(theirs) & set(own)


def test_excel_export_and_reimport(seeded, tmp_path):
    exported = seeded.execute("export", format="excel")
    assert exported.ok, exported.errors
    path = exported.data["path"]
    wb = load_workbook(path)
    assert wb.sheetnames == SHEETS
    ws = wb["Safety Plan"]
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["order", "PDP Phases", "Safety Lifecycle Phase", "Safety Deliverables"]
    assert headers.count("Tech Lead - Functional Safety") == 2 and len(headers) == 72 and headers[-1] == "uid"
    # 0.20.0: the derived Work product column (the hyperlink) before the Work product URLs field, then the uid
    assert headers[-5:] == ["Status (Pending/ Ongoing/ Blocked/ Alternative/Out-of-Scope/ Done)", "Jira key",
                            "Work product", "Work product URLs", "uid"]
    assert headers[3:7] == ["Safety Deliverables", "Points", "Default points", "Sprint"]  # 0.11.0, after deliverable
    assert ws.freeze_panes == "A2"
    assert ws["D2"].value == "Safety Plan" and ws["K2"].value == 1  # deliverable, DMS flag (K since 0.11.2)
    assert ws.data_validations.dataValidation  # dropdowns from the lookups
    # the uid is the last column of every sheet, hidden, present in the data (3.1 of the hidden-identity plan)
    uids = {r["id"]: r["uid"] for r in seeded.execute("list", entity=PLAN).data}
    assert ws.column_dimensions["BT"].hidden is True and not ws.column_dimensions["BS"].hidden  # uid, 72nd
    assert [ws.cell(row=r, column=72).value for r in (2, 3, 4)] == [uids[1], uids[2], uids[3]]
    for name, width in (("Sprints", 10), ("Sprint Items", 17), ("Weekly Status", 6), ("Holidays", 5),
                        ("Parameters", 6)):
        sheet = wb[name]
        assert sheet.cell(row=1, column=width).value == "uid" and sheet.cell(row=1, column=width + 1).value is None
        assert sheet.column_dimensions[get_column_letter(width)].hidden is True
    items = wb["Sprint Items"]
    assert items.cell(row=1, column=6).value == "Safety Plan row uid" and items.column_dimensions["F"].hidden is True

    with Brain(tmp_path / "re.db") as other:
        result = other.execute("import", format="excel", path=path, mode="replace")
        assert result.ok, result.errors
        rows = other.execute("list", entity=PLAN).data
        assert [r["order"] for r in rows] == [1, 2, 3]
        row = rows[0]
        assert row["deliverable"] == "Safety Plan" and row["loc_dms"] and row["reviewer_tech_lead_functional_safety"]
        assert row["author_tech_lead_functional_safety"] is False and row["status"] == "Done"
        assert {r["id"]: r["uid"] for r in rows} == uids  # export then import on a fresh database keeps every uid
        # a workbook saved by 0.8.x has no uid column: it still imports, and the new rows are stamped
        wb_old = load_workbook(path)
        wb_old["Safety Plan"].delete_cols(72)  # the uid column, last (72nd since 0.20.0)
        old = tmp_path / "plan-0.8.xlsx"
        wb_old.save(old)
        with Brain(tmp_path / "old.db") as third:
            result = third.execute("import", format="excel", path=str(old), mode="replace")
            assert result.ok, result.errors
            assert result.data["entities"][PLAN] == {"inserted": 3, "updated": 0, "deleted": 0, "skipped": 0}
            stamped = {r["uid"] for r in third.execute("list", entity=PLAN).data}
            assert len(stamped) == 3 and all(stamped) and not stamped & set(uids.values())


def test_locked_workbook_falls_back_to_sibling(seeded, tmp_path, monkeypatch):
    import asi_kletz.io.exporters as ex

    calls = []

    def fake_save(path):
        calls.append(path)
        if len(calls) == 1:
            raise PermissionError("locked")
        path.write_bytes(b"ok")

    path, warning = ex._save_with_fallback(fake_save, tmp_path / "view.xlsx", ".xlsx")
    assert path != tmp_path / "view.xlsx" and path.suffix == ".xlsx" and "locked" in warning


def test_run_script(seeded, tmp_path):
    script = tmp_path / "steps.jsonl"
    script.write_text(
        '# comment\n'
        '{"verb": "create", "args": {"entity": "safety_plan_row", "data": {"deliverable": "from script"}}}\n'
        '{"verb": "count", "args": {"entity": "safety_plan_row"}}\n'
        '{"verb": "get", "args": {"entity": "safety_plan_row", "id": 999}}\n',
        encoding="utf-8",
    )
    result = seeded.execute("run", path=str(script))
    assert not result.ok and result.data["steps"] == 3 and result.data["failed"] == [3]
    assert result.data["results"][1]["data"]["count"] == 4
    assert seeded.storage.audit_tail(3)[-1]["actor"] == "python"
    stop = seeded.execute("run", path=str(script), stop_on_error=True)
    assert stop.data["steps"] == 3


@pytest.mark.skipif(not WORKBOOK.is_file(), reason="real workbook not present")
def test_seed_from_real_workbook(tmp_path):
    with Brain(tmp_path / "seed.db") as b:
        result = b.execute("import", format="excel", path=str(WORKBOOK), mode="replace")
        assert result.ok, result.errors
        counts = b.execute("status").data["entities"]
        assert list(counts) == ENTITIES and counts[PLAN] >= 48
        plan = b.execute("list", entity=PLAN, filter=["order=1"]).data[0]
        assert plan["lifecycle_phase"] == "1. Concept Development" and plan["ref_qf_17"] is True
        out = b.execute("export", format="excel", out=str(tmp_path / "regen.xlsx"))
        assert out.ok and out.data["entities"][PLAN] == counts[PLAN]
        assert load_workbook(out.data["path"]).sheetnames == SHEETS


@pytest.mark.skipif(not WORKBOOK.is_file(), reason="real workbook not present")
def test_source_workbook_columns_match_the_schema():
    """The table shows the sheet's columns in source order by construction, not by coincidence."""
    from asi_kletz.schema import default_registry
    from asi_kletz.schema.definitions import WORK_PRODUCT_HEADER

    ws = load_workbook(WORKBOOK, read_only=True)["Safety Plan"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    while headers and headers[-1] is None:
        headers.pop()
    # the owner's original sheet has the 63 source columns (no points, default or sprint); once the tool has
    # regenerated it (export excel) it carries every visible field, and the hidden uid column last, which is
    # neither; the table shows whichever in source order by construction
    if headers and headers[-1] == "uid":
        headers.pop()
    headers = [h for h in headers if h != WORK_PRODUCT_HEADER]  # 0.20.0: the export's own column, no field
    schema = default_registry.get(PLAN)
    source = [f.header for f in schema.source_fields]
    visible = [f.header for f in schema.visible_fields]
    # the owner's sheet, or a view regenerated by any 0.11.x: its columns are visible fields, in schema order
    assert headers == source or (set(headers) <= set(visible) and headers == [h for h in visible if h in set(headers)])
    assert headers[:4] == source[:4] and len(source) == 63 and set(source) < set(visible)


def test_the_files_carry_points_and_the_import_recalculates_sprint_and_deadline(seeded, tmp_path):
    """0.11.0: points and defaults travel; sprint and deadline are the tool's, warned about once when a file
    disagrees; a Done row keeps the file's deadline as history."""
    from tests.conftest import SEED_WEIGHTS

    assert seeded.execute("create", entity="sprint", data={
        "number": 2, "label": "SPRINT 2", "start": "2026-10-27", "end": "2026-11-09", "status": "Active",
    }).ok
    assert seeded.execute("plan", action="weights", what="load", file=str(SEED_WEIGHTS)).ok
    assert seeded.execute("update", entity=PLAN, id=2, set=["points=4"]).ok
    rows = seeded.execute("list", entity=PLAN).data
    assert [(r["points"], r["default_points"], r["sprint"], r["deadline"]) for r in rows] == [
        (None, 5, None, "Phase 1 exit"), (4, 3, 2, "2026-11-09"), (None, 2, 2, "2026-11-09"),
    ]
    dump = tmp_path / "plan.json"
    assert seeded.execute("dump", entity=PLAN, out=str(dump)).ok
    dumped = json.loads(dump.read_text(encoding="utf-8"))["rows"]
    assert all(k in dumped[0] for k in ("points", "default_points", "sprint", "deadline"))  # JSON carries everything
    assert dumped[1]["sprint"] == 2 and dumped[1]["deadline"] == "2026-11-09"
    workbook = seeded.execute("export", format="excel", out=str(tmp_path / "plan.xlsx")).data["path"]
    ws = load_workbook(workbook)["Safety Plan"]
    assert [ws.cell(row=3, column=c).value for c in (5, 6, 7, 67)] == [4, 3, 2, "2026-11-09"]  # a TEXT field
    assert ws.cell(row=2, column=67).value == "Phase 1 exit"  # the Done row's history, in the deadline column

    with Brain(tmp_path / "fresh.db") as other:
        # the workbook brings its Sprints sheet: the chain arrives with the plan, the recompute runs once after
        # every sheet, and the file's sprint and deadline agree with it, so nothing is warned about
        result = other.execute("import", format="excel", path=workbook, mode="replace")
        assert result.ok, result.errors
        assert result.warnings == []
        rows = other.execute("list", entity=PLAN).data
        assert [(r["points"], r["default_points"], r["sprint"], r["deadline"]) for r in rows] == [
            (None, 5, None, "Phase 1 exit"), (4, 3, 2, "2026-11-09"), (None, 2, 2, "2026-11-09"),
        ]
        # a dump whose rows carry another sprint and deadline: warned about once, the calculated values stay
        for row in dumped[1:]:
            row["sprint"], row["deadline"] = 99, "2099-01-01"
        dumped[0]["deadline"] = "Phase 1 exit (revised)"  # a Done row's deadline is history: taken as is
        forged = tmp_path / "forged.json"
        forged.write_text(json.dumps({"entity": PLAN, "rows": dumped}), encoding="utf-8")
        again = other.execute("load", path=str(forged), mode="merge")
        assert again.ok and again.warnings == [
            f"{PLAN}: 2 row(s) carried a sprint, sprint start, sprint end or deadline that differs from the "
            "calculated value; the file's values were ignored"
        ]
        rows = other.execute("list", entity=PLAN).data
        assert [(r["sprint"], r["deadline"]) for r in rows] == [
            (None, "Phase 1 exit (revised)"), (2, "2026-11-09"), (2, "2026-11-09"),
        ]
        same = other.execute("load", path=str(dump), mode="merge")  # the original dump agrees: no warning
        assert same.ok and same.warnings == []
        with_sheet = other.execute("import", format="excel", path=workbook, mode="merge", entity=PLAN)
        assert with_sheet.ok and with_sheet.warnings == []  # one sheet of the same workbook: still agrees
