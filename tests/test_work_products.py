"""0.20.0, the owner, 2026-10-29: "please add the ability to add url parameters to each deliverable... each
deliverable may have multiple url... make sure that the excel export includes links to the work products
wherever they may be... on multiple urls just go to first; list of urls is sequential while lists of names are
alphabetical... include links to work products in the tool too... use microsoft homepage as placeholder url
for all".

The work product URLs of a plan row: the pure functions, the verbs (the API's array, the CLI's text), the
validation, the standards' ``row_N_url`` / ``row_N_urls``, the Excel export's Work product columns (the
hyperlink to the first, the list one per line, ``no ticket attached`` for a row without a key) and its import,
the dump round trip, the placeholder verb and its dry run, and the schema 9 to 10 migration."""

from __future__ import annotations

import json
import sqlite3

import pytest
from openpyxl import load_workbook

from asi_kletz import Brain
from asi_kletz.schema.definitions import NO_TICKET, WORK_PRODUCT_HEADER, WORK_PRODUCT_URLS_HEADER
from asi_kletz.schema.urls import (
    UrlsError,
    first_url,
    is_url,
    parse_urls,
    url_host,
    urls_json,
    urls_lines,
    urls_of,
    urls_text,
)
from asi_kletz.services.standards_service import ROW_PARAM
from tests.conftest import PLAN
from tests.test_cli import run
from tests.test_standards import standards_dir

A = "https://dms.example.com/doc/17"
B = "http://jama.example.com/items/2554"
MS = "https://www.microsoft.com/"
LINKS_SOURCE = """---
title: "Links (test)"
kind: plan
doc_id: QF-17
---
# Links (test)

| Deliverable | Work product |
| :--- | :--- |
| Safety Plan | {{ row_1_url }}<br>{{ row_1_urls }} |
| HARA | {{ row_2_url }} |
"""


@pytest.fixture
def standards(tmp_path, monkeypatch):
    root = standards_dir(tmp_path, monkeypatch)
    (root / "qf-17-links").mkdir()
    (root / "qf-17-links" / "qf-17-links.md").write_text(LINKS_SOURCE, encoding="utf-8")
    return root


# -- the pure functions -------------------------------------------------------------------------------------


def test_parse_urls_reads_every_form_and_keeps_the_order():
    assert parse_urls([B, A]) == [B, A]  # the API's array: sequential, never sorted
    assert parse_urls(json.dumps([B, A])) == [B, A]  # its JSON text
    assert parse_urls(f"{B}; {A}") == [B, A]  # the CLI's text, `;` between URLs
    assert parse_urls(f"{B}\n{A}\n") == [B, A]  # one per line
    assert parse_urls(f"  {B} ;; ") == [B]
    assert parse_urls(None) == [] and parse_urls("") == [] and parse_urls("[]") == []
    assert urls_json([B, A]) == json.dumps([B, A], separators=(", ", ": ")) and urls_json([]) is None
    assert urls_of([B, A]) == [B, A] and urls_of("not a url") == [] and urls_of(None) == []
    assert first_url([B, A]) == B and first_url(None) is None and first_url("[]") is None
    assert urls_text([B, A]) == f"{B}; {A}" and urls_lines([B, A]) == f"{B}\n{A}"
    assert url_host(MS) == "www.microsoft.com" and url_host(B) == "jama.example.com"
    assert is_url(A) and is_url("HTTP://x.io") and not is_url("ftp://x.io") and not is_url("x.io/path")
    assert not is_url("https://") and not is_url("https://a b.io") and not is_url(None)


@pytest.mark.parametrize(("value", "message"), [
    ("ftp://x.io/a", "urls, line 1 ('ftp://x.io/a'): not an absolute http(s) URL; write an absolute http(s) URL, "
                     "one per line"),
    ("dms/doc/17", "urls, line 1 ('dms/doc/17'): not an absolute http(s) URL; write an absolute http(s) URL, one "
                   "per line"),
    ([A, A], f"urls: {A!r} is listed twice (lines 1 and 2)"),
    ([A, None], "urls, line 2: an empty entry; write an absolute http(s) URL, one per line"),
    ("[oops", "urls: not a JSON array of URLs: Expecting value: line 1 column 2 (char 1)"),
    ('{"a": 1}', "urls, line 1 ('{\"a\": 1}'): not an absolute http(s) URL; write an absolute http(s) URL, one per "
                 "line"),
])
def test_parse_urls_refuses_and_names_the_line(value, message):
    with pytest.raises(UrlsError) as info:
        parse_urls(value)
    assert info.value.errors == [message]


def test_parse_urls_reports_every_problem_at_once():
    with pytest.raises(UrlsError) as info:
        parse_urls(f"{A}; nope; {A}", "urls")
    assert info.value.errors == [
        "urls, line 2 ('nope'): not an absolute http(s) URL; write an absolute http(s) URL, one per line",
        f"urls: {A!r} is listed twice (lines 1 and 3)",
    ]


# -- the verbs ------------------------------------------------------------------------------------------------


def test_a_row_holds_its_urls_in_order_as_an_array(seeded):
    updated = seeded.execute("update", entity=PLAN, id=1, data={"urls": [B, A]})
    assert updated.ok and updated.warnings == [] and updated.data["urls"] == [B, A]  # an array, in order
    assert seeded.execute("get", entity=PLAN, id=1).data["urls"] == [B, A]
    assert seeded.execute("update", entity=PLAN, id=1, set=[f"urls={A}; {B}"]).data["urls"] == [A, B]  # the text
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": f"{B}\n{A}"}).data["urls"] == [B, A]
    created = seeded.execute("create", entity=PLAN, data={"order": 4, "deliverable": "TSC", "urls": [MS]})
    assert created.ok and created.data["urls"] == [MS]
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": ""}).data["urls"] is None  # cleared
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": []}).data["urls"] is None
    # the audit carries the stored JSON; the stored text is the canonical array
    raw = seeded.storage.fetchone("SELECT urls FROM safety_plan WHERE id = ?", (4,))
    assert raw["urls"] == json.dumps([MS], separators=(", ", ": "))


def test_the_urls_rule_refuses_what_it_must(seeded):
    bad = seeded.execute("update", entity=PLAN, id=1, data={"urls": ["ftp://x.io", A, A]})
    assert not bad.ok and bad.errors == [
        "urls, line 1 ('ftp://x.io'): not an absolute http(s) URL; write an absolute http(s) URL, one per line",
        f"urls: {A!r} is listed twice (lines 2 and 3)",
    ]
    assert seeded.execute("get", entity=PLAN, id=1).data["urls"] is None  # nothing changed
    bad = seeded.execute("create", entity=PLAN, data={"order": 9, "deliverable": "x", "urls": "not a url"})
    assert not bad.ok and "not an absolute http(s) URL" in bad.errors[0]
    assert seeded.execute("count", entity=PLAN).data["count"] == 3


def test_the_no_ticket_text_is_no_key(seeded):
    """The workbook prints ``no ticket attached`` for a row without a key; typed or imported, it is no key."""
    typed = seeded.execute("update", entity=PLAN, id=1, data={"jira_key": "No Ticket  Attached"})
    assert typed.ok and typed.data["jira_key"] is None
    assert seeded.execute("update", entity=PLAN, id=1, data={"jira_key": "HWL-1"}).data["jira_key"] == "HWL-1"
    assert seeded.execute("update", entity=PLAN, id=1, set=[f"jira_key={NO_TICKET}"]).data["jira_key"] is None


def test_the_cli_sets_urls_as_text_and_as_json(tmp_path, capsys):
    db = str(tmp_path / "cli.db")
    code, res = run(capsys, "--db", db, "create", PLAN, "--set", "deliverable=Safety Plan", "--set", "order=1")
    assert code == 0
    code, res = run(capsys, "--db", db, "update", PLAN, "1", "--set", f"urls={A}; {B}")
    assert code == 0 and res["data"]["urls"] == [A, B]
    code, res = run(capsys, "--db", db, "update", PLAN, "1", "--data", json.dumps({"urls": [B]}))
    assert code == 0 and res["data"]["urls"] == [B]
    code, res = run(capsys, "--db", db, "update", PLAN, "1", "--set", "urls=nope")
    assert code == 1 and "not an absolute http(s) URL" in res["errors"][0]
    code, res = run(capsys, "--db", db, "plan", "urls", "--placeholder", MS, "--check")
    assert code == 0 and res["data"]["check"] is True and res["data"]["would_fill"] == []
    code, res = run(capsys, "--db", db, "create", PLAN, "--set", "deliverable=HARA", "--set", "order=2")
    code, res = run(capsys, "--db", db, "plan", "urls", "--placeholder", MS)
    assert code == 0 and res["data"]["filled"] == [2] and res["data"]["already"] == 1
    code, res = run(capsys, "--db", db, "list", PLAN)
    assert [r["urls"] for r in res["data"]] == [[B], [MS]]


# -- the standards ---------------------------------------------------------------------------------------------


def test_the_standards_get_the_first_url_and_all_of_them(seeded, standards):
    assert ROW_PARAM.match("row_1_url") and ROW_PARAM.match("row_12_urls") and not ROW_PARAM.match("row_1_link")
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [B, A]}).ok
    rows = seeded.execute("list", entity=PLAN).data
    values = seeded.standards.row_values(rows, {}, [])
    assert values["row_1_url"] == B and values["row_1_urls"] == f"{B}\n{A}"  # the first is the link; all, a line each
    assert "row_2_url" not in values and "row_2_urls" not in values  # a row with none leaves the blanks
    listed = seeded.execute("standards", action="list").data["standards"]
    links = next(s for s in listed if s["folder"] == "qf-17-links")
    assert links["row_parameters"] == 3 and links["document_parameters"] == []
    preview = seeded.execute("plan", action="parameters", row=1, standard="qf-17-links")
    assert preview.ok, preview.errors
    assert preview.data["standards"][0]["values"] == {"row_1_url": B, "row_1_urls": f"{B}\n{A}"}
    assert preview.data["standards"][0]["unfilled"] == ["row_2_url"]


# -- the files -------------------------------------------------------------------------------------------------


def test_the_excel_export_links_the_work_product_and_says_no_ticket_attached(seeded, tmp_path):
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS, A], "jira_key": "HWL-2554"}).ok
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [B]}).ok
    exported = seeded.execute("export", format="excel", out=str(tmp_path / "view.xlsx"))
    assert exported.ok, exported.errors
    ws = load_workbook(exported.data["path"])["Safety Plan"]
    headers = [c.value for c in ws[1]]
    key_col = headers.index("Jira key") + 1
    link_col, urls_col = headers.index(WORK_PRODUCT_HEADER) + 1, headers.index(WORK_PRODUCT_URLS_HEADER) + 1
    assert (link_col, urls_col) == (key_col + 1, key_col + 2) and headers[urls_col] == "uid"
    # row 1: the Work product cell is a hyperlink to the first URL, its text the kind and host (0.22.0: the
    # placeholder reads `placeholder` alone), the Hyperlink style
    link = ws.cell(row=2, column=link_col)
    assert link.value == "placeholder" and link.hyperlink.target == MS and link.style == "Hyperlink"
    assert ws.cell(row=2, column=urls_col).value == f"{MS}\n{A}"  # all of them, one per line
    assert ws.cell(row=2, column=key_col).value == "HWL-2554"
    assert ws.cell(row=3, column=link_col).value == "other, jama.example.com"
    assert ws.cell(row=3, column=link_col).hyperlink.target == B
    assert ws.cell(row=3, column=urls_col).value == B
    assert ws.cell(row=3, column=key_col).value == NO_TICKET  # no ticket: the text, not an empty cell
    # row 3: no URL at all, both cells empty; no ticket either
    assert ws.cell(row=4, column=link_col).value is None and ws.cell(row=4, column=link_col).hyperlink is None
    assert ws.cell(row=4, column=urls_col).value is None and ws.cell(row=4, column=key_col).value == NO_TICKET
    # the import reads the list back in order, the no-ticket text as no key, and never names the derived column
    with Brain(tmp_path / "re.db") as other:
        imported = other.execute("import", format="excel", path=exported.data["path"], mode="replace")
        assert imported.ok, imported.errors
        assert not any(WORK_PRODUCT_HEADER in w for w in imported.warnings)
        rows = other.execute("list", entity=PLAN).data
        assert [r["urls"] for r in rows] == [[MS, A], [B], None]
        assert [r["jira_key"] for r in rows] == ["HWL-2554", None, None]


def test_urls_round_trip_through_the_dump_and_a_bad_list_is_dropped_with_a_warning(seeded, tmp_path):
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [B, A]}).ok
    dump = tmp_path / "dump.json"
    assert seeded.execute("dump", out=str(dump), entity=PLAN).ok
    payload = json.loads(dump.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 10 and payload["rows"][0]["urls"] == [B, A]
    payload["rows"][1]["urls"] = ["ftp://x.io", A]  # bad: dropped, warned, the row still loads
    payload["rows"][2]["urls"] = f"{MS}; {A}"  # the text form in a file: parsed too
    dump.write_text(json.dumps(payload), encoding="utf-8")
    with Brain(tmp_path / "re.db") as other:
        loaded = other.execute("load", path=str(dump), mode="replace")
        assert loaded.ok, loaded.errors
        rows = other.execute("list", entity=PLAN).data
        assert [r["urls"] for r in rows] == [[B, A], None, [MS, A]]
        assert any("('ftp://x.io'): not an absolute http(s) URL" in w and "urls ignored" in w for w in loaded.warnings)


# -- the placeholder --------------------------------------------------------------------------------------------


def test_plan_urls_fills_the_rows_without_a_url_once(seeded):
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [B, A]}).ok
    before = len(seeded.storage.audit_tail(50))
    check = seeded.execute("plan", action="urls", placeholder=MS, check=True)
    assert check.ok and check.data["check"] is True
    assert check.data["would_fill"] == [1, 3] and check.data["filled"] == []
    assert check.data["already"] == 1 and [r["id"] for r in check.data["rows_filled"]] == [1, 3]
    assert [r["urls"] for r in seeded.execute("list", entity=PLAN).data] == [None, [B, A], None]  # nothing written
    assert len(seeded.storage.audit_tail(50)) == before
    filled = seeded.execute("plan", action="urls", placeholder=MS)
    assert filled.ok and filled.data["filled"] == [1, 3] and filled.data["would_fill"] == []
    assert [r["urls"] for r in seeded.execute("list", entity=PLAN).data] == [[MS], [B, A], [MS]]  # row 2 untouched
    entry = seeded.storage.audit_tail(1)[0]
    assert entry["action"] == "urls_placeholder" and entry["payload"] == {"url": MS, "rows": [1, 3]}
    again = seeded.execute("plan", action="urls", placeholder=MS)
    assert again.ok and again.data["filled"] == [] and again.data["already"] == 3  # idempotent
    assert seeded.storage.audit_tail(1)[0] == entry
    assert seeded.execute("plan", action="urls", placeholder="microsoft.com").errors == [
        "'microsoft.com' is not an absolute http(s) URL"
    ]
    assert seeded.execute("plan", action="urls").errors == [
        "plan urls needs --placeholder <url>, --audit or --replace <old-url> <new-url>"
    ]


# -- the migration ----------------------------------------------------------------------------------------------


def test_a_schema_9_database_gains_the_urls_column_and_keeps_every_row(tmp_path, seeded):
    db = tmp_path / "kletz.db"
    with Brain(db) as before:
        before.execute("create", entity=PLAN, data={"order": 1, "deliverable": "Safety Plan", "jira_key": "HWL-1"})
        before.execute("create", entity="parameter", data={"name": "vehicle_type", "value": "Wheel Loader"})
        rows_before = before.execute("list", entity=PLAN).data
    conn = sqlite3.connect(db)
    conn.execute("ALTER TABLE safety_plan DROP COLUMN urls")  # what a 0.19.0 file looks like
    conn.commit()
    conn.close()
    with Brain(db) as after:
        assert after.migration["created"] == [] and after.migration["altered"] == {"safety_plan": ["urls"]}
        assert after.status()["schema_version"] == 10
        rows = after.execute("list", entity=PLAN).data
        assert [{**r, "urls": None} for r in rows_before] == rows and rows[0]["urls"] is None
        assert after.storage.audit_tail(1)[0]["action"] != "update"  # the migration edits no row
        assert after.execute("update", entity=PLAN, id=1, data={"urls": [MS]}).data["urls"] == [MS]


def test_the_excel_deliverable_cell_links_to_the_first_url(seeded, tmp_path):
    """0.21.0, the owner: "turn the name of the safety deliverable into a clickable URL". The Safety
    Deliverables cell of a row with a URL is a hyperlink to the first one, in the Hyperlink style, its text
    the name as before; a row without a URL keeps a plain cell; the import reads the names back."""
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS, A]}).ok
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [B]}).ok
    exported = seeded.execute("export", format="excel", out=str(tmp_path / "view.xlsx"))
    assert exported.ok, exported.errors
    ws = load_workbook(exported.data["path"])["Safety Plan"]
    headers = [c.value for c in ws[1]]
    name_col = headers.index("Safety Deliverables") + 1
    first, second, third = (ws.cell(row=r, column=name_col) for r in (2, 3, 4))
    assert first.value == "Safety Plan" and first.hyperlink.target == MS and first.style == "Hyperlink"
    assert second.hyperlink.target == B and second.style == "Hyperlink" and second.value
    assert third.value and third.hyperlink is None and third.style != "Hyperlink"
    assert first.alignment.wrap_text  # still the multiline cell it was
    with Brain(tmp_path / "re.db") as other:
        imported = other.execute("import", format="excel", path=exported.data["path"], mode="replace")
        assert imported.ok, imported.errors
        rows = other.execute("list", entity=PLAN).data
        assert rows[0]["deliverable"] == "Safety Plan" and [r["urls"] for r in rows] == [[MS, A], [B], None]
