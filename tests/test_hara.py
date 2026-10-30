"""The HARA (0.28.0): the nine entities and their rules, the workbook reader, the one computation (the MPLr of
every assessment row from the rating table and the vehicle code), the derived Summary, the check against the
workbook's typed values, the export in the workbook's shape, the verb, the screen.

Everything here runs on a **synthetic** workbook with neutral invented names (a gate carrier, a lift, a stop),
written by openpyxl in the exact sheet shapes of the owner's HARA workbook as the dissection in Roy's Brain
describes them: the Document Information pairs, the Function List with an annex letter per column, the Use
Case List naming the vehicles, the rating table with its letters in the row above the header, the context
sheets of both generations, the Hazardous Events with "Role: Name" approver columns, the ids cut short by
``LEFT``. No real HARA text is here.
"""

from __future__ import annotations

import datetime as dt
import gc
import json
import tkinter as tk
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from asi_kletz import Brain
from asi_kletz.automation.cli import main
from asi_kletz.errors import ValidationError
from asi_kletz.hara import DEFAULT_FORMATS, FORMATS, read_hara_workbook
from asi_kletz.hara.workbook_reader import guess_vehicle_code
from asi_kletz.schema.definitions import CUMULATIVE_CODE, HARA_CONTEXTS, HARA_ENTITIES
from asi_kletz.services.hara_service import (
    HARA_ENTITY,
    NOT_FOUND,
    PART_ENTITIES,
    equivalents,
    highest,
    norm_key,
    norm_mplr,
    ratings_of,
    split_event,
)
from asi_kletz.version import SCHEMA_VERSION
from tests.conftest import DB_0_23_0, EMPTY_COUNTS, HARA_TABLES

ENTITIES = [HARA_ENTITY, *PART_ENTITIES]
LETTERS = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"
VEHICLES = {"A": "Tiny Carrier", "B": "Big Carrier", "G": "Gate Carriers(>=2000kgs)", "I": "Small Carrier"}
TITLE = "Hazard Analysis and Risk Assessment Report (HARA)"
GOAL_1 = "Carrier shall keep its distance from persons"
GOAL_2 = "Carrier shall not strike the gate"
ASSESSMENT_HEADER = [
    "Sr.#",
    "Hazard ID",
    "Hazards",
    "Scenarios ID ",
    "Scenarios",
    "Hazardous Events ID",
    "Hazardous Events",
    "Functions",
    "Use Cases",
    "Failure Mode",
    "MPLr",
    "Cummulative Hazard MPLr",
    "Comments",
    "SG Global ID",
    "Safety Constraint(Safety Goal)",
    "SG MPLr",
    "MTTFd???",
]
# (event, function, use case, failure mode, the MPLr the workbook's formula answered, cumulative, sg mplr)
BF_BUC = [
    ("H-1.1.1", "lift", "Travel/Roading *", "Uncommanded Activation", "c", "d", "d"),
    ("H-1.1.1", "stop", "Travel/Roading *", "Failure to apply on Demand", "d", None, None),
    ("H-1.1.2", "lift", "Travel/Roading", "Failure to apply on Demand", "-", None, None),  # no asterisk typed
    ("H-2.1.1", "lift", "Travel/Roading *", "Uncommanded Activation", "d", "d", "d"),  # the cumulative column's d
]
WIF_BUC = [
    ("H-1.1.1", "tilt", "Gating", "Uncommanded Activation", "b", None, None),
    ("H-2.1.1", "tilt", "Travel/Roading *", "Uncommanded Activation", "-", None, None),
]
COUNTS = {
    "vehicles": 3,
    "functions": 3,
    "use_cases": 2,
    "failure_modes": 2,
    "assumptions": 2,
    "ratings": 4,
    "assessments": 6,
    "approvals": 2,
}
STEP1 = {"losses": 2, "hazards": 8, "constraints": 2}


def _tk_available() -> bool:
    try:
        root = tk.Tk()
        root.withdraw()
        root.destroy()
        del root
        gc.collect()
        return True
    except tk.TclError:
        return False


def synthetic_workbook(path: Path) -> Path:
    """The synthetic HARA workbook, in the owner's sheet shapes."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Cover")
    ws["B10"] = TITLE
    ws["B13"] = "Gate Carrier ASAM"
    ws["H35"] = "RA1"
    ws = wb.create_sheet("Dashboard")
    ws["B2"] = "Dashboard"
    ws = wb.create_sheet("Document Information")
    ws.append([])
    ws.append([None, "Document Information", "1. System Information", "1. System Information", "1. System Information"])
    ws.append([])
    ws.append([None, "System Name", "Gate Carrier ASAM", "Company", "Example Robots"])
    ws.append([None, "Department", "Functional Safety Team", "Responsible", "Pat Example"])
    ws.append([None, "Hazard Analysis date", dt.datetime(2026, 1, 15), "Supporting document", "Hazard Analysis (STPA)"])
    ws.append([None, "Project", "Carrier Platform", "", "Risk Assessment- ISO 19014"])
    ws.append([None, "Vehicle code", None, "Vehicle Name", "Gate Carrier"])
    ws.append([])
    ws.append([None, "Description", "Description", "Description", "Description"])
    ws.append(
        [None, "System Description\nA synthetic carrier moving through gates on a fenced site.", None, None, None]
    )
    ws.append([])
    ws.append([None, "Worksheet", "Description"])
    ws.append([None, "HARA (BF-BUC)", "HARA for Base Functions and Base Use Cases"])
    ws = wb.create_sheet("Vocabulary")
    ws.append([])
    ws.append(["Sr#", "Type", "Terms", "Definition"])
    ws.append([1, "Function", "lift", "Raising the carrier's load."])
    ws.append([2, "Use case", "Gating", "Passing through a gate."])
    ws = wb.create_sheet("Scope")
    ws.append([])
    ws.append([None, "Scope: ", "This risk assessment covers the carriers of the example projects.", None, None])
    ws.append([])
    ws.append([])
    ws.append([None, "Annex", "Vehicles on ISO 19014", "ASI Vehicles ", "ASI Projects"])
    ws.append([None, "B", VEHICLES["B"], None, None])
    ws.append([None, "G", VEHICLES["G"], "x (GC-1)", "Gates"])
    ws.append([None, "I", VEHICLES["I"], "In Future", None])
    ws = wb.create_sheet("Function List")
    ws.append([None, "As per ISO 19014"])
    ws.append([None, None, None, "MPLr"])
    ws.append([None, "Sr.#", "Annex", *LETTERS])
    names = [None] * (len(LETTERS))
    for letter, name in VEHICLES.items():
        names[LETTERS.index(letter)] = name
    ws.append([None, None, "Base Functions", *names, "Foundation Vehicle Cummulative", "Gate Carrier Specific MPLr"])

    def marks(*letters: str) -> list:
        return ["x" if letter in letters else None for letter in LETTERS]

    ws.append([None, 1, "lift", *marks("B", "G"), "c", "b"])
    ws.append([None, 2, "stop", *marks("G", "I"), "d", "c"])
    ws.append([None, None, "Work Implement Functions", *[None] * len(LETTERS)])
    ws.append([None, 1, "tilt", *marks("G"), "b", None])
    ws = wb.create_sheet("Use Case List")
    ws.append([])
    ws.append([None, "Vehicles", None, VEHICLES["B"], VEHICLES["G"], VEHICLES["I"]])
    ws.append([None, "Usecases as per ISO 19014"])
    ws.append([None, "Sr.#", "Use Cases (ISO19014)", "Sites", "Gates", "Sheds"])
    ws.append([None, 1, "Travel/Roading *", "x", "x", "x"])
    ws.append([None, 2, "Gating", None, "x", None])
    ws.append([None, "*", "Applies to Base Functions"])
    ws = wb.create_sheet("Failure modes")
    ws.append([])
    ws.append([None, "As per ISO 19014-5"])
    ws.append([None, "Sr.#", "Failure Modes"])
    ws.append([None, 1, "Failure to apply on Demand"])
    ws.append([None, 2, "Uncommanded Activation"])
    ws = wb.create_sheet("Assumption List")
    ws.append([None, None, "As per ISO 19014", None, None, None, "As per ASI"])
    ws.append([None, "Sr#", "Vehicles", "Assumption", None, "Sr#", "Vehicles", "Assumption"])
    ws.append([None, 1, VEHICLES["B"], "Sites are fenced.", None, 1, "All", "Power is off when parked."])
    ws = wb.create_sheet("Loss")
    ws.append([])
    ws.append([None, "1.1 Loss", "1.1 Loss", "1.1 Loss", "1.1 Loss"])
    ws.append([])
    ws.append([None, "ID", "Loss", "Detail", "Relation"])
    ws.append([None, "L-1", "Death or injury", "People near the carrier", "H-1"])
    ws.append([None, "L-2", "Damage to the gate", "Property", "H-2"])
    ws = wb.create_sheet("Hazard List")
    ws.append([])
    ws.append([None, "1.2 Hazard", "1.2 Hazard", "1.2 Hazard", "1.2 Hazard"])
    ws.append([])
    ws.append([None, "ID", "Hazard", "Detail", "Relation"])
    ws.append([None, "H-1", "Carrier too close to persons", "Physical State: too close", "L-1", "OK"])
    ws.append([None, "H-2", "Carrier damages the gate", "Physical State: contact", "L-2", "NON FuSa"])
    ws = wb.create_sheet("Scenarios List")
    ws.append(["Scenario ID", "Scenario", "Relations", "NA"])
    ws.append(["H-1.1", "H-1.1: Direct interactions", "H-1", None])
    ws.append(["H-2.1", "H-2.1: Gate strike", "H-2", "x"])
    ws = wb.create_sheet("Hazardous Events")
    ws.append(
        [
            "Hazardous Event ID",
            "Hazardous Event",
            "Relation",
            "Carrier Approval",
            "Robot Chief: Pat Example",
            "Software Lead: Sam Example",
            "Comments (Use Initials)",
        ]
    )
    ws.append(["H-1.1.1", "H-1.1.1: Carrier approaches a person", "H-1.1", None, "Y", "N", "PE: check"])
    ws.append(["H-1.1.2", "H-1.1.2: Carrier fails to see a person", "H-1.1", None, None, None, None])
    ws.append(["H-1.1.1", "H-1.1.10: Carrier reverses into a person", "H-1.1", None, None, None, None])  # LEFT cut it
    ws.append(["H-2.1.1", "H-2.1.1: Carrier strikes the gate", "H-2.1.", None, None, None, None])  # a typo in Relation
    ws = wb.create_sheet("RA(ISO 19014)")
    ws.append(["Sr.#", "Annex", None, None, "B", "G", "I", None, None])
    ws.append(
        [
            None,
            "Base Functions",
            "Use Case",
            "Failure type",
            VEHICLES["B"],
            VEHICLES["G"],
            VEHICLES["I"],
            "Foundation Vehicle Cummulative",
            "Comments",
        ]
    )
    ws.append([1, "lift", "Travel/Roading *", "Uncommanded Activation", "b", "c", "a", "d", "a note"])
    ws.append([2, "lift", "Travel/Roading *", "Failure to apply on Demand", None, "b", None, "c", None])
    ws.append([3, "stop", "Travel/Roading *", "Failure to apply on Demand", "c", "d", "c", "d", None])
    ws.append([4, "tilt", "Gating", "Uncommanded Activation", None, "b", None, "c", None])
    ws.append([5, "stop", None, None, "c", "c", "c", "c", None])  # a function-level row: no key
    ws = wb.create_sheet("RA(for all vehicles)")
    ws.append(["Sr.#", "Annex", None, None, None, None, None, "Vehicle (PLr)"])
    ws.append(
        [
            None,
            "Base Functions",
            "Use Case",
            "Failure type",
            VEHICLES["A"],
            VEHICLES["B"],
            VEHICLES["G"],
            VEHICLES["I"],
            "Foundation Vehicle Cummulative",
            "Comments",
        ]
    )
    ws.append([1, "lift", "Travel/Roading", "Uncommanded Activation", "a", "b", "c", "a", "d", None])
    ws.append([2, "stop", None, None, "c", "c", "d", "c", "d", None])

    def context_sheet(title: str, rows: list[tuple], garbage: bool = False) -> None:
        ws = wb.create_sheet(title)
        ws.append(ASSESSMENT_HEADER)
        last = None
        for n, (event, function, use_case, mode, mplr, cumulative, sg) in enumerate(rows, start=1):
            hazard, scenario, _ = split_event(event)
            first = hazard != last
            last = hazard
            ws.append(
                [
                    n if first else "#VALUE!",
                    hazard if first else None,
                    f"{hazard} text" if first else None,
                    scenario if first else None,
                    f"{scenario}: text" if first else None,
                    event,
                    f"{event}: text",
                    function,
                    use_case,
                    mode,
                    use_case if garbage else mplr,
                    cumulative,
                    "why" if cumulative else None,
                    "GID-1" if cumulative else None,
                    (GOAL_1 if hazard == "H-1" else GOAL_2) if cumulative else None,
                    sg,
                    None,
                ]
            )

    context_sheet("HARA(BF-BUC)", BF_BUC)
    context_sheet("HARA(WIF-BUC)", WIF_BUC)
    context_sheet("HARA-1 (BF-BUC)", BF_BUC, garbage=True)
    ws = wb.create_sheet("System-Level Constraint")
    ws.append([])
    ws.append([None, "1.3 System-Level Constraint"] * 1)
    ws.append([])
    ws.append([None, "ID", "System-Level Constraint", "Detail", "Relation"])
    ws.append([None, "SC-1", "Person separation", f"As per EARs Syntax:\n\nSafety Goal: {GOAL_1}", "H-1"])
    ws.append([None, "SC-2", "Gate protection", f"As per EARs Syntax:\n\nSafety Goal: {GOAL_2}", "H-2"])
    wb.save(path)
    return path


@pytest.fixture
def workbook(tmp_path):
    return synthetic_workbook(tmp_path / "synthetic_hara.xlsx")


@pytest.fixture
def loaded(brain, workbook):
    result = brain.execute("hara", action="new", target="from_workbook", **{"from": str(workbook)}, name="Gate HARA")
    assert result.ok, result.errors
    return brain, result.data, result.warnings


# -- the pieces ----------------------------------------------------------------------------------------
def test_the_nine_entities_are_registered_with_their_sheets_and_lookups(brain):
    assert len(brain.registry) == 28 and [e.name for e in HARA_ENTITIES] == ENTITIES
    sheets = [brain.registry.get(e).sheet for e in ENTITIES]
    assert sheets == [
        "HARA",
        "HARA vehicles",
        "HARA functions",
        "HARA use cases",
        "HARA failure modes",
        "HARA assumptions",
        "HARA ratings",
        "HARA assessments",
        "HARA approvals",
    ]
    assert all(brain.registry.get(e).field_names[-1] == "uid" for e in ENTITIES)
    assert all("hara" in brain.registry.get(e).field_names for e in PART_ENTITIES)
    assert brain.registry.get("hara_assessment").field("mplr").calculated
    assert brain.lookups["hara_context"] == HARA_CONTEXTS == ("BF-BUC", "BF-WIUC", "WIF-BUC", "WIF-WIUC")
    assert brain.lookups["stpa_source_kind"] == ("confluence", "hand", "workbook")
    assert SCHEMA_VERSION == 12 and brain.status()["schema_version"] == 12
    assert FORMATS == ("xlsx", "md") and DEFAULT_FORMATS == ("xlsx", "md")


def test_the_scale_the_keys_and_the_codes():
    assert [norm_mplr(v) for v in ("b", " C ", "QM", "qm", 0, "0", None, "", "-", "Travel/Roading *")] == [
        "b",
        "c",
        "QM",
        "QM",
        "",
        "",
        "",
        "",
        "-",
        "",
    ]
    assert highest(["b", "-", "", "d", "c"]) == "d" and highest(["QM", "a"]) == "a" and highest(["-", ""]) == ""
    assert equivalents("d") == {"mplr": "d", "plr": "d", "agplr": "d", "sil": "2"}
    assert equivalents("QM") == {"mplr": "QM", "plr": "QM", "agplr": "-", "sil": "QM"}
    assert equivalents("a")["sil"] == "-" and equivalents("e")["sil"] == "3"
    assert norm_key("Travel/Roading *", " lift") == norm_key("travel/roading", "LIFT")
    assert split_event("H-1.1.10") == ("H-1", "H-1.1", "H-1.1.10")
    assert split_event("H-4") == ("H-4", "", "H-4") and split_event("H-4.2") == ("H-4", "H-4.2", "H-4.2")
    assert split_event("nonsense") == ("", "", "nonsense")
    assert ratings_of({"g": "b", "B": 0}) == {"G": "b", "B": ""}
    assert ratings_of("G: b\nB = c; i: QM") == {"G": "b", "B": "c", "I": "QM"}
    assert ratings_of('{"G": "d"}') == {"G": "d"} and ratings_of("") == {} and ratings_of(None) == {}
    with pytest.raises(ValueError):
        ratings_of("not a rating line")
    assert guess_vehicle_code({"vehicle_name": "Gate Carrier"}, [{"letter": "G", "iso_vehicle": VEHICLES["G"]}]) == "G"
    assert guess_vehicle_code({"vehicle_name": "Boat"}, [{"letter": "G", "iso_vehicle": VEHICLES["G"]}]) is None


# -- the reader ------------------------------------------------------------------------------------------
def test_reader_reads_every_sheet_of_the_synthetic_workbook(workbook):
    data = read_hara_workbook(workbook)
    head = data["hara"]
    assert head["name"] == TITLE and head["system"] == "Gate Carrier ASAM" and head["vehicle_code"] == "G"
    assert (
        head["vehicle_name"] == "Gate Carrier" and head["analysis_date"] == "2026-01-15" and head["revision"] == "RA1"
    )
    assert head["standard"] == "Risk Assessment- ISO 19014" and head["project"] == "Carrier Platform"
    assert head["description"].startswith("System Description") and head["scope"].startswith("This risk assessment")
    assert "Pat Example" not in json.dumps(data)  # the responsible and the approvers are roles, never names
    assert [v["letter"] for v in data["vehicles"]] == ["B", "G", "I"]
    functions = {f["name"]: f for f in data["functions"]}
    assert functions["lift"]["kind"] == "base" and functions["lift"]["vehicles"] == ["B", "G"]
    assert functions["lift"]["cumulative_mplr"] == "c" and functions["lift"]["specific_mplr"] == "b"
    assert functions["lift"]["definition"] == "Raising the carrier's load."
    assert functions["tilt"]["kind"] == "work_implement" and functions["tilt"]["vehicles"] == ["G"]
    use_cases = {u["name"]: u for u in data["use_cases"]}
    assert use_cases["Travel/Roading *"]["base"] and use_cases["Travel/Roading *"]["vehicles"] == ["B", "G", "I"]
    assert use_cases["Gating"]["vehicles"] == ["G"] and use_cases["Gating"]["definition"] == "Passing through a gate."
    assert [m["name"] for m in data["failure_modes"]] == ["Failure to apply on Demand", "Uncommanded Activation"]
    assert [(a["source"], a["vehicle"]) for a in data["assumptions"]] == [("ISO 19014", VEHICLES["B"]), ("ASI", "All")]
    ratings = {(r["function"], r["use_case"], r["failure_mode"]): r for r in data["ratings"]}
    lift = ratings[("lift", "Travel/Roading *", "Uncommanded Activation")]
    assert lift["ratings"] == {"B": "b", "G": "c", "I": "a", "A": "a"}  # A from RA(for all vehicles), the rest primary
    assert lift["cumulative"] == "d" and lift["comment"] == "a note"
    assert ratings[("lift", "Travel/Roading *", "Failure to apply on Demand")]["ratings"] == {
        "B": "",
        "G": "b",
        "I": "",
    }
    assert len(data["ratings"]) == 4
    assert [a["context"] for a in data["assessments"]] == ["BF-BUC"] * 4 + ["WIF-BUC"] * 2
    first = data["assessments"][0]
    assert first["hazard"] == "H-1" and first["scenario"] == "H-1.1" and first["event"] == "H-1.1.1"
    assert first["stated_mplr"] == "c" and first["cumulative_mplr"] == "d" and first["sg_mplr"] == "d"
    assert first["safety_goal"] == "SC-1" and first["comment"] == "why" and first["sg_global_id"] == "GID-1"
    assert data["assessments"][2]["stated_mplr"] == NOT_FOUND and data["assessments"][3]["safety_goal"] == "SC-2"
    assert data["approvals"] == [
        {"event": "H-1.1.1", "role": "robot_chief", "decision": "Y"},
        {"event": "H-1.1.1", "role": "software_lead", "decision": "N"},
    ]
    stpa = data["stpa"]
    assert stpa["analysis"]["source_kind"] == "workbook" and stpa["analysis"]["source_bundle"] == str(workbook)
    assert [x["code"] for x in stpa["losses"]] == ["L-1", "L-2"]
    hazards = {h["code"]: h for h in stpa["hazards"]}
    assert list(hazards) == ["H-1", "H-2", "H-1.1", "H-2.1", "H-1.1.1", "H-1.1.2", "H-1.1.10", "H-2.1.1"]
    assert hazards["H-1"]["losses"] == ["L-1"] and hazards["H-1"]["note"] == "Physical State: too close\n[OK]"
    assert (
        hazards["H-2.1"]["retired"]
        and hazards["H-2.1"]["parent"] == "H-2"
        and hazards["H-1.1"]["text"] == "Direct interactions"
    )
    assert hazards["H-1.1.10"]["text"] == "Carrier reverses into a person" and hazards["H-1.1.10"]["parent"] == "H-1.1"
    assert hazards["H-2.1.1"]["parent"] == "H-2.1" and hazards["H-1.1.1"]["note"] == "approval comment: PE: check"
    assert hazards["H-1"]["cite"].endswith("Hazard List!B5") and hazards["H-1.1.1"]["cite"].endswith(
        "Hazardous Events!A2"
    )
    constraints = {c["code"]: c for c in stpa["constraints"]}
    assert constraints["SC-1"]["text"] == GOAL_1 and constraints["SC-1"]["note"] == "Person separation\nhazard: H-1"
    joined = "\n".join(data["warnings"])
    assert "1 id(s) were cut short" in joined and "H-2.1.1 relates to 'H-2.1.'" in joined
    assert "RA(ISO 19014): 1 rated row(s) carry no use case" in joined and "G taken from the vehicle name" in joined


def test_reader_second_generation_keeps_no_stated_rating(workbook):
    data = read_hara_workbook(workbook, generation=2, vehicle_code="b")
    assert data["hara"]["vehicle_code"] == "B" and data["hara"]["source_generation"] == 2
    assert [a["context"] for a in data["assessments"]] == ["BF-BUC"] * 4
    assert all(a["stated_mplr"] == "" for a in data["assessments"])
    joined = "\n".join(data["warnings"])
    assert "HARA-1 (BF-BUC): 4 MPLr cell(s) hold text that is no rating" in joined
    assert "no sheet for BF-WIUC, WIF-BUC, WIF-WIUC" in joined
    with pytest.raises(Exception, match="generation must be 1 or 2"):
        read_hara_workbook(workbook, generation=3)


def test_reader_refuses_what_is_not_a_hara_workbook(tmp_path):
    wb = Workbook()
    wb.active.title = "Nothing"
    wb.save(tmp_path / "other.xlsx")
    with pytest.raises(Exception, match="no sheet named like"):
        read_hara_workbook(tmp_path / "other.xlsx")
    with pytest.raises(Exception, match="workbook not found"):
        read_hara_workbook(tmp_path / "absent.xlsx")


# -- the load and the one computation ------------------------------------------------------------------
def test_from_workbook_loads_every_part_and_computes_the_mplr(loaded):
    brain, row, warnings = loaded
    assert row["name"] == "Gate HARA" and row["slug"] == "gate_hara" and row["vehicle_code"] == "G"
    assert row["analysis"] == "gate_hara" and row["counts"] == COUNTS and row["step1"] == STEP1
    assert row["mplr"] == {"rows": 6, "changed": 6, "not_found": 1, "empty": 0, "vehicle_code": "G", "ratings": 4}
    assert row["source_generation"] == 1 and len(row["source_sha256"]) == 64
    assessments = brain.execute("list", entity="hara_assessment").data
    computed = [(a["event"], a["function"], a["use_case"], a["mplr"], a["stated_mplr"]) for a in assessments]
    assert computed == [
        ("H-1.1.1", "lift", "Travel/Roading *", "c", "c"),
        ("H-1.1.1", "stop", "Travel/Roading *", "d", "d"),
        ("H-1.1.2", "lift", "Travel/Roading", "b", "-"),  # the asterisk forgiven: the workbook's formula found nothing
        ("H-2.1.1", "lift", "Travel/Roading *", "c", "d"),  # the workbook read the cumulative column
        ("H-1.1.1", "tilt", "Gating", "b", "b"),
        ("H-2.1.1", "tilt", "Travel/Roading *", "-", "-"),
    ]
    assert [a["order"] for a in assessments] == [1, 2, 3, 4, 5, 6]
    analysis = brain.execute("stpa", action="get", target="gate_hara").data
    assert analysis["source_kind"] == "workbook" and analysis["system"] == "Gate Carrier ASAM"
    assert {k: len(v) for k, v in analysis["parts"].items()} == {**STEP1, "ucas": 0, "scenarios": 0}
    hazards = {h["code"]: h for h in analysis["parts"]["hazards"]}
    assert hazards["H-1.1.10"]["parent"] == "H-1.1" and hazards["H-2.1"]["retired"] is True
    listed = brain.execute("hara", action="list").data
    assert [h["slug"] for h in listed] == ["gate_hara"] and listed[0]["counts"] == COUNTS
    assert brain.status()["entities"] == {
        **EMPTY_COUNTS,
        "stpa_analysis": 1,
        "stpa_loss": 2,
        "stpa_hazard": 8,
        "stpa_constraint": 2,
        "hara": 1,
        **{f"hara_{k[:-1] if k != 'use_cases' else 'use_case'}": v for k, v in COUNTS.items() if k != "failure_modes"},
        "hara_failure_mode": 2,
    }
    audit = [a["action"] for a in brain.storage.audit_tail(400)]
    assert "hara.new" in audit and "stpa.new" in audit
    assert any("cut short" in w for w in warnings)


def test_summary_derives_the_sheet_and_check_says_where_the_workbook_differs(loaded):
    brain, row, _ = loaded
    summary = brain.execute("hara", action="summary", target="gate_hara").data
    by = {e["hazard"]: e for e in summary["hazards"]}
    assert list(by) == ["H-1", "H-2"]
    h1 = by["H-1"]
    assert h1["contexts"] == {"BF-BUC": "d", "BF-WIUC": "", "WIF-BUC": "b", "WIF-WIUC": ""}
    assert (h1["mplr"], h1["plr"], h1["agplr"], h1["sil"]) == ("d", "d", "d", "2")
    assert h1["typed_cumulative"]["BF-BUC"] == "d" and h1["typed_mplr"] == "d" and h1["agrees"] is True
    assert (
        h1["safety_goal"] == "SC-1"
        and h1["safety_goal_text"] == GOAL_1
        and h1["text"] == "Carrier too close to persons"
    )
    h2 = by["H-2"]
    assert h2["contexts"]["BF-BUC"] == "c" and h2["contexts"]["WIF-BUC"] == "" and h2["mplr"] == "c"
    assert h2["typed_mplr"] == "d" and h2["agrees"] is False and h2["safety_goal"] == "SC-2"
    assert (summary["agree"], summary["differ"], summary["untyped"]) == (1, 1, 0)
    check = brain.execute("hara", action="check", target="gate_hara")
    assert check.ok and check.warnings == ["BF-BUC: 2 of 4 rows differ from the workbook"]
    bf = check.data["contexts"]["BF-BUC"]
    assert {k: bf[k] for k in ("rows", "agree", "differ", "not_found", "unstated")} == {
        "rows": 4,
        "agree": 2,
        "differ": 2,
        "not_found": 0,
        "unstated": 0,
    }
    assert [(x["event"], x["computed"], x["stated"]) for x in bf["examples"]] == [
        ("H-1.1.2", "b", "-"),
        ("H-2.1.1", "c", "d"),
    ]
    wif = check.data["contexts"]["WIF-BUC"]
    assert (wif["rows"], wif["agree"], wif["differ"], wif["not_found"]) == (2, 2, 0, 1)
    assert check.data["summary"] == {"agree": 1, "differ": 1, "untyped": 0}
    assert check.data["hazards_differing"] == [{"hazard": "H-2", "computed": "c", "typed": "d"}]
    # the asterisk is forgiven on the lists too: "Travel/Roading" is the listed "Travel/Roading *"
    assert check.data["unknown"] == {"functions": [], "use_cases": [], "failure_modes": [], "events": []}
    validate = brain.execute("hara", action="validate", target="gate_hara")
    assert validate.ok and validate.warnings == []
    assert brain.execute(
        "create",
        entity="hara_assessment",
        data={
            "hara": "gate_hara",
            "context": "BF-WIUC",
            "event": "H-9.9.9",
            "function": "fly",
            "use_case": "Gating",
            "failure_mode": "Uncommanded Activation",
        },
    ).ok
    unknown = brain.execute("hara", action="check", target="gate_hara").data["unknown"]
    assert unknown == {"functions": ["fly"], "use_cases": [], "failure_modes": [], "events": ["H-9.9.9"]}
    assert brain.execute("hara", action="validate", target="gate_hara").warnings == [
        "functions: 'fly' is used but not listed",
        "events: 'H-9.9.9' is used but not listed",
    ]


def test_recompute_follows_the_vehicle_code_and_a_rating_edit(loaded):
    brain, row, _ = loaded
    assert brain.execute("update", entity="hara", id=row["id"], data={"vehicle_code": CUMULATIVE_CODE}).ok
    done = brain.execute("hara", action="recompute", target="gate_hara")
    assert done.ok and done.data["vehicle_code"] == CUMULATIVE_CODE and done.data["changed"] == 0  # the update did it
    mplr = {(a["event"], a["function"]): a["mplr"] for a in brain.execute("list", entity="hara_assessment").data}
    assert mplr[("H-1.1.1", "lift")] == "d" and mplr[("H-1.1.1", "stop")] == "d" and mplr[("H-1.1.1", "tilt")] == "c"
    assert brain.execute("update", entity="hara", id=row["id"], data={"vehicle_code": "i"}).ok  # upper-cased
    assert brain.execute("get", entity="hara", id=row["id"]).data["vehicle_code"] == "I"
    brain.execute("hara", action="recompute", target="gate_hara")
    mplr = {(a["event"], a["function"]): a["mplr"] for a in brain.execute("list", entity="hara_assessment").data}
    assert mplr[("H-1.1.1", "lift")] == "a" and mplr[("H-1.1.1", "tilt")] == ""  # an empty rating cell
    rating = next(r for r in brain.execute("list", entity="hara_rating").data if r["function"] == "tilt")
    assert rating["ratings"] == {"B": "", "G": "b", "I": ""}
    edited = brain.execute("update", entity="hara_rating", id=rating["id"], data={"ratings": "G: b\nI: e"})
    assert edited.ok and edited.data["ratings"] == {"G": "b", "I": "e"}
    brain.execute("hara", action="recompute", target="gate_hara")
    mplr = {(a["event"], a["function"]): a["mplr"] for a in brain.execute("list", entity="hara_assessment").data}
    assert mplr[("H-1.1.1", "tilt")] == "e"
    assert brain.execute("update", entity="hara", id=row["id"], data={"vehicle_code": ""}).ok
    cleared = brain.execute("hara", action="recompute", target="gate_hara").data
    assert cleared["empty"] == 5 and cleared["not_found"] == 1
    assert (
        brain.execute("hara", action="validate", target="gate_hara").warnings[0]
        == "no vehicle code: no MPLr is computed"
    )
    refused = brain.execute("update", entity="hara", id=row["id"], data={"vehicle_code": "GX"})
    assert not refused.ok and "neither an annex letter" in refused.errors[0]


# -- the rules -----------------------------------------------------------------------------------------
def test_the_rules_hold_on_every_verb(loaded):
    brain, row, _ = loaded
    refused = brain.execute("create", entity="hara_function", data={"hara": "gate_hara", "name": " LIFT "})
    assert not refused.ok and "already a HARA Function" in refused.errors[0]
    refused = brain.execute("create", entity="hara_vehicle", data={"hara": "nowhere", "letter": "Z"})
    assert not refused.ok and "no HARA has the slug" in refused.errors[0]
    made = brain.execute("create", entity="hara_vehicle", data={"hara": "gate_hara", "letter": "z"})
    assert made.ok and made.data["letter"] == "Z" and made.data["order"] == 4
    refused = brain.execute(
        "create",
        entity="hara_assessment",
        data={"hara": "gate_hara", "context": "BF-BUC", "event": "H-1.1.2", "mplr": "d"},
    )
    assert not refused.ok and "mplr: calculated" in refused.errors[0]
    made = brain.execute(
        "create",
        entity="hara_assessment",
        data={
            "hara": "gate_hara",
            "context": "BF-BUC",
            "event": "H-1.1.2",
            "function": "stop",
            "use_case": "Travel/Roading *",
            "failure_mode": "Failure to apply on Demand",
            "stated_mplr": 0,
        },
    )
    assert (
        made.ok and made.data["hazard"] == "H-1" and made.data["scenario"] == "H-1.1" and made.data["stated_mplr"] == ""
    )
    assert made.data["mplr"] is None  # the Result is the row as written; the after hook computed it since
    assert brain.execute("get", entity="hara_assessment", id=made.data["id"]).data["mplr"] == "d"
    assert brain.execute("hara", action="recompute", target="gate_hara").data["changed"] == 0
    refused = brain.execute("update", entity="hara_assessment", id=made.data["id"], data={"mplr": "a"})
    assert not refused.ok and "mplr: calculated" in refused.errors[0]
    refused = brain.execute(
        "create", entity="hara_assessment", data={"hara": "gate_hara", "context": "SIDEWAYS", "event": "H-1.1.1"}
    )
    assert not refused.ok and "context" in refused.errors[0]
    refused = brain.execute(
        "create",
        entity="hara_rating",
        data={
            "hara": "gate_hara",
            "function": "Lift",
            "use_case": "travel/roading",
            "failure_mode": "Uncommanded Activation",
        },
    )
    assert not refused.ok and "already exists" in refused.errors[0]
    refused = brain.execute(
        "create",
        entity="hara_rating",
        data={
            "hara": "gate_hara",
            "function": "lift",
            "use_case": "Gating",
            "failure_mode": "Uncommanded Activation",
            "ratings": "what",
        },
    )
    assert not refused.ok and "ratings:" in refused.errors[0]
    approval = brain.execute(
        "create",
        entity="hara_approval",
        data={"hara": "gate_hara", "event": "H-1.1.2", "role": "Systems Architect", "decision": "yes"},
    )
    assert approval.ok and approval.data["role"] == "systems_architect" and approval.data["decision"] == "Y"
    refused = brain.execute(
        "create",
        entity="hara_approval",
        data={"hara": "gate_hara", "event": "H-1.1.2", "role": "x", "decision": "maybe"},
    )
    assert not refused.ok and "decision" in refused.errors[0]
    refused = brain.execute("update", entity="hara", id=row["id"], data={"slug": "other"})
    assert not refused.ok and "the parts point at it" in refused.errors[0]
    twin = brain.execute("hara", action="new", target="blank", name="Gate HARA")
    assert not twin.ok and "already the slug" in twin.errors[0]
    blank = brain.execute("hara", action="new", target="blank", name="Empty", vehicle="g")
    assert (
        blank.ok
        and blank.data["slug"] == "empty"
        and blank.data["vehicle_code"] == "G"
        and blank.data["analysis"] is None
    )
    refused = brain.execute("update", entity="hara", id=blank.data["id"], data={"analysis": "nowhere"})
    assert not refused.ok and "no STPA analysis has the slug" in refused.errors[0]
    with pytest.raises(ValidationError):
        brain.service("hara_function").create({"hara": "gate_hara", "name": ""})
    events = [e for e in brain.events.last(50) if e["topic"] == "hara.changed"]
    assert events and all(e.get("hara") in ("gate_hara", "empty") for e in events)


def test_delete_takes_the_parts_and_a_fresh_step_1_along(loaded):
    brain, row, _ = loaded
    refused = brain.execute("hara", action="delete", target="gate_hara")
    assert not refused.ok and "delete with --cascade" in refused.errors[0]
    assert brain.execute(
        "create",
        entity="stpa_uca",
        data={
            "analysis": "gate_hara",
            "code": "UCA1",
            "controller": "Carrier",
            "control_action": "lift",
            "category": "providing",
            "hazards": ["H-1"],
        },
    ).ok
    kept = brain.execute("hara", action="delete", target="gate_hara", cascade=True)
    assert kept.ok and kept.data["analysis_kept"] == "gate_hara" and kept.data["deleted"]["assessments"] == 6
    assert brain.execute("stpa", action="get", target="gate_hara").ok  # step 3 had begun: the analysis stays
    assert brain.execute("hara", action="list").data == []
    assert brain.execute("count", entity="hara_assessment").data["count"] == 0
    assert brain.execute("stpa", action="delete", target="gate_hara", cascade=True).ok
    again = brain.execute("hara", action="new", target="from_workbook", **{"from": row["source"]}, name="Gate HARA")
    assert again.ok
    gone = brain.execute("hara", action="delete", target="gate_hara", cascade=True)
    assert gone.ok and gone.data["analysis_kept"] is None and gone.data["deleted"]["stpa"]["hazards"] == 8
    assert brain.status()["entities"] == EMPTY_COUNTS


def test_from_workbook_can_point_at_an_analysis_already_in_kletz_and_clone(brain, workbook):
    assert brain.execute("stpa", action="new", target="blank", name="Step one", slug="step_one").ok
    made = brain.execute(
        "hara",
        action="new",
        target="from_workbook",
        **{"from": str(workbook)},
        name="Gate HARA",
        analysis="step_one",
        vehicle="B",
        generation=2,
    )
    assert made.ok and made.data["analysis"] == "step_one" and made.data["vehicle_code"] == "B"
    assert made.data["step1"] == {"losses": 0, "hazards": 0, "constraints": 0}
    assert made.data["source_generation"] == 2 and made.data["counts"]["assessments"] == 4
    assert brain.execute("count", entity="stpa_analysis").data["count"] == 1  # the workbook's step 1 was not loaded
    summary = brain.execute("hara", action="summary", target="gate_hara").data
    assert (
        summary["hazards"][0]["text"] == "" and summary["hazards"][0]["contexts"]["BF-BUC"] == "c"
    )  # B: lift b, stop c
    copy = brain.execute("hara", action="new", target="clone", **{"from": "gate_hara"}, name="Gate HARA copy")
    assert copy.ok and copy.data["slug"] == "gate_hara_copy" and copy.data["counts"] == made.data["counts"]
    assert copy.data["analysis"] == "step_one" and copy.data["mplr"]["rows"] == 4
    refused = brain.execute("hara", action="new", target="from_workbook", name="x")
    assert not refused.ok and "needs --from" in refused.errors[0]
    refused = brain.execute("hara", action="new", target="unknown_builder", name="x")
    assert not refused.ok and "unknown builder" in refused.errors[0]
    builders = brain.execute("hara", action="builders").data["builders"]
    assert [b["name"] for b in builders] == ["blank", "from_workbook", "clone"] and all(
        b["available"] for b in builders
    )


# -- the export -----------------------------------------------------------------------------------------
def test_export_writes_the_workbook_shape_and_the_report(loaded, tmp_path):
    brain, row, _ = loaded
    out = tmp_path / "exported"
    done = brain.execute("hara", action="export", target="gate_hara", out=str(out))
    assert done.ok and sorted(Path(p).name for p in done.data["files"]) == ["gate_hara.md", "gate_hara.xlsx"]
    wb = load_workbook(out / "gate_hara.xlsx")
    assert wb.sheetnames == [
        "Document Information",
        "Scope",
        "Function List",
        "Use Case List",
        "Failure modes",
        "Assumption List",
        "Loss",
        "Hazard List",
        "Scenarios List",
        "Hazardous Events",
        "RA(ISO 19014)",
        "HARA(BF-BUC)",
        "HARA(BF-WIUC)",
        "HARA(WIF-BUC)",
        "HARA(WIF-WIUC)",
        "System-Level Constraint",
        "Summary",
    ]
    info = {r[0]: r[1] for r in wb["Document Information"].iter_rows(min_row=2, values_only=True)}
    assert info["Vehicle code"] == "G" and info["STPA analysis"] == "gate_hara" and "Pat Example" not in str(info)
    events = list(wb["Hazardous Events"].iter_rows(values_only=True))
    assert events[0][3:5] == ("robot_chief", "software_lead") and events[1][3:5] == ("Y", "N")
    bf = list(wb["HARA(BF-BUC)"].iter_rows(values_only=True))
    assert bf[0][9] == "MPLr (G)" and bf[0][10] == "Stated MPLr"
    assert [(r[4], r[6], r[9], r[10]) for r in bf[1:]] == [
        ("H-1.1.1", "lift", "c", "c"),
        ("H-1.1.1", "stop", "d", "d"),
        ("H-1.1.2", "lift", "b", "-"),
        ("H-2.1.1", "lift", "c", "d"),
    ]
    assert bf[1][14] == f"SC-1: {GOAL_1}"
    summary = list(wb["Summary"].iter_rows(values_only=True))
    assert summary[0][:3] == ("Hazard", "Safety Goal", "BF-BUC") and summary[1][0].startswith("H-1: ")
    assert summary[1][6:10] == ("d", "d", "d", "2") and summary[2][6] == "c" and summary[2][-1] == "no"
    ratings = list(wb["RA(ISO 19014)"].iter_rows(values_only=True))
    assert ratings[0][4:8] == ("A", "B", "G", "I") and ratings[1][1:4] == (
        "lift",
        "Travel/Roading *",
        "Uncommanded Activation",
    )
    report = (out / "gate_hara.md").read_text(encoding="utf-8")
    assert report.startswith("# Gate HARA") and "## Summary (computed)" in report and "| H-2 | SC-2 |" in report
    assert "| BF-BUC | 4 | 2 | 2 | 0 | 0 |" in report and "computed c, stated d" in report
    audit = brain.storage.audit_tail(3)
    assert audit[0]["action"] == "hara.export" and audit[0]["payload"]["slug"] == "gate_hara"
    refused = brain.execute("hara", action="export", target="gate_hara", out=str(out), format="pdf")
    assert not refused.ok and "must be one of" in refused.errors[0]
    default = brain.execute("hara", action="export", target="gate_hara")
    assert default.ok and Path(default.data["out"]) == Path(brain.out_dir) / "hara"


# -- the CLI and the dump -----------------------------------------------------------------------------
def test_cli_and_dump_round_trip(loaded, tmp_path, capsys):
    brain, row, _ = loaded
    code = main(["--db", brain.db_path, "--no-forward", "--compact", "hara", "check", "gate_hara"])
    assert code == 0
    printed = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert printed["ok"] and printed["data"]["contexts"]["BF-BUC"]["differ"] == 2
    code = main(
        [
            "--db",
            brain.db_path,
            "--no-forward",
            "hara",
            "export",
            "gate_hara",
            "--out",
            str(tmp_path / "cli"),
            "--format",
            "md",
        ]
    )
    assert code == 0 and (tmp_path / "cli" / "gate_hara.md").is_file()
    dump = tmp_path / "dump.json"
    assert brain.execute("dump", out=str(dump)).ok
    payload = json.loads(dump.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 12 and len(payload["entities"]["hara_assessment"]) == 6
    assert brain.execute("hara", action="delete", target="gate_hara", cascade=True).ok
    back = brain.execute("load", path=str(dump), mode="merge")
    assert back.ok, back.errors
    assert brain.execute("hara", action="list").data[0]["counts"] == COUNTS
    mplr = [a["mplr"] for a in brain.execute("list", entity="hara_assessment").data]
    assert mplr == ["c", "d", "b", "c", "b", "-"]  # the calculated field came back through the load's recompute


def test_a_0_23_0_database_gains_the_hara_tables(tmp_path):
    import shutil

    db = tmp_path / "old.db"
    shutil.copy(DB_0_23_0, db)
    with Brain(db) as b:
        assert b.migration["created"][-9:] == HARA_TABLES and b.status()["schema_version"] == 12
        assert b.execute("hara", action="list").data == []


# -- the screen ----------------------------------------------------------------------------------------
@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_the_hara_screen_lists_parts_summary_and_runs_the_verbs(loaded, tmp_path):
    from asi_kletz.gui.hara_screen import SUMMARY_COLUMNS, HaraScreen
    from asi_kletz.gui.main_window import MainWindow
    from asi_kletz.gui.screens import HARA_SCREEN

    brain, row, _ = loaded
    window = MainWindow(brain, interactive=False)
    window.withdraw()
    try:
        screen = window.show(HARA_SCREEN)
        assert isinstance(screen, HaraScreen) and window.factory.title(HARA_SCREEN) == "HARA"
        assert list(screen.tree.get_children()) == ["gate_hara"]
        screen.select("gate_hara")
        assert len(screen.summary_tree.get_children()) == 2 and screen.summary_tree["columns"] == SUMMARY_COLUMNS
        assert screen.summary_tree.set("H-1", "mplr") == "d" and screen.summary_tree.set("H-2", "agrees") == "✗"
        assert len(screen.part_trees["hara_assessment"].get_children()) == 6
        assert len(screen.part_trees["hara_rating"].get_children()) == 4
        differing = [
            i
            for i in screen.part_trees["hara_assessment"].get_children()
            if "differ" in screen.part_trees["hara_assessment"].item(i, "tags")
        ]
        assert len(differing) == 2
        assert screen.check().ok and "BF-BUC: 4 rows, 2 agree, 2 differ" in screen.log.get("1.0", "end")
        assert screen.recompute().ok and screen.validate().ok
        exported = screen.export(str(tmp_path / "gui"))
        assert exported.ok and (tmp_path / "gui" / "gate_hara.xlsx").is_file()
        dash = window.show("dashboard")
        assert dash._stpa_vars["hara"].get() == "1"
        window.show(HARA_SCREEN)
        assert screen.delete().ok and list(screen.tree.get_children()) == []
        assert screen.summary_tree.get_children() == ()
    finally:
        window.close()
        del window
        gc.collect()
