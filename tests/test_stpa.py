"""STPA analyses read from a Confluence page (asi_hopper/docs/confluence-integration-plan.md, Phase D; 0.24.0,
moved here from the sister tool with the control structures).

Everything here runs on a **synthetic** STPA page with neutral invented names, written in the exact
Markdown shape the Atlassian MCP server renders (the four sections, the glued bold runs, the escaped
hazard lists, a struck duplicate, CA counterparts, CSTR references), and on a synthetic Hopper bundle
built around it in a temporary folder the way ``snapshot_pages`` writes one. No real page text is
here. The replay test reads asi_hopper's own fixtures, which are synthetic too.
"""

from __future__ import annotations

import hashlib
import json
import shutil
import sqlite3
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from asi_kletz import Brain
from asi_kletz.automation.cli import main
from asi_kletz.errors import NotFound, ValidationError
from asi_kletz.services.stpa_service import ANALYSIS_ENTITY, PART_ENTITIES, code_key, match_controller
from asi_kletz.stpa import StpaFactory, read_bundle
from asi_kletz.stpa.confluence_reader import (
    cite,
    codes_in,
    constraint_codes,
    embolden_loose_codes,
    read_constraint_cell,
    read_constraints,
    read_stpa_page,
    read_uca_cell,
    row_shift,
)
from asi_kletz.stpa.stpa_factory import DEFAULT_FORMATS, FORMATS
from asi_kletz.version import SCHEMA_VERSION, __version__
from tests.conftest import DB_0_23_0, EMPTY_COUNTS, REPO_ROOT, STPA_TABLES

HOPPER_FIXTURES = REPO_ROOT.parent / "asi_hopper" / "tests" / "fixtures"
STPA_ENTITIES = (ANALYSIS_ENTITY, *PART_ENTITIES)
PARTS = ("losses", "hazards", "ucas", "constraints", "scenarios")
PAGE = "900000001"
URL = f"https://example.invalid/wiki/pages/viewpage.action?pageId={PAGE}"
TITLE = "Conveyor Gate Safety Function STPA"

# the synthetic page (tests/fixtures/conveyor-gate-stpa.md): four `##` sections, a context table whose cells
# glue bold runs to their text, escaped hazard lists, a struck UCA that duplicates a later one, a link cell,
# a loss-scenario table with CA counterparts, beliefs, CSTR references and an unbolded row
PAGE_FILE = REPO_ROOT / "tests" / "fixtures" / "conveyor-gate-stpa.md"
PAGE_BODY = PAGE_FILE.read_text(encoding="utf-8")

COUNTS = {"losses": 2, "hazards": 5, "ucas": 8, "scenarios": 10, "constraints": 6}
COMMENTS = [
    {
        "id": "c1",
        "kind": "inline",
        "body": "Should the hold expire at all?",
        "author_id": None,
        "created": "2026-10-20T10:00:00Z",
        "status": "current",
        "title": None,
        "resolution": "open",
        "selection": "When the supervisor's hold has expired.",
        "replies": [
            {
                "id": "c1r",
                "kind": "inline",
                "body": "Yes, after the shift.",
                "author_id": None,
                "created": "2026-10-21T10:00:00Z",
                "status": "current",
                "title": None,
                "resolution": None,
                "selection": None,
                "replies": [],
            },
        ],
    },
    {
        "id": "c2",
        "kind": "footer",
        "body": "Reviewed for the gate release.",
        "author_id": None,
        "created": "2026-10-22T10:00:00Z",
        "status": "current",
        "title": None,
        "resolution": None,
        "selection": None,
        "replies": [],
    },
]


def write_bundle(
    folder: Path,
    body: str = PAGE_BODY,
    *,
    version: int = 7,
    page_id: str = PAGE,
    comments: list | None = None,
    taken_at: str = "2026-10-29T12:00:00-06:00",
) -> Path:
    """A Confluence bundle in the shape ``asi_hopper``'s ``snapshot_pages`` writes: ``manifest.json``,
    ``pages.json``, ``pages/<id>.md`` byte for byte, the body hashed."""
    from asi_hopper.model.pagetext import sections

    (folder / "pages").mkdir(parents=True, exist_ok=True)
    data = body.encode("utf-8")
    sha = hashlib.sha256(data).hexdigest()
    (folder / "pages" / f"{page_id}.md").write_bytes(data)
    entry = {
        "id": page_id,
        "title": TITLE,
        "status": "current",
        "version": version,
        "updated": "2026-10-28T09:00:00Z",
        "url": f"https://example.invalid/wiki/pages/viewpage.action?pageId={page_id}",
        "body_file": f"pages/{page_id}.md",
        "body_sha256": sha,
        "body_chars": len(body),
        "sections": [
            {"heading": s.heading, "anchor": s.anchor, "sha256": s.sha256, "chars": s.chars} for s in sections(body)
        ],
        "comments": comments if comments is not None else COMMENTS,
    }
    entry["comment_count"] = len(entry["comments"])
    pages = json.dumps([entry], indent=2, ensure_ascii=False) + "\n"
    (folder / "pages.json").write_text(pages, encoding="utf-8", newline="\n")
    manifest = {
        "kind": "confluence",
        "tool": "asi_hopper 0.1.9",
        "record_version": 1,
        "site": "example.invalid",
        "account": {},
        "taken_at": taken_at,
        "source": {"kind": "pages", "ids": [page_id]},
        "comments": True,
        "count": 1,
        "pages": [{"id": page_id, "title": TITLE, "version": version, "updated": entry["updated"], "body_sha256": sha}],
        "files": {"pages": "pages.json", "bodies": "pages/", "csv": "pages.csv", "readable": "pages.md"},
        "sha256": {"pages.json": hashlib.sha256(pages.encode("utf-8")).hexdigest(), f"pages/{page_id}.md": sha},
        "folder": str(folder),
    }
    (folder / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8", newline="\n")
    return folder


def changed_body() -> str:
    """The page a version later: UCA4's text changed, a new hazard, UCA8 gone, CSTR4 no longer spelled out."""
    body = PAGE_BODY.replace("**UCA4 **While a person is in the opening.", "**UCA4 **While anyone is in the opening.")
    body = body.replace("~~H4 - An old hazard, retired~~", "~~H4 - An old hazard, retired~~\n\nH5 - A new hazard")
    body = body.replace("**UCA8 **Held after maintenance has left. *\\[H2\\]*", "")
    return body.replace("**CSTR4 **An expired hold must close the gate without a new request.", "**CSTR4**")


@pytest.fixture
def bundle(tmp_path):
    return write_bundle(tmp_path / "bundle" / "conveyor-gate-stpa-2026-10-29")


@pytest.fixture
def loaded(brain, bundle):
    row, warnings = brain.stpa_factory.create("from_bundle", "test", path=str(bundle), name="Conveyor Gate")
    return brain, row, warnings


def small_structure(brain) -> dict:
    """A three-box structure whose names match the page's controllers and control actions, nearly."""
    row, _ = brain.control_structures.create_from_dataset(
        {
            "structure": {"name": "Conveyor Gate Structure", "slug": "gate", "system": "Conveyor gate"},
            "controllers": [
                {"key": "supervisor", "order": 1, "name": "Supervisor", "kind": "human_controller"},
                {"key": "gate_controller", "order": 2, "name": "Gate Controller"},
                {"key": "gate", "order": 3, "name": "Gate", "kind": "controlled_process"},
            ],
            "edges": [
                {"key": "e1", "kind": "control_action", "source": "supervisor", "target": "gate_controller"},
                {"key": "e2", "kind": "control_action", "source": "gate_controller", "target": "gate"},
            ],
            "signals": [
                {"edge": "e1", "signal_id": "CA-01", "name": "Hold request override"},
                {"edge": "e2", "signal_id": "CA-10", "name": "Close gate"},
                {"edge": "e2", "signal_id": "CA-11", "name": "Open gate"},
            ],
        },
        "test",
    )
    return row


# -- entities --------------------------------------------------------------------------------------
def test_the_six_entities_are_registered_with_their_sheets_and_lookups(brain):
    assert len(brain.registry) == 19 and set(s.name for s in brain.registry) >= set(STPA_ENTITIES)
    sheets = [brain.registry.get(e).sheet for e in STPA_ENTITIES]
    assert sheets == ["STPA analyses", "STPA losses", "STPA hazards", "STPA UCAs", "STPA constraints", "STPA scenarios"]
    assert all(brain.registry.get(e).field_names[-1] == "uid" for e in STPA_ENTITIES)
    assert all("cite" in brain.registry.get(e).field_names for e in PART_ENTITIES)
    assert brain.lookups["stpa_uca_category"] == ("not_providing", "providing", "timing_order", "duration")
    assert brain.lookups["stpa_source_kind"] == ("confluence", "hand")
    assert brain.lookups["stpa_status"] == ("Draft", "Reviewed", "Baselined")
    assert SCHEMA_VERSION == 11 and brain.status()["schema_version"] == 11
    assert brain.status()["evidence_dir"].endswith("evidence")


# -- the reader ------------------------------------------------------------------------------------
def test_reader_counts_codes_categories_and_cites():
    data = read_stpa_page(PAGE_BODY, page_id=PAGE, url=URL, version=7, title=TITLE)
    assert {k: len(data[k]) for k in COUNTS} == COUNTS
    assert data["analysis"] == {
        "name": TITLE,
        "system": "",
        "source_kind": "confluence",
        "source_page_id": PAGE,
        "source_url": URL,
        "source_version": 7,
    }
    assert [h["code"] for h in data["losses"]] == ["L1", "L2"]
    assert data["losses"][1]["text"] == "The conveyor line stops for longer than a shift"
    assert [h["code"] for h in data["hazards"]] == ["H1", "H2", "H2.1", "H3", "H4"]
    by_code = {h["code"]: h for h in data["hazards"]}
    assert (
        by_code["H2.1"]["parent"] == "H2" and by_code["H2"]["parent"] == "" and by_code["H2"]["losses"] == ["L1", "L2"]
    )
    assert by_code["H4"]["retired"] is True and by_code["H4"]["text"] == "An old hazard, retired"
    assert by_code["H1"]["text"] == "The gate closes while a person is in the opening"
    ucas = {u["code"]: u for u in data["ucas"]}
    assert list(ucas) == [f"UCA{n}" for n in range(1, 9)]
    assert [u["category"] for u in data["ucas"]] == [
        "not_providing",
        "not_providing",
        "not_providing",
        "providing",
        "timing_order",
        "not_providing",
        "timing_order",
        "duration",
    ]
    assert [u["controller"] for u in data["ucas"]] == ["Gate Controller"] * 5 + ["Supervisor"] * 3
    assert ucas["UCA1"]["control_action"] == "Close Gate" and ucas["UCA8"]["control_action"] == "Hold Request"
    assert ucas["UCA1"]["text"] == "When the belt is about to run." and ucas["UCA1"]["hazards"] == ["H2", "H2.1"]
    assert ucas["UCA2"]["hazards"] == ["H2", "H9"] and ucas["UCA4"]["hazards"] == ["H1"]
    assert ucas["UCA1"]["cite"] == f"page {PAGE} v7 § Context Table ({URL}#Context-Table)"
    assert by_code["H1"]["cite"] == f"page {PAGE} v7 § Losses and Hazards ({URL}#Losses-and-Hazards)"
    assert data["scenarios"][0]["cite"] == f"page {PAGE} v7 § Loss Scenarios ({URL}#Loss-Scenarios)"
    assert data["sections"] == ["Losses and Hazards", "Control Structure", "Context Table", "Loss Scenarios"]
    assert cite("1", None, None, "") == "page 1" and cite("1", 2, None, "u") == "page 1 v2 (u)"
    assert codes_in(r"text *\[H1, H4.3\]* more [L1]") == ["H1", "H4.3", "L1"]


def test_reader_retired_duplicate_link_cell_and_unknown_hazard():
    data = read_stpa_page(PAGE_BODY, page_id=PAGE, url=URL, version=7)
    ucas = {u["code"]: u for u in data["ucas"]}
    assert ucas["UCA3"]["retired"] is True and ucas["UCA3"]["duplicate_of"] == "UCA6"
    assert ucas["UCA3"]["text"] == "When the line restarts after a fault." and ucas["UCA3"]["hazards"] == ["H2"]
    assert all(not u["retired"] for c, u in ucas.items() if c != "UCA3")
    assert all(not u["duplicate_of"] for c, u in ucas.items() if c != "UCA3")
    assert "UCA6" in ucas  # the twin named in "(duplicate of **UCA6**)" is not read as a UCA of the first cell
    warnings = data["warnings"]
    assert any("UCA2: cites hazard H9, which the page does not list" == w for w in warnings)
    link = [w for w in warnings if "Belt Drive / Run Belt" in w]
    assert len(link) == 1 and "page 800000002 (Belt Drive STPA)" in link[0] and "no UCA read" in link[0]
    assert any(w.startswith("Loss Scenarios row 10: empty UCA cell") for w in warnings)
    assert not any("no losses" in w for w in warnings)
    # a page without losses says so, and a page without the sections says so
    bare = read_stpa_page("## Losses and Hazards\n\nH1 - one\n", page_id="1")
    assert bare["hazards"][0]["code"] == "H1" and bare["losses"] == []
    assert any("no losses listed" in w for w in bare["warnings"])
    assert any("Context Table" in w and "no section" in w for w in bare["warnings"])
    assert bare["hazards"][0]["cite"] == "page 1 § Losses and Hazards (Losses-and-Hazards)"
    empty = read_stpa_page("", page_id="2")
    assert {k: len(empty[k]) for k in COUNTS} == dict.fromkeys(COUNTS, 0) and len(empty["warnings"]) == 3


def test_reader_scenarios_types_beliefs_counterparts_and_constraints():
    data = read_stpa_page(PAGE_BODY, page_id=PAGE, url=URL, version=7)
    rows = {s["row"]: s for s in data["scenarios"]}
    assert list(rows) == list(range(1, 11)) and [s["order"] for s in data["scenarios"]] == list(range(1, 11))
    assert rows[1]["uca"] == "UCA1" and rows[1]["counterpart"] == "" and rows[1]["types"] == [1, 2]
    assert rows[1]["beliefs"] == "The gate is already closed." and rows[1]["constraints"] == ["CSTR1", "CSTR2"]
    assert rows[1]["causal_factors"] == "The gate position sensor is stuck. The run request is not received."
    assert rows[2]["uca"] == "UCA1" and rows[2]["counterpart"] == "CA1" and rows[2]["types"] == [3, 4]
    assert rows[2]["constraints"] == ["CSTR1", "CSTR2", "CSTR3"] and rows[2]["beliefs"] == ""
    assert rows[4]["uca"] == "UCA2" and rows[4]["counterpart"] == "CA2" and rows[4]["constraints"] == ["CSTR1", "CSTR3"]
    assert rows[5]["uca"] == "UCA3" and rows[5]["retired"] is True and rows[5]["constraints"] == ["CSTR1"]
    assert rows[8]["beliefs"] == "Maintenance has left." and rows[8]["constraints"] == ["CSTR3", "CSTR6"]
    assert rows[9]["uca"] == "UCA7" and rows[9]["counterpart"] == "" and rows[9]["constraints"] == ["CSTR5", "CSTR6"]
    assert rows[10]["uca"] == "" and rows[10]["types"] == [3, 4] and rows[10]["constraints"] == []
    constraints = {c["code"]: c for c in data["constraints"]}
    assert list(constraints) == [f"CSTR{n}" for n in range(1, 7)]
    assert constraints["CSTR1"]["text"] == "The gate controller must confirm the gate position before every run."
    assert constraints["CSTR3"]["text"] == "The gate drive must report a failed close within one second."
    assert constraints["CSTR6"]["text"] == "Maintenance presence must be shown at the supervisor's desk."
    assert read_constraint_cell("**CSTR1 **and **CSTR2 ****CSTR3 **text") == [
        ("CSTR1", ""),
        ("CSTR2", ""),
        ("CSTR3", "text"),
    ]
    assert read_constraint_cell("~~**CSTR1**~~") == [("CSTR1", "")] and read_constraint_cell("") == []
    assert read_constraint_cell("CSTR9 spelled without bold and") == [("CSTR9", "spelled without bold")]


# -- the factory ------------------------------------------------------------------------------------
def test_from_bundle_loads_through_the_services(loaded, bundle):
    brain, row, warnings = loaded
    assert row["slug"] == "conveyor_gate" and row["counts"] == COUNTS and row["status"] == "Draft"
    assert row["source_kind"] == "confluence" and row["source_page_id"] == PAGE and row["source_version"] == 7
    assert row["source_url"] == URL and row["source_bundle"] == str(bundle)
    assert row["source_body_sha256"] == hashlib.sha256(PAGE_BODY.encode("utf-8")).hexdigest()
    assert row["source_taken_at"] == "2026-10-29T12:00:00-06:00"
    # the reader's warnings first, then the services': the dangling H9 is a warning, never a refusal
    assert warnings[0].startswith("UCA2: cites hazard H9") and any("Belt Drive" in w for w in warnings)
    assert any(w == "stpa_uca: UCA2: hazards names H9, which the analysis has no hazard for" for w in warnings)
    assert any(w == "stpa_scenario: row 10: no UCA" for w in warnings)
    assert not any("duplicate of UCA6" in w for w in warnings)  # the twin listed after it is known to the load
    ucas = brain.service("stpa_uca").list({"analysis": "conveyor_gate"})
    assert [u["code"] for u in ucas] == [f"UCA{n}" for n in range(1, 9)]
    assert ucas[2]["retired"] is True and ucas[2]["duplicate_of"] == "UCA6" and ucas[0]["hazards"] == ["H2", "H2.1"]
    assert all(u["uid"] and u["cite"].startswith(f"page {PAGE} v7 § Context Table") for u in ucas)
    hazards = brain.service("stpa_hazard").list({"analysis": "conveyor_gate"})
    assert [h["parent"] for h in hazards] == [None, None, "H2", None, None]
    tail = brain.storage.audit_tail(1)[0]
    assert tail["action"] == "stpa.new" and tail["payload"] == {"slug": "conveyor_gate", **COUNTS}
    event = brain.events.last(1)[0]
    assert event["topic"] == "stpa.changed" and event["action"] == "new" and event["analysis"] == "conveyor_gate"
    # the bundle is read again the same way, and the page can be picked
    again = read_bundle(bundle, PAGE)
    assert {k: len(again[k]) for k in COUNTS} == COUNTS and len(again["comments"]) == 2
    with pytest.raises(NotFound, match="no page 1 in the bundle"):
        read_bundle(bundle, "1")
    with pytest.raises(NotFound, match="not a Hopper bundle"):
        read_bundle(bundle.parent)
    # a second analysis of the same page and version is refused
    with pytest.raises(ValidationError, match="already holds page"):
        brain.stpa_factory.create("from_bundle", "test", path=str(bundle), name="Twice")


def test_blank_and_clone_builders_and_the_catalog(loaded):
    brain, row, _ = loaded
    catalog = brain.stpa_factory.catalog()
    assert [b["name"] for b in catalog] == ["blank", "from_bundle", "from_confluence", "clone"]
    assert all(b["available"] for b in catalog) and StpaFactory.builders() == [b["name"] for b in catalog]
    blank, warnings = brain.stpa_factory.create("blank", "test", name="By Hand", system="A gate")
    assert blank["slug"] == "by_hand" and blank["source_kind"] == "hand" and blank["system"] == "A gate"
    assert blank["counts"] == dict.fromkeys(PARTS, 0) and warnings == []
    clone, warnings = brain.stpa_factory.create("clone", "test", path="conveyor_gate", name="Gate rev B")
    assert clone["slug"] == "gate_rev_b" and clone["counts"] == COUNTS and clone["source_kind"] == "hand"
    assert clone["source_page_id"] is None and clone["note"] == "duplicate of conveyor_gate"
    cloned = brain.service("stpa_uca").list({"analysis": "gate_rev_b"})
    original = brain.service("stpa_uca").list({"analysis": "conveyor_gate"})
    assert [c["code"] for c in cloned] == [o["code"] for o in original]
    assert {c["uid"] for c in cloned}.isdisjoint({o["uid"] for o in original}) and cloned[0]["cite"] == original[0][
        "cite"
    ]
    with pytest.raises(Exception, match="needs --name"):
        brain.stpa_factory.create("blank", "test")
    with pytest.raises(NotFound, match="unknown builder"):
        brain.stpa_factory.create("magic", "test")
    duplicate, _ = brain.stpa.duplicate("conveyor_gate", "Gate rev C")
    assert duplicate["slug"] == "gate_rev_c" and brain.storage.audit_tail(1)[0]["action"] == "stpa.duplicate"


@pytest.mark.skipif(not (HOPPER_FIXTURES / "index.json").is_file(), reason="asi_hopper fixtures are not checked out")
def test_from_confluence_on_hoppers_replay_gateway(brain, tmp_path, monkeypatch):
    monkeypatch.delenv("ASI_HOPPER_CONFIG", raising=False)
    monkeypatch.setenv("ASI_HOPPER_GATEWAY", "replay")
    monkeypatch.setenv("ASI_HOPPER_FIXTURES", str(HOPPER_FIXTURES))
    with Brain(tmp_path / "replay.db") as b:
        made = b.execute("stpa", action="new", target="from_confluence", page="399343803", name="Synthetic")
        assert made.ok, made.errors
        # the synthetic page of asi_hopper's fixtures: 5 hazards, 8 UCAs (one retired), 12 rows, 7 constraints; the
        # bundle landed in the evidence folder beside the database, through Kletz's own Confluence door
        assert made.data["counts"] == {"losses": 0, "hazards": 5, "ucas": 8, "scenarios": 12, "constraints": 7}
        assert made.data["source_page_id"] == "399343803" and made.data["source_version"] == 45
        bundle = Path(made.data["source_bundle"])
        assert bundle.parent == tmp_path / "evidence" and (bundle / "manifest.json").is_file()
        assert (bundle / "pages" / "399343803.md").is_file()
        assert any("Braking STPA" in w for w in made.warnings)
        ucas = b.service("stpa_uca").list()
        assert sum(1 for u in ucas if u["retired"]) == 1 and [u["duplicate_of"] for u in ucas if u["retired"]] == [
            "UCA7"
        ]
        # a refresh against the same page reports nothing
        report = b.execute("stpa", action="refresh", target="synthetic")
        assert report.ok, report.errors
        assert report.data["same"] is True and report.data["applied"] is False
        assert report.data["new"] == {} and report.data["changed"] == {} and report.data["retired"] == {}
        assert report.data["source"]["source_version"] == 45
        # no session on a gateway that has no recording: a failure that names the fix, never a browser
        assert "needs --page" in b.execute("stpa", action="new", target="from_confluence").errors[0]
        assert "digits" in b.execute("stpa", action="new", target="from_confluence", page="abc").errors[0]
        missing = b.execute("stpa", action="new", target="from_confluence", page="1")
        assert not missing.ok and "Confluence read failed" in missing.errors[0]


# -- the rules ---------------------------------------------------------------------------------------
def test_rules_of_the_parts(loaded):
    brain, row, _ = loaded
    uca = brain.service("stpa_uca")
    with pytest.raises(ValidationError, match="no STPA analysis has the slug"):
        uca.create({"analysis": "nope", "code": "UCA9"})
    with pytest.raises(ValidationError, match="already a STPA UCA"):
        uca.create({"analysis": "conveyor_gate", "code": "UCA1"})
    with pytest.raises(ValidationError, match="category"):
        uca.create({"analysis": "conveyor_gate", "code": "UCA9", "category": "sideways"})
    with pytest.raises(ValidationError, match="code: a value is required"):
        uca.create({"analysis": "conveyor_gate", "code": " "})
    made, warnings = uca.create(
        {"analysis": "conveyor_gate", "code": "UCA9", "hazards": ["H1", "H77"], "duplicate_of": "UCA55"}
    )
    assert made["order"] == 9 and made["hazards"] == ["H1", "H77"]
    assert warnings == [
        "UCA9: hazards names H77, which the analysis has no hazard for",
        "UCA9: duplicate of UCA55, which is not a UCA of the analysis",
    ]
    assert brain.events.last(1)[0]["topic"] == "stpa.changed"
    # a code other rows point at is immutable; a free one moves
    uca1 = next(u for u in uca.list({"analysis": "conveyor_gate"}) if u["code"] == "UCA1")
    with pytest.raises(ValidationError, match="other rows point at it"):
        uca.update(uca1["id"], {"code": "UCA100"})
    uca.update(made["id"], {"code": "UCA10"})
    hazard = brain.service("stpa_hazard")
    h5, warnings = hazard.create({"analysis": "conveyor_gate", "code": "H2.5", "text": "sub", "losses": ["L9"]})
    assert h5["parent"] == "H2" and warnings == ["H2.5: losses names L9, which the analysis has no loss for"]
    h2 = next(h for h in hazard.list({"analysis": "conveyor_gate"}) if h["code"] == "H2")
    with pytest.raises(ValidationError, match="other rows point at it"):
        hazard.update(h2["id"], {"code": "H20"})
    scenario = brain.service("stpa_scenario")
    with pytest.raises(ValidationError, match="not a scenario type"):
        scenario.create({"analysis": "conveyor_gate", "uca": "UCA1", "types": [5]})
    s, warnings = scenario.create(
        {
            "analysis": "conveyor_gate",
            "row": 11,
            "uca": "UCA99",
            "types": ["2", 2, "1"],
            "constraints": ["CSTR1", "CSTR99"],
        }
    )
    assert s["types"] == [2, 1]
    assert warnings == [
        "row 11: UCA UCA99 is not a UCA of the analysis",
        "row 11: constraints names CSTR99, which the analysis has no constraint for",
    ]
    assert scenario.create({"analysis": "conveyor_gate", "row": 12})[1] == ["row 12: no UCA"]
    constraint = brain.service("stpa_constraint")
    cstr1 = next(c for c in constraint.list({"analysis": "conveyor_gate"}) if c["code"] == "CSTR1")
    with pytest.raises(ValidationError, match="other rows point at it"):
        constraint.update(cstr1["id"], {"code": "CSTR100"})
    # the analysis: the source fields are read only once read from Confluence; the slug while parts exist
    analysis = brain.service(ANALYSIS_ENTITY)
    with pytest.raises(ValidationError, match="read from Confluence"):
        analysis.update(row["id"], {"source_version": 8})
    with pytest.raises(ValidationError, match="the parts point at it"):
        analysis.update(row["id"], {"slug": "other"})
    with pytest.raises(ValidationError, match="no control structure has the slug"):
        analysis.update(row["id"], {"control_structure": "nope"})
    updated, _ = analysis.update(row["id"], {"status": "Reviewed", "note": "kept"})
    assert updated["status"] == "Reviewed"
    with pytest.raises(ValidationError, match="status"):
        analysis.update(row["id"], {"status": "Shipped"})
    assert code_key("UCA10") > code_key("UCA2") and code_key("H4.3") > code_key("H4") and code_key("x") == ("x", ())


def test_delete_is_refused_without_cascade_and_cascades_with_it(loaded):
    brain, row, _ = loaded
    refused = brain.execute("stpa", action="delete", target="conveyor_gate")
    assert not refused.ok and "delete with --cascade" in refused.errors[0]
    gone = brain.execute("stpa", action="delete", target=row["id"], cascade=True)
    assert gone.ok and gone.data["deleted"] == {ANALYSIS_ENTITY: 1, **COUNTS}
    assert all(brain.service(e).count() == 0 for e in STPA_ENTITIES)
    assert any(e["action"] == "stpa.delete" for e in brain.storage.audit_tail(5))
    assert brain.execute("stpa", action="delete", target="conveyor_gate").errors[0].startswith("STPA analysis")


def test_validate_reports_what_the_rows_say(loaded):
    brain, row, _ = loaded
    report = brain.stpa.validate("conveyor_gate")
    assert report["ok"] and report["errors"] == [] and report["counts"] == COUNTS
    assert "UCA2: hazard H9 is not a hazard of the analysis" in report["warnings"]
    assert "UCA4: no loss scenario names it" in report["warnings"]
    assert "UCA8: no loss scenario names it" in report["warnings"]
    assert "row 10: no UCA" in report["warnings"]
    assert not any("UCA3" in w and "no loss scenario" in w for w in report["warnings"])  # retired: not expected
    # a constraint only referenced, two UCAs with one text, a dangling parent
    brain.execute("create", entity="stpa_constraint", data={"analysis": "conveyor_gate", "code": "CSTR7"})
    uca5 = next(u for u in brain.service("stpa_uca").list() if u["code"] == "UCA5")
    brain.execute("update", entity="stpa_uca", id=uca5["id"], data={"text": "While a person is in the opening."})
    h1 = next(h for h in brain.service("stpa_hazard").list() if h["code"] == "H1")
    brain.service("stpa_hazard").repo.update(h1["id"], {"parent": "H0"})  # behind the service
    brain.service("stpa_uca").repo.update(uca5["id"], {"category": "sideways"})
    report = brain.execute("stpa", action="validate", target="conveyor_gate")
    assert not report.ok and report.errors == [
        "UCA5: category 'sideways' is not one of not_providing, providing, timing_order, duration"
    ]
    assert "CSTR7: only referenced, never spelled out" in report.warnings
    assert "UCA4, UCA5 have the same text" in report.warnings
    assert "H1: parent H0 is not a hazard of the analysis" in report.warnings


# -- the verb --------------------------------------------------------------------------------------
def test_every_stpa_verb_answers_a_result(brain, bundle, tmp_path, capsys):
    assert brain.execute("stpa", action="list").data == []
    builders = brain.execute("stpa", action="builders").data["builders"]
    assert [b["name"] for b in builders] == ["blank", "from_bundle", "from_confluence", "clone"]
    assert "needs a builder" in brain.execute("stpa", action="new").errors[0]
    assert "needs --from" in brain.execute("stpa", action="new", target="from_bundle").errors[0]
    assert "needs --name" in brain.execute("stpa", action="new", target="blank").errors[0]
    assert "needs an analysis" in brain.execute("stpa", action="get").errors[0]
    assert "not found" in brain.execute("stpa", action="get", target="nope").errors[0]
    made = brain.execute(
        "stpa", action="new", target="from_bundle", **{"from": str(bundle)}, name="Conveyor Gate", actor="agent"
    )
    assert made.ok, made.errors
    assert made.data["counts"] == COUNTS and made.verb == "stpa" and brain.storage.audit_tail(1)[0]["actor"] == "agent"
    listed = brain.execute("stpa", action="list")
    assert [r["slug"] for r in listed.data] == ["conveyor_gate"] and listed.data[0]["counts"] == COUNTS
    got = brain.execute("stpa", action="get", target=made.data["uid"])
    assert got.ok and list(got.data["parts"]) == list(PARTS) and len(got.data["parts"]["ucas"]) == 8
    assert brain.execute("stpa", action="get", target=str(made.data["id"])).data["slug"] == "conveyor_gate"
    assert "needs --structure" in brain.execute("stpa", action="link", target="conveyor_gate").errors[0]
    assert (
        "must be one of"
        in brain.execute(
            "stpa", action="export", target="conveyor_gate", format=["pdf", "exe"], out=str(tmp_path)
        ).errors[0]
    )
    assert (
        "not read from Confluence"
        in brain.execute(
            "stpa",
            action="refresh",
            target=brain.execute("stpa", action="new", target="blank", name="Hand").data["slug"],
        ).errors[0]
    )
    # the command line prints the same Result: the acceptance form, then validate and export
    code = main(["--db", brain.db_path, "stpa", "new", "clone", "--from", "conveyor_gate", "--name", "Gate B"])
    printed = json.loads(capsys.readouterr().out)
    assert code == 0 and printed["ok"] and printed["data"]["slug"] == "gate_b" and printed["data"]["counts"] == COUNTS
    code = main(["--db", brain.db_path, "stpa", "validate", "gate_b"])
    printed = json.loads(capsys.readouterr().out)
    assert code == 0 and printed["ok"] and "UCA4: no loss scenario names it" in printed["warnings"]
    code = main(["--db", brain.db_path, "stpa", "export", "gate_b", "--out", str(tmp_path / "cli"), "--format", "xlsx"])
    printed = json.loads(capsys.readouterr().out)
    assert code == 0 and printed["data"]["formats"] == ["xlsx"] and (tmp_path / "cli" / "gate_b.xlsx").is_file()
    code = main(["--db", brain.db_path, "stpa", "delete", "gate_b"])
    assert code == 1 and "cascade" in json.loads(capsys.readouterr().out)["errors"][0]
    assert "stpa" in brain.status()["commands"]
    assert brain.execute("stpa", action="list").data[0]["slug"] == "conveyor_gate"


# -- refresh -----------------------------------------------------------------------------------------
def test_refresh_reports_nothing_on_the_same_bundle_and_the_changes_on_a_new_one(loaded, bundle, tmp_path):
    brain, row, _ = loaded
    same = brain.execute("stpa", action="refresh", target="conveyor_gate", **{"from": str(bundle)})
    assert same.ok, same.errors
    assert same.data["same"] is True and same.data["applied"] is False
    assert all(same.data[k] == {} for k in ("new", "changed", "retired", "removed"))
    assert same.data["counts"]["ucas"] == {"stored": 8, "page": 8}
    later = write_bundle(
        tmp_path / "bundle" / "conveyor-gate-stpa-2026-11-05",
        changed_body(),
        version=8,
        taken_at="2026-11-05T12:00:00-06:00",
    )
    # Kletz-side fields, to be kept: a note on UCA4, a link on UCA1, the analysis's status
    uca4 = next(u for u in brain.service("stpa_uca").list() if u["code"] == "UCA4")
    uca1 = next(u for u in brain.service("stpa_uca").list() if u["code"] == "UCA1")
    brain.execute("update", entity="stpa_uca", id=uca4["id"], data={"note": "discussed at the gate review"})
    brain.service("stpa_uca").repo.update(uca1["id"], {"cs_controller": "gate_controller", "cs_signal": "CA-10"})
    brain.execute("update", entity=ANALYSIS_ENTITY, id=row["id"], data={"status": "Reviewed"})
    audit_before = len(brain.storage.audit_tail(100))
    report = brain.execute("stpa", action="refresh", target="conveyor_gate", **{"from": str(later)})
    assert report.ok, report.errors
    data = report.data
    assert data["same"] is False and data["applied"] is False
    assert data["new"] == {"hazards": ["H5"]} and data["retired"] == {"ucas": ["UCA8"]} and data["removed"] == {}
    assert list(data["changed"]) == ["ucas", "constraints"]
    assert data["changed"]["ucas"] == [
        {
            "key": "UCA4",
            "fields": {
                "text": {
                    "before": "While a person is in the opening.",
                    "after": "While anyone is in the opening.",
                }
            },
        }
    ]
    assert (
        data["changed"]["constraints"][0]["key"] == "CSTR4"
        and data["changed"]["constraints"][0]["fields"]["text"]["after"] == ""
    )
    assert data["source"]["source_version"] == 8 and data["source"]["source_bundle"] == str(later)
    assert len(brain.storage.audit_tail(100)) == audit_before  # a report writes nothing
    assert brain.service("stpa_uca").get(uca4["id"])["text"] == "While a person is in the opening."
    applied = brain.execute(
        "stpa", action="refresh", target="conveyor_gate", apply=True, **{"from": str(later)}, actor="agent"
    )
    assert applied.ok, applied.errors
    assert applied.data["applied"] is True
    assert applied.data["touched"] == {
        "created": 1,
        "updated": 2,
        "retired": 1,
        "removed": 0,
        "cited": sum(COUNTS.values()) - 3,  # every object but the two updated and the one retired
    }
    after4 = brain.service("stpa_uca").get(uca4["id"])
    assert after4["text"] == "While anyone is in the opening." and after4["note"] == "discussed at the gate review"
    assert after4["cite"] == f"page {PAGE} v8 § Context Table ({URL}#Context-Table)"
    after1 = brain.service("stpa_uca").get(uca1["id"])
    assert after1["cs_controller"] == "gate_controller" and after1["cs_signal"] == "CA-10"
    assert after1["cite"].startswith(f"page {PAGE} v8")  # the baseline moved on every object
    assert next(u for u in brain.service("stpa_uca").list() if u["code"] == "UCA8")["retired"] is True
    assert next(h for h in brain.service("stpa_hazard").list() if h["code"] == "H5")["text"] == "A new hazard"
    assert next(c for c in brain.service("stpa_constraint").list() if c["code"] == "CSTR4")["text"] is None
    head = brain.stpa.resolve("conveyor_gate")
    assert head["source_version"] == 8 and head["source_bundle"] == str(later) and head["status"] == "Reviewed"
    assert head["source_body_sha256"] == hashlib.sha256(changed_body().encode("utf-8")).hexdigest()
    actions = [e["action"] for e in brain.storage.audit_tail(6)]
    assert actions[0] == "stpa.refresh" and actions.count("update") >= 3 and "create" in actions
    assert brain.storage.audit_tail(1)[0]["actor"] == "agent"
    event = brain.events.last(1)[0]
    assert event["topic"] == "stpa.changed" and event["action"] == "refresh"
    # applied twice: nothing left to do
    again = brain.execute("stpa", action="refresh", target="conveyor_gate", **{"from": str(later)})
    assert again.data["same"] is True


# -- link --structure ----------------------------------------------------------------------------------
def test_link_structure_proposes_exact_matches_and_applies_them(loaded):
    brain, row, _ = loaded
    small_structure(brain)
    proposed = brain.execute("stpa", action="link", target="conveyor_gate", structure="gate")
    assert proposed.ok, proposed.errors
    data = proposed.data
    assert data["structure"] == "gate" and data["applied"] is None
    assert (data["exact"], data["partial"], data["none"]) == (5, 3, 0)
    by_code = {u["code"]: u for u in data["ucas"]}
    assert by_code["UCA1"]["status"] == "exact" and by_code["UCA1"]["cs_controller"] == "gate_controller"
    assert by_code["UCA1"]["cs_signal"] == "CA-10"
    assert by_code["UCA6"]["status"] == "partial" and by_code["UCA6"]["cs_controller"] == "supervisor"
    assert by_code["UCA6"]["cs_signal"] is None and by_code["UCA6"]["candidates"]["signals"] == [
        "CA-01 Hold request override"
    ]
    assert len(proposed.warnings) == 3 and proposed.warnings[0].startswith("UCA6: partial match")
    assert brain.stpa.resolve("conveyor_gate")["control_structure"] is None  # nothing written yet
    assert all(u["cs_controller"] is None for u in brain.service("stpa_uca").list())
    applied = brain.execute("stpa", action="link", target="conveyor_gate", structure="gate", apply=True)
    assert applied.ok and applied.data["applied"] == 5
    ucas = {u["code"]: u for u in brain.service("stpa_uca").list()}
    assert ucas["UCA1"]["cs_controller"] == "gate_controller" and ucas["UCA1"]["cs_signal"] == "CA-10"
    assert ucas["UCA6"]["cs_controller"] is None  # a partial match is never written
    assert brain.stpa.resolve("conveyor_gate")["control_structure"] == "gate"
    assert brain.storage.audit_tail(1)[0]["action"] == "stpa.link"
    # the structure the analysis names is the default; validate names the controller the structure lacks
    again = brain.execute("stpa", action="link", target="conveyor_gate")
    assert again.ok and again.data["structure"] == "gate" and again.data["exact"] == 5
    brain.service("stpa_uca").repo.update(ucas["UCA6"]["id"], {"controller": "Night Shift Lead"})
    report = brain.stpa.validate("conveyor_gate")
    assert "UCA6: controller 'Night Shift Lead' is not in structure gate" in report["warnings"]
    brain.service("stpa_uca").repo.update(ucas["UCA1"]["id"], {"cs_controller": "ghost"})
    assert "UCA1: cs_controller 'ghost' is not in gate" in brain.stpa.validate("conveyor_gate")["errors"]
    assert match_controller("", [])[0] is None and match_controller("x", [{"name": "x", "key": "x"}])[0] is not None
    assert "not found" in brain.execute("stpa", action="link", target="conveyor_gate", structure="nope").errors[0]


# -- the views -----------------------------------------------------------------------------------------
def test_the_report_and_the_trace_render(loaded, tmp_path):
    brain, row, _ = loaded
    catalog = brain.execute("document").data["documents"]
    assert next(s for s in catalog if s["key"] == "stpa_report")["params"][0]["name"] == "analysis"
    assert "stpa_trace" in [s["key"] for s in brain.execute("diagram").data["diagrams"]]
    built = brain.execute("document", name="stpa_report", param=["analysis=conveyor_gate"], out=str(tmp_path / "doc"))
    assert built.ok, built.errors
    assert built.warnings == [] and built.data["pages"] >= 2
    md = Path(built.data["paths"]["md"]).read_text(encoding="utf-8")
    for heading in (
        "Purpose",
        "Source",
        "Losses",
        "Hazards",
        "Unsafe control actions",
        "Loss scenarios",
        "Constraints",
        "Trace",
        "Citations",
        "Review comments",
    ):
        assert heading in md, heading
    assert "### 5.1 Gate Controller" in md and "### 5.2 Supervisor" in md
    assert "| UCA3 | Close Gate | Not providing | When the line restarts after a fault. |" in md
    assert "| H2 | retired; duplicate of UCA6 |" in md
    assert f"page {PAGE} v7 § Context Table ({URL}#Context-Table)" in md
    assert "| CSTR1 | The gate controller must confirm the gate position before every run. | 1, 2, 4, 5 |" in md
    assert "Should the hold expire at all?" in md and "| reply | 2026-10-21 |" in md and "Yes, after the shift." in md
    assert (
        f"[{PAGE}]({URL})" in md
        and "![Figure 1: STPA trace: Conveyor Gate](figures/stpa_trace_conveyor_gate.png)" in md
    )
    assert (tmp_path / "doc" / "figures" / "stpa_trace_conveyor_gate.png").is_file()
    assert "| H2.1 | H2 | Nearby workers are unaware that the belt is about to run |  |  |" in md
    assert "| H4 |  | An old hazard, retired |  | yes |" in md
    # the trace: four columns, left to right, no lanes without a linked structure
    drawn = brain.execute("diagram", name="stpa_trace", param=["analysis=conveyor_gate"], out=str(tmp_path / "pic"))
    assert drawn.ok, drawn.errors
    assert drawn.data["nodes"] == sum(COUNTS[k] for k in ("hazards", "ucas", "scenarios", "constraints"))
    from asi_kletz.generate import BuildContext, DiagramFactory

    diagram = DiagramFactory.build("stpa_trace", brain, {"analysis": "conveyor_gate"}, BuildContext(brain))
    assert diagram.direction == "LR" and diagram.lanes == []
    kinds = {n.kind for n in diagram.nodes.values()}
    assert kinds == {"io", "process", "action", "document"}
    assert "«retired»" in diagram.nodes["u_uca3"].label and diagram.nodes["u_uca3"].fill is not None
    dashed = [e for e in diagram.edges if e.style == "dashed"]
    assert len(dashed) == 1 and (dashed[0].src, dashed[0].dst) == ("u_uca3", "u_uca6")
    assert drawn.data["edges"] == len(diagram.edges)
    # linked to a structure: a swimlane per controller, hazards and constraints in lanes of their own
    small_structure(brain)
    brain.execute("stpa", action="link", target="conveyor_gate", structure="gate", apply=True)
    diagram = DiagramFactory.build("stpa_trace", brain, {"analysis": "conveyor_gate"}, BuildContext(brain))
    assert diagram.direction == "TB" and [ln.title for ln in diagram.lanes] == [
        "Hazards",
        "Gate Controller",
        "Not linked",
        "Constraints",
    ]
    assert diagram.nodes["u_uca1"].lane == "lane_gate_controller" and diagram.nodes["u_uca6"].lane == "lane_unlinked"
    assert diagram.nodes["h_h1"].lane == "lane_hazards" and diagram.nodes["c_cstr1"].lane == "lane_constraints"
    lanes = brain.execute("diagram", name="stpa_trace", param=["analysis=conveyor_gate"], out=str(tmp_path / "lanes"))
    assert lanes.ok, lanes.errors
    empty = brain.execute("stpa", action="new", target="blank", name="Empty").data["slug"]
    assert brain.execute("diagram", name="stpa_trace", param=[f"analysis={empty}"], out=str(tmp_path / "e")).ok
    assert brain.execute("document", name="stpa_report", param=[f"analysis={empty}"], out=str(tmp_path / "e")).ok
    assert "missing required param" in brain.execute("diagram", name="stpa_trace", out=str(tmp_path)).errors[0]


def test_export_writes_the_six_sheets_the_report_and_the_trace(loaded, tmp_path):
    brain, row, _ = loaded
    assert DEFAULT_FORMATS == ("xlsx", "md", "png") and FORMATS == ("xlsx", "md", "docx", "pdf", "svg", "png")
    exported = brain.execute("stpa", action="export", target="conveyor_gate", out=str(tmp_path / "out"), actor="agent")
    assert exported.ok, exported.errors
    data = exported.data
    assert data["formats"] == list(DEFAULT_FORMATS) and sorted(data["paths"]) == ["md", "png", "xlsx"]
    assert all(Path(p).is_file() for p in data["files"]) and data["primary"].endswith("conveyor_gate_trace.png")
    wb = load_workbook(data["paths"]["xlsx"])
    assert wb.sheetnames == [
        "STPA analyses",
        "STPA losses",
        "STPA hazards",
        "STPA UCAs",
        "STPA constraints",
        "STPA scenarios",
    ]
    ucas = wb["STPA UCAs"]
    headers = [c.value for c in ucas[1]]
    assert headers[:4] == ["analysis", "order", "code", "controller"] and headers[-1] == "uid"
    assert ucas.max_row == 1 + COUNTS["ucas"]
    assert ucas.column_dimensions[get_column_letter(len(headers))].hidden  # the uid column, last and hidden
    assert ucas.cell(row=2, column=headers.index("hazards") + 1).value == '["H2", "H2.1"]'
    assert wb["STPA analyses"].max_row == 2 and wb["STPA scenarios"].max_row == 1 + COUNTS["scenarios"]
    tail = brain.storage.audit_tail(1)[0]
    assert (
        tail["action"] == "stpa.export"
        and tail["actor"] == "agent"
        and sorted(tail["payload"]["paths"]) == ["md", "png", "xlsx"]
    )
    event = brain.events.last(1)[0]
    assert event["topic"] == "stpa.exported" and event["analysis"] == "conveyor_gate"
    everything = brain.execute("stpa", action="export", target="conveyor_gate", out=str(tmp_path / "all"), all=True)
    assert everything.ok and sorted(everything.data["paths"]) == sorted(FORMATS)
    assert everything.data["primary"].endswith("conveyor_gate_report.pdf")
    two = brain.execute("stpa", action="export", target="conveyor_gate", format=["svg", "docx"])
    assert two.ok and two.data["formats"] == ["svg", "docx"] and two.data["out"].endswith("stpa")
    assert Path(two.data["out"]).parent == Path(brain.out_dir)


# -- transfers and the migration ------------------------------------------------------------------------
def test_dump_load_and_the_workbook_carry_the_six_entities(loaded, tmp_path):
    brain, row, _ = loaded
    dump = tmp_path / "dump.json"
    assert brain.execute("dump", out=str(dump)).ok
    payload = json.loads(dump.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 11 and {len(payload["entities"][e]) for e in PART_ENTITIES} == {2, 5, 8, 6, 10}
    assert brain.execute("stpa", action="delete", target="conveyor_gate", cascade=True).ok
    loaded_back = brain.execute("load", path=str(dump), mode="merge")
    assert loaded_back.ok, loaded_back.errors
    assert brain.execute("stpa", action="list").data[0]["counts"] == COUNTS
    back = next(u for u in brain.service("stpa_uca").list() if u["code"] == "UCA1")
    assert back["hazards"] == ["H2", "H2.1"] and back["uid"]
    exported = brain.execute("export", format="excel", out=str(tmp_path / "view.xlsx"))
    wb_path = Path(exported.data["path"])  # 0.27.0: the workbook is written under its stamped name
    assert exported.ok and wb_path.name.startswith("view_")
    assert "STPA UCAs" in load_workbook(wb_path).sheetnames
    assert brain.execute("import", format="excel", path=str(wb_path), mode="replace").ok
    assert brain.execute("stpa", action="list").data[0]["counts"] == COUNTS


def test_a_schema_10_database_opens_as_0_24_0_with_the_six_stpa_tables_added(tmp_path, bundle):
    db = tmp_path / "kletz.db"
    shutil.copy(DB_0_23_0, db)
    con = sqlite3.connect(db)
    before = {
        name: [dict(zip([c[0] for c in cur.description], r, strict=True)) for r in cur.fetchall()]
        for name in ("safety_plan", "parameter")
        for cur in [con.execute(f'SELECT * FROM "{name}"')]
    }
    tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    con.close()
    assert "stpa_analyses" not in tables and "safety_plan" in tables and len(before["safety_plan"]) == 3
    with Brain(db) as b:
        assert b.migration["created"][-6:] == STPA_TABLES  # after the seven control-structure tables
        assert b.migration["altered"] == {} and b.migration["uids_stamped"] == EMPTY_COUNTS
        versions = [
            m["app_version"] for m in b.storage.fetchall("SELECT app_version FROM schema_migrations ORDER BY version")
        ]
        assert versions == ["0.23.0", __version__] and b.status()["schema_version"] == 11
        assert b.execute("stpa", action="list").data == []
        assert b.execute("list", entity="safety_plan_row").data[0]["deliverable"] == "Safety Plan"
        made = b.execute("stpa", action="new", target="from_bundle", **{"from": str(bundle)}, name="Conveyor Gate")
        assert made.ok and made.data["counts"] == COUNTS
    con = sqlite3.connect(db)
    for name, rows in before.items():
        cur = con.execute(f'SELECT * FROM "{name}"')
        assert [dict(zip([c[0] for c in cur.description], r, strict=True)) for r in cur.fetchall()] == rows
    con.close()
    with Brain(db) as b:
        assert b.migration["created"] == [] and len(b.storage.fetchall("SELECT 1 FROM schema_migrations")) == 2


# -- the window -----------------------------------------------------------------------------------------
def _tk_available() -> bool:
    import gc
    import tkinter as tk

    try:
        root = tk.Tk()
        root.withdraw()
        root.destroy()
        del root
        gc.collect()
        return True
    except tk.TclError:
        return False


@pytest.mark.skipif(not _tk_available(), reason="no display for tkinter")
def test_the_stpa_screen_creates_validates_refreshes_and_exports(brain, bundle, tmp_path):
    import gc

    from asi_kletz.gui.dashboard import STPA_TILES
    from asi_kletz.gui.main_window import MainWindow
    from asi_kletz.gui.screens import STPA_SCREEN

    assert STPA_TILES[1] == ("stpa_analysis", "STPA ANALYSES", STPA_SCREEN)  # the analyses, never the parts
    window = MainWindow(brain, interactive=False)
    window.withdraw()
    try:
        keys = [k for k, _ in window.factory.items()]
        assert STPA_SCREEN in keys and not any(k.startswith("entity:") for k in keys)
        dash = window.show("dashboard")
        assert "stpa_analysis" in dash._stpa_vars and "stpa_uca" not in dash._stpa_vars
        assert window.server.show_screen(STPA_SCREEN).ok  # /screen {"show": "stpa"}
        screen = window.show(STPA_SCREEN)
        assert (
            window.title().endswith("STPA analyses") and screen.rows == [] and screen.builder_var.get() == "from_bundle"
        )
        assert list(screen.builder_box["values"]) == ["blank", "from_bundle", "from_confluence", "clone"]
        assert screen.validate() is None and "select an analysis" in window.status_text()
        screen.name_var.set("Conveyor Gate")
        screen.from_var.set(str(bundle))
        created = screen.create()
        assert created.ok, created.errors
        assert screen.selected == "conveyor_gate" and screen.tree.exists("conveyor_gate")
        assert brain.storage.audit_tail(1)[0]["actor"] == "gui"
        assert len(screen.part_trees["stpa_uca"].get_children()) == 8
        assert len(screen.part_trees["stpa_scenario"].get_children()) == 10
        columns = set(screen.tree.tk.splitlist(screen.tree["columns"]))
        assert not columns & {"id", "order", "uid"}
        for tree in screen.part_trees.values():
            assert not set(tree.tk.splitlist(tree["columns"])) & {"id", "order", "uid"}
        assert "[ok] stpa new from_bundle" in screen.log.get("1.0", "end")
        validated = screen.validate()
        assert validated.ok and "UCA4: no loss scenario names it" in screen.log.get("1.0", "end")
        exported = screen.export(str(tmp_path / "exp"))
        assert exported.ok and (tmp_path / "exp" / "conveyor_gate.xlsx").is_file()
        # Refresh from Confluence runs the report off the Tk thread, shows it, then applies (confirm says yes)
        later = write_bundle(tmp_path / "bundle" / "later", changed_body(), version=8)
        assert screen.start_refresh(str(later)) is True and screen.refreshing
        assert screen.start_refresh(str(later)) is False  # one at a time
        deadline = time.time() + 20
        while screen.refreshing and time.time() < deadline:
            window.update()
            time.sleep(0.02)
        assert not screen.refreshing and screen.last_report is not None
        assert screen.last_report["new"] == {"hazards": ["H5"]} and screen.last_report["retired"] == {"ucas": ["UCA8"]}
        log = screen.log.get("1.0", "end")
        assert "new hazards: H5" in log and "UCA4 text:" in log and "[ok] stpa refresh conveyor_gate --apply" in log
        assert brain.stpa.resolve("conveyor_gate")["source_version"] == 8
        assert "refreshed" in window.status_text()
        # the same bundle again: nothing to apply
        assert screen.start_refresh(str(later))
        deadline = time.time() + 20
        while screen.refreshing and time.time() < deadline:
            window.update()
            time.sleep(0.02)
        assert screen.last_report["same"] is True and "unchanged" in window.status_text()
        deleted = screen.delete()
        assert deleted.ok and screen.rows == [] and screen.selected is None
    finally:
        window.close()
        del window
        gc.collect()


# -- 0.10.1: the shapes the real page has and the first synthetic page lacked ----------------------------
QUIRKS_FILE = REPO_ROOT / "tests" / "fixtures" / "conveyor-gate-stpa-quirks.md"
QUIRKS_BODY = QUIRKS_FILE.read_text(encoding="utf-8")
QUIRKS_COUNTS = {"losses": 0, "hazards": 4, "ucas": 14, "scenarios": 6, "constraints": 6}


def test_reader_reads_rows_whose_controller_cell_was_merged_with_the_row_above():
    """A controller cell merged with the row above is dropped by the renderer, so the row is short and its
    cells shift left: the first cell holding UCA content says where the categories start, the controller is
    carried down, the control action is the row's first cell."""
    data = read_stpa_page(QUIRKS_BODY, page_id="900000002", url=URL, version=3)
    assert {k: len(data[k]) for k in QUIRKS_COUNTS} == QUIRKS_COUNTS
    ucas = {u["code"]: u for u in data["ucas"]}
    assert ucas["UCA7"]["controller"] == "Gate Controller" and ucas["UCA7"]["control_action"] == "Hold Gate"
    assert ucas["UCA7"]["category"] == "providing" and ucas["UCA8"]["category"] == "timing_order"
    assert ucas["UCA10"]["control_action"] == "Signal Lamp" and ucas["UCA10"]["category"] == "not_providing"
    assert ucas["UCA13"]["category"] == "providing" and ucas["UCA14"]["category"] == "duration"
    assert ucas["UCA6"]["controller"] == "Supervisor" and ucas["UCA6"]["control_action"] == "Hold Request"
    shifted = [w for w in data["warnings"] if "merged with the row above" in w]
    assert shifted == [
        "Context Table row 2: the controller cell is merged with the row above; read as Gate Controller / Hold Gate",
        "Context Table row 3: the controller cell is merged with the row above; read as Gate Controller / Signal Lamp",
    ]
    assert row_shift(("Gate Controller", "Close Gate", "**UCA1 **x", "", "", "")) == 2
    assert row_shift(("Hold Gate", "(duplicate of **UCA6** here)", "**UCA7 **x", "", "", "")) == 1
    assert row_shift(("**UCA9 **x *\\[H1\\]*", "", "", "", "", "")) == 0
    assert row_shift(("Belt Drive", "Run Belt", "[Other](https://example.invalid/wiki/pages/1)", "", "", "")) == 2
    assert row_shift(("", "", "", "", "", "")) is None


def test_reader_never_creates_a_uca_from_a_duplicate_note_and_reads_unbolded_codes():
    data = read_stpa_page(QUIRKS_BODY, page_id="900000002", url=URL, version=3)
    ucas = {u["code"]: u for u in data["ucas"]}
    # "(duplicate of **UCA9**)" after UCA3 in the same cell: UCA3's twin, not a UCA of the cell
    assert ucas["UCA3"]["retired"] is True and ucas["UCA3"]["duplicate_of"] == "UCA9"
    assert ucas["UCA9"]["controller"] == "Supervisor"  # read where it is defined, once
    # "(duplicate of **UCA6** here)" alone in a cell: a warning, nothing created, UCA6 stays the Supervisor's
    assert ucas["UCA6"]["controller"] == "Supervisor" and ucas["UCA6"]["text"] == "When maintenance is in the opening."
    assert (
        "Context Table row 2 (Gate Controller / Hold Gate): the not_providing cell is only a duplicate note of UCA6, "
        "which is defined elsewhere; nothing read"
    ) in data["warnings"]
    # "(duplicate of the signal lamp row)": no code at all, a warning, nothing created
    assert (
        "Context Table row 5 (Belt Drive / Run Belt): the not_providing cell is a duplicate note without a code "
        "('(duplicate of the signal lamp row)'); nothing read"
    ) in data["warnings"]
    # an unbolded code right after a hazard list is a UCA of its own, not the tail of the one before it
    assert ucas["UCA11"]["text"] == "When a line restarts {after a fault}." and ucas["UCA11"]["hazards"] == ["H2.2"]
    assert ucas["UCA12"]["text"] == "When the lamp is masked by dust." and ucas["UCA12"]["hazards"] == ["H2.1"]
    assert embolden_loose_codes("**UCA22** t *\\[H2, H4.5\\]*UCA28 more *\\[H4.5\\]*") == (
        "**UCA22** t *\\[H2, H4.5\\]***UCA28 **more *\\[H4.5\\]*"
    )
    assert embolden_loose_codes("UCA8 plain") == "**UCA8 **plain" and embolden_loose_codes("**UCA8 **x") == "**UCA8 **x"
    # a code read twice: the later one is skipped with a warning naming both rows, the load never fails
    assert [u["code"] for u in data["ucas"]].count("UCA6") == 1
    assert (
        "Context Table row 6 (Belt Drive / Lamp Test): UCA6 was already read in row 4; the later one is skipped"
        in data["warnings"]
    )
    assert ucas["UCA10"]["hazards"] == ["H2.1", "H9"]
    assert "UCA10: cites hazard H9, which the page does not list" in data["warnings"]
    # the scenario table's unbolded "UCA12 ..." row names UCA12, which the context table now holds
    assert not any("not in the context table" in w for w in data["warnings"])
    cell, warnings = read_uca_cell(
        "**UCA1 **x (duplicate of",
        who="row 9",
        category="providing",
        controller="c",
        action="a",
        where="w",
        hazard_codes=set(),
        seen={},
        row_number=9,
        order=1,
    )
    assert [u["code"] for u in cell] == ["UCA1"] and warnings == ["row 9: UCA1 says 'duplicate of' but names no UCA"]
    cell, warnings = read_uca_cell(
        "(duplicate of",
        who="row 9",
        category="providing",
        controller="c",
        action="a",
        where="w",
        hazard_codes=set(),
        seen={},
        row_number=9,
        order=1,
    )
    assert cell == [] and warnings == ["row 9: the providing cell is a duplicate note that names no UCA; nothing read"]


def test_reader_tolerates_constraint_lists_duplicate_notes_and_broken_escapes():
    data = read_stpa_page(QUIRKS_BODY, page_id="900000002", url=URL, version=3)
    rows = {s["row"]: s for s in data["scenarios"]}
    assert rows[2]["constraints"] == ["CSTR1", "CSTR2", "CSTR3"]  # **CSTR\[1,2\]** expands to references
    assert rows[3]["constraints"] == ["CSTR3"]  # duplicate **CSTR3**: a reference
    assert rows[4]["constraints"] == ["CSTR4", "CSTR2", "CSTR1", "CSTR3"]  # a lone code, "see also", a duplicate list
    assert rows[5]["constraints"] == ["CSTR5"]  # **CSTR\[5 **: broken escape, read as a reference with a warning
    assert rows[6]["constraints"] == ["CSTR5", "CSTR6"]  # **duplicate CSTR5 and CSTR6**
    constraints = {c["code"]: c for c in data["constraints"]}
    assert constraints["CSTR4"]["text"] == "The closing state must be reported every cycle."
    assert constraints["CSTR2"]["text"] == "A stuck position sensor must be treated as a fault."  # row 1 spelled it
    assert constraints["CSTR5"]["text"] == "" and constraints["CSTR6"]["text"] == ""
    assert (
        "Loss Scenarios row 5: broken constraint list CSTR\\[5 (no closing bracket); read as CSTR5"
        in (data["warnings"])
    )
    assert constraint_codes("CSTR\\[66, 10, 11\\] and CSTR 7") == (["CSTR66", "CSTR10", "CSTR11", "CSTR7"], [])
    codes, warnings = constraint_codes("CSTR\\[63, 64")
    assert codes == ["CSTR63", "CSTR64"] and warnings and "no closing bracket" in warnings[0]
    assert read_constraint_cell("duplicate **CSTR61**") == [("CSTR61", "")]
    assert read_constraint_cell("**duplicate CSTR73 and CSTR74**") == [("CSTR73", ""), ("CSTR74", "")]
    assert read_constraint_cell("**CSTR81 **text (more) **see CSTR55****duplicate CSTR56**") == [
        ("CSTR81", "text (more)"),
        ("CSTR55", ""),
        ("CSTR56", ""),
    ]
    assert read_constraints("**CSTR\\[63")[1] and read_constraints("")[0] == []
    # the same UCA named twice across rows and a "(same as UCA5 CF)" note are text, not references
    assert rows[2]["causal_factors"].endswith("(same as UCA5 CF)") and rows[1]["uca"] == rows[2]["uca"] == "UCA1"


def test_the_quirky_page_loads_through_the_factory_without_a_hard_failure(brain, tmp_path):
    bundle = write_bundle(tmp_path / "quirks", QUIRKS_BODY, version=3, page_id="900000002", comments=[])
    made = brain.execute("stpa", action="new", target="from_bundle", **{"from": str(bundle)}, name="Quirks")
    assert made.ok, made.errors
    assert made.data["counts"] == QUIRKS_COUNTS
    assert sum(1 for u in brain.service("stpa_uca").list() if u["retired"]) == 1
    assert [u["duplicate_of"] for u in brain.service("stpa_uca").list() if u["duplicate_of"]] == ["UCA9"]
    assert any("was already read in row 4" in w for w in made.warnings)
    assert any("only a duplicate note of UCA6" in w for w in made.warnings)
    assert brain.execute("stpa", action="validate", target="quirks").ok
    same = brain.execute("stpa", action="refresh", target="quirks", **{"from": str(bundle)})
    assert same.ok and same.data["same"] is True
