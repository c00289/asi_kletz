"""0.22.0, Phase A of the Confluence integration plan (asi_hopper/docs/confluence-integration-plan.md): "URLs
that know what they are". The owner, 2026-10-29: "when we add confluence read only into kletz we should be able
to associate objects with [confluence] urls; obviously the next candidate is replacing the ... placeholder with a
real url once known".

Every kind from the host and the path, the optional host map of kletz.toml, the labels, ``plan urls --audit`` on
the seeded database, ``plan urls --replace`` in the three forms (Python, the wire format, the CLI) with its dry
run and its refusals, the Excel Work product cell and the form's link lines. Nothing here reaches Confluence
or Jira: the kind is read off the URL."""

from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from asi_kletz import Brain
from asi_kletz.brain import ENV_CONFIG, load_config
from asi_kletz.schema.definitions import WORK_PRODUCT_HEADER
from asi_kletz.schema.urls import URL_KINDS, jira_issue_key, kind_label, url_kind, url_kinds
from tests.conftest import PLAN
from tests.test_cli import run

MS = "https://www.microsoft.com/"
WIKI = "https://asirobots.atlassian.net/wiki/spaces/STPA/pages/399343803/Locator+Safety+Function+STPA"
WIKI_VIEW = "https://asirobots.atlassian.net/wiki/pages/viewpage.action?pageId=399343803"
BROWSE = "https://asirobots.atlassian.net/browse/HWL-2554"
BOARD = "https://asirobots.atlassian.net/jira/software/c/projects/HWL/boards/12?selectedIssue=HWL-2572"
ISSUES = "https://asirobots.atlassian.net/jira/software/c/projects/HWL/issues/HWL-10"
SHAREPOINT = "https://asirobots.sharepoint.com/sites/FuSa/Shared%20Documents/plan.docx"
JAMA = "https://jama.example.com/items/2554"
DMS = "https://docs.dms.example.org/doc/17"
HOSTS = {"jama": "jama.example.com", "dms": ["dms.example.com", "*.dms.example.org"]}


# -- the pure functions -------------------------------------------------------------------------------------


def test_every_kind_is_decided_from_the_host_and_the_path():
    assert URL_KINDS == ("confluence", "jira", "sharepoint", "placeholder", "other")
    assert url_kind(WIKI) == "confluence" and url_kind(WIKI_VIEW) == "confluence"
    assert url_kind("https://asirobots.atlassian.net/wiki") == "confluence"
    assert url_kind(BROWSE) == "jira" and url_kind(BOARD) == "jira" and url_kind(ISSUES) == "jira"
    assert url_kind("HTTPS://ASIROBOTS.ATLASSIAN.NET/browse/HWL-1") == "jira"  # the host in any case
    assert url_kind("https://asirobots.atlassian.net/") == "other"  # the site's root is neither
    assert url_kind("https://other.atlassian.net/wiki/spaces/X") == "other"  # another site is not ours
    assert url_kind(SHAREPOINT) == "sharepoint" and url_kind("https://a.b.sharepoint.com/x") == "sharepoint"
    assert url_kind("https://sharepoint.com/") == "other"  # the bare domain is nobody's site
    assert url_kind(MS) == "placeholder" and url_kind("http://www.microsoft.com") == "placeholder"
    assert url_kind("https://learn.microsoft.com/") == "other"  # only the homepage is the placeholder
    assert url_kind(JAMA) == "other" and url_kind(DMS) == "other"  # no Jama kind without the owner's map
    assert url_kind("nope") == "other" and url_kind(None) == "other" and url_kind("") == "other"
    assert url_kinds([MS, WIKI, WIKI_VIEW, BROWSE]) == ["placeholder", "confluence", "jira"]  # each once, in order
    assert url_kinds(json.dumps([JAMA])) == ["other"] and url_kinds(None) == []


def test_the_optional_host_map_names_more_kinds():
    assert url_kind(JAMA, HOSTS) == "jama" and url_kind(DMS, HOSTS) == "dms"
    assert url_kind("https://dms.example.com/doc/1", HOSTS) == "dms"  # the list's first host
    assert url_kind("https://dms.example.org/doc/1", HOSTS) == "other"  # `*.` names the hosts under, not the domain
    assert url_kind("https://JAMA.example.com/x", HOSTS) == "jama"
    assert url_kind(WIKI, HOSTS) == "confluence" and url_kind(MS, HOSTS) == "placeholder"  # the five still decide
    assert url_kind(WIKI, {"wiki": "asirobots.atlassian.net"}) == "wiki"  # the owner's word first
    assert url_kinds([JAMA, DMS, MS], HOSTS) == ["jama", "dms", "placeholder"]


def test_kind_label_reads_the_kind_and_the_host():
    assert kind_label(WIKI) == "Confluence, asirobots.atlassian.net"
    assert kind_label(BROWSE) == "Jira, asirobots.atlassian.net"
    assert kind_label(SHAREPOINT) == "SharePoint, asirobots.sharepoint.com"
    assert kind_label(MS) == "placeholder"  # alone: the row still points nowhere
    assert kind_label(JAMA) == "other, jama.example.com"
    assert kind_label(JAMA, HOSTS) == "jama, jama.example.com" and kind_label(DMS, HOSTS) == "dms, docs.dms.example.org"


def test_jira_issue_key_is_read_off_the_url():
    assert (
        jira_issue_key(BROWSE) == "HWL-2554"
        and jira_issue_key(BOARD) == "HWL-2572"
        and jira_issue_key(ISSUES) == "HWL-10"
    )
    assert jira_issue_key("https://asirobots.atlassian.net/browse/hwl-7") == "HWL-7"
    assert jira_issue_key("https://asirobots.atlassian.net/jira/software/c/projects/HWL/boards/12") is None
    assert jira_issue_key(WIKI) is None and jira_issue_key(None) is None


# -- the config ----------------------------------------------------------------------------------------------


def test_the_host_map_comes_from_kletz_toml(tmp_path, monkeypatch):
    toml = tmp_path / "kletz.toml"
    toml.write_text(
        'db = "k.db"\n[urls]\nhosts = { jama = "jama.example.com", dms = ["dms.example.com", "*.dms.example.org"], '
        "bad = 7 }\n",
        encoding="utf-8",
    )
    monkeypatch.setenv(ENV_CONFIG, str(toml))
    config = load_config()
    loaded = {"jama": ["jama.example.com"], "dms": ["dms.example.com", "*.dms.example.org"]}  # every value a list
    assert config["url_hosts"] == loaded
    assert config["warnings"] == ["[urls] hosts.bad = 7 in kletz.toml is ignored: write a host or a list of hosts"]
    with Brain(tmp_path / "k.db") as brain:
        assert brain.url_hosts == loaded and brain.work_products.hosts == loaded
        assert brain.work_products.kind(JAMA) == "jama" and brain.work_products.kind(WIKI) == "confluence"
        assert brain.status()["config_warnings"] == config["warnings"]
    assert load_config(tmp_path / "absent.toml") == {}  # no file: no map, nothing else either


# -- the audit -----------------------------------------------------------------------------------------------


def test_plan_urls_audit_names_the_findings_and_writes_nothing(seeded):
    # row 1: the placeholder first and a Jira URL that names another ticket than the row's key
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS, BROWSE], "jira_key": "HWL-1"}).ok
    # row 2: the work product location names Confluence, the one URL is a Jama link
    assert seeded.execute(
        "update",
        entity=PLAN,
        id=2,
        data={
            "urls": [JAMA],
            "work_product_location": "Confluence (STPA space)",
        },
    ).ok
    # row 3: Confluence named and a Confluence page linked, the Jira URL agreeing with the key: nothing to say
    assert seeded.execute(
        "update",
        entity=PLAN,
        id=3,
        data={
            "urls": [WIKI, ISSUES],
            "work_product_location": "confluence",
            "jira_key": "hwl-10",
        },
    ).ok
    before = len(seeded.storage.audit_tail(50))
    result = seeded.execute("plan", action="urls", audit=True)
    assert result.ok and result.warnings == [], result.errors  # findings are data, not errors
    assert [(f["id"], f["code"]) for f in result.data] == [(1, "placeholder"), (1, "jira_key"), (2, "no_confluence")]
    first, second, third = result.data
    assert first["finding"] == f"the placeholder {MS} is still present"
    assert first["kinds"] == ["placeholder", "jira"] and first["urls"] == [MS, BROWSE]
    assert first["deliverable"] == "Safety Plan" and first["order"] == 1 and len(first["uid"]) == 36
    assert second["finding"] == "the Jira URL names HWL-2554 and the row's Jira key is HWL-1"
    assert (
        third["finding"] == "the work product location says 'Confluence (STPA space)' and no URL is a Confluence page"
    )
    assert third["kinds"] == ["other"] and third["deliverable"] == "Hazard Analysis and Risk Assessment"
    assert len(seeded.storage.audit_tail(50)) == before  # read only
    # only some rows: ids or uids, repeatable; an unknown row is refused
    uid2 = seeded.execute("get", entity=PLAN, id=2).data["uid"]
    some = seeded.execute("plan", action="urls", audit=True, rows=[3, uid2])
    assert some.ok and [(f["id"], f["code"]) for f in some.data] == [(2, "no_confluence")]
    assert seeded.execute("plan", action="urls", audit=True, rows=3).data == []
    assert seeded.execute("plan", action="urls", audit=True, rows=[9]).errors == ["no plan row 9"]
    # a row with no Jira key: the Jira URL's key disagrees with "empty"; a row whose URLs are gone while its
    # location still names Confluence has no Confluence page either
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [BROWSE], "jira_key": None}).ok
    assert seeded.execute("update", entity=PLAN, id=3, data={"urls": []}).ok
    findings = seeded.execute("plan", action="urls", audit=True).data
    assert [f["finding"] for f in findings if f["id"] == 1] == [
        "the Jira URL names HWL-2554 and the row's Jira key is empty"
    ]
    assert [(f["id"], f["code"]) for f in findings] == [(1, "jira_key"), (2, "no_confluence"), (3, "no_confluence")]
    assert findings[-1]["kinds"] == [] and findings[-1]["urls"] == []


def test_the_audit_uses_the_owners_host_map(tmp_path, monkeypatch):
    toml = tmp_path / "kletz.toml"
    toml.write_text('[urls]\nhosts = { jama = "jama.example.com" }\n', encoding="utf-8")
    monkeypatch.setenv(ENV_CONFIG, str(toml))
    with Brain(tmp_path / "k.db") as brain:
        brain.execute("create", entity=PLAN, data={"order": 1, "deliverable": "HARA", "urls": [JAMA, MS]})
        findings = brain.execute("plan", action="urls", audit=True).data
        assert [f["code"] for f in findings] == ["placeholder"] and findings[0]["kinds"] == ["jama", "placeholder"]


# -- the replacement ------------------------------------------------------------------------------------------


def test_plan_urls_replace_puts_the_new_url_where_the_old_one_stood(seeded):
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [JAMA, MS]}).ok  # the placeholder second
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [MS]}).ok
    assert seeded.execute("update", entity=PLAN, id=3, data={"urls": [JAMA]}).ok  # no placeholder: untouched
    before = len(seeded.storage.audit_tail(50))
    result = seeded.execute("plan", action="urls", replace=[MS, WIKI])
    assert result.ok and result.warnings == [], result.errors
    assert result.data["old"] == MS and result.data["new"] == WIKI and result.data["kind"] == "confluence"
    assert result.data["replaced"] == [1, 2] and result.data["would_replace"] == [] and result.data["refused"] == []
    assert result.data["rows"] == 3 and result.data["check"] is False
    assert [(r["id"], r["position"], r["deliverable"]) for r in result.data["rows_replaced"]] == [
        (1, 2, "Safety Plan"),
        (2, 1, "Hazard Analysis and Risk Assessment"),
    ]
    assert [r["urls"] for r in seeded.execute("list", entity=PLAN).data] == [[JAMA, WIKI], [WIKI], [JAMA]]  # in place
    entries = seeded.storage.audit_tail(50)
    assert len(entries) == before + 1  # one entry for the whole step, no per-row update
    assert entries[0]["action"] == "urls_replace" and entries[0]["entity"] == PLAN
    assert entries[0]["payload"] == {"old": MS, "new": WIKI, "rows": [1, 2], "refused": []}
    event = seeded.events.last(1)[0]
    assert event["topic"] == "plan.urls_replaced" and event["old"] == MS and event["new"] == WIKI
    assert event["rows"] == 2 and event["actor"] == "python"
    # a page that moves is corrected once, the same way
    moved = seeded.execute("plan", action="urls", replace=[WIKI, WIKI_VIEW])
    assert moved.ok and moved.data["replaced"] == [1, 2]
    assert [r["urls"] for r in seeded.execute("list", entity=PLAN).data] == [[JAMA, WIKI_VIEW], [WIKI_VIEW], [JAMA]]
    # the audit says every row is off the placeholder now
    assert seeded.execute("plan", action="urls", audit=True).data == []


def test_plan_urls_replace_dry_run_refusals_and_the_wire_format(seeded):
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS, WIKI]}).ok  # the new URL already here
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [MS]}).ok
    before = len(seeded.storage.audit_tail(50))
    check = seeded.execute("plan", action="urls", replace=[MS, WIKI], check=True)
    assert check.ok and check.data["check"] is True
    assert check.data["would_replace"] == [2] and check.data["replaced"] == [] and check.data["refused"] == [1]
    assert check.warnings == [f"row 1 (Safety Plan): {WIKI} is already on the row; left as it is"]
    assert [r["urls"] for r in seeded.execute("list", entity=PLAN).data] == [[MS, WIKI], [MS], None]  # nothing written
    assert len(seeded.storage.audit_tail(50)) == before
    # the wire format, an agent's payload: the refused row is a warning, the other row changes
    wired = seeded.execute_dict(
        {
            "verb": "plan",
            "args": {"action": "urls", "replace": [MS, WIKI]},
            "actor": "agent",
        }
    )
    assert wired.ok and wired.data["replaced"] == [2] and wired.data["refused"] == [1] and len(wired.warnings) == 1
    assert [r["urls"] for r in seeded.execute("list", entity=PLAN).data] == [[MS, WIKI], [WIKI], None]
    entry = seeded.storage.audit_tail(1)[0]
    assert entry["actor"] == "agent" and entry["payload"]["rows"] == [2] and entry["payload"]["refused"] == [1]
    # nothing to do: no write, no audit entry, no event
    events = len(seeded.events.last(50))
    again = seeded.execute("plan", action="urls", replace=[MS, WIKI])
    assert again.ok and again.data["replaced"] == [] and again.data["refused"] == [1]
    assert seeded.storage.audit_tail(1)[0] == entry and len(seeded.events.last(50)) == events
    # the refusals: both URLs absolute http(s), two of them, different
    assert seeded.execute("plan", action="urls", replace=["microsoft.com", WIKI]).errors == [
        "'microsoft.com' is not an absolute http(s) URL"
    ]
    assert seeded.execute("plan", action="urls", replace=[MS, "wiki/page"]).errors == [
        "'wiki/page' is not an absolute http(s) URL"
    ]
    assert seeded.execute("plan", action="urls", replace=[MS, MS]).errors == [
        f"the old and the new URL are the same ({MS})"
    ]
    assert seeded.execute("plan", action="urls", replace=[MS]).errors == [
        "plan urls --replace needs two URLs: the old and the new"
    ]
    assert seeded.execute("plan", action="urls", replace=MS).errors == [
        "plan urls --replace needs two URLs: the old and the new"
    ]
    assert seeded.execute("plan", action="urls", check=True).errors == [
        "plan urls needs --placeholder <url>, --audit or --replace <old-url> <new-url>"
    ]


def test_the_cli_audits_and_replaces(tmp_path, capsys):
    db = str(tmp_path / "cli.db")
    code, _res = run(capsys, "--db", db, "create", PLAN, "--set", "deliverable=Safety Plan", "--set", "order=1")
    assert code == 0
    code, _res = run(
        capsys,
        "--db",
        db,
        "create",
        PLAN,
        "--set",
        "deliverable=HARA",
        "--set",
        "order=2",
        "--set",
        "work_product_location=Confluence",
    )
    assert code == 0
    code, res = run(capsys, "--db", db, "plan", "urls", "--placeholder", MS)
    assert code == 0 and res["data"]["filled"] == [1, 2]
    code, res = run(capsys, "--db", db, "plan", "urls", "--audit")
    assert (
        code == 0
        and res["ok"]
        and [(f["id"], f["code"]) for f in res["data"]]
        == [
            (1, "placeholder"),
            (2, "placeholder"),
            (2, "no_confluence"),
        ]
    )
    code, res = run(capsys, "--db", db, "plan", "urls", "--audit", "--rows", "2")
    assert code == 0 and [f["id"] for f in res["data"]] == [2, 2]
    code, res = run(capsys, "--db", db, "plan", "urls", "--replace", MS, WIKI, "--check")
    assert code == 0 and res["data"]["would_replace"] == [1, 2] and res["data"]["replaced"] == []
    code, res = run(capsys, "--db", db, "plan", "urls", "--replace", MS, WIKI)
    assert code == 0 and res["data"]["replaced"] == [1, 2] and res["data"]["kind"] == "confluence"
    code, res = run(capsys, "--db", db, "list", PLAN)
    assert [r["urls"] for r in res["data"]] == [[WIKI], [WIKI]]
    code, res = run(capsys, "--db", db, "plan", "urls", "--audit")
    assert code == 0 and res["data"] == []
    code, res = run(capsys, "--db", db, "plan", "urls", "--replace", WIKI, "nope")
    assert code == 1 and res["errors"] == ["'nope' is not an absolute http(s) URL"]
    code, res = run(capsys, "--db", db, "schema")
    plan = next(c for c in res["data"]["commands"] if c["verb"] == "plan")
    replace = next(a for a in plan["arguments"] if a["name"] == "replace")
    assert replace["nargs"] == 2 and replace["flag"] is True


# -- the views -----------------------------------------------------------------------------------------------


def test_the_excel_work_product_cell_reads_the_kind_and_the_host(seeded, tmp_path):
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [WIKI, MS]}).ok
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [MS, WIKI]}).ok
    assert seeded.execute("update", entity=PLAN, id=3, data={"urls": [JAMA]}).ok
    exported = seeded.execute("export", format="excel", out=str(tmp_path / "view.xlsx"))
    assert exported.ok, exported.errors
    ws = load_workbook(exported.data["path"])["Safety Plan"]
    headers = [c.value for c in ws[1]]
    col = headers.index(WORK_PRODUCT_HEADER) + 1
    first, second, third = (ws.cell(row=r, column=col) for r in (2, 3, 4))
    assert first.value == "Confluence, asirobots.atlassian.net" and first.hyperlink.target == WIKI
    assert first.style == "Hyperlink"
    assert second.value == "placeholder" and second.hyperlink.target == MS  # the row still points nowhere
    assert third.value == "other, jama.example.com" and third.hyperlink.target == JAMA
    with Brain(tmp_path / "re.db") as other:  # the derived column is still skipped by the import
        imported = other.execute("import", format="excel", path=exported.data["path"], mode="replace")
        assert imported.ok, imported.errors
        assert not any(WORK_PRODUCT_HEADER in w for w in imported.warnings)
        assert [r["urls"] for r in other.execute("list", entity=PLAN).data] == [[WIKI, MS], [MS, WIKI], [JAMA]]


def test_the_excel_cell_uses_the_owners_host_map(tmp_path, monkeypatch):
    toml = tmp_path / "kletz.toml"
    toml.write_text('[urls]\nhosts = { jama = "jama.example.com" }\n', encoding="utf-8")
    monkeypatch.setenv(ENV_CONFIG, str(toml))
    with Brain(tmp_path / "k.db") as brain:
        brain.execute("create", entity=PLAN, data={"order": 1, "deliverable": "HARA", "urls": [JAMA]})
        exported = brain.execute("export", format="excel", out=str(tmp_path / "view.xlsx"))
        assert exported.ok, exported.errors
        ws = load_workbook(exported.data["path"])["Safety Plan"]
        col = [c.value for c in ws[1]].index(WORK_PRODUCT_HEADER) + 1
        assert ws.cell(row=2, column=col).value == "jama, jama.example.com"


def test_the_forms_link_lines_carry_the_kind():
    """The text of each link line of the URLs editor (the widget test in ``test_gui_links_and_lookup.py`` sees
    them on a real screen); the owner's map reaches it through the screen."""
    pytest.importorskip("tkinter")
    from asi_kletz.gui.widgets import THE_LINK, link_text

    assert link_text(WIKI, 0) == f"confluence: {WIKI}{THE_LINK}"
    assert link_text(MS, 1) == f"placeholder: {MS}" and link_text(BROWSE, 2) == f"jira: {BROWSE}"
    assert link_text(JAMA, 1) == f"other: {JAMA}" and link_text(JAMA, 1, HOSTS) == f"jama: {JAMA}"
