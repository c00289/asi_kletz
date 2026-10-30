"""The export fills the Word original in place when the standard's folder holds one (0.16.0): the owner,
2026-10-29, "can we make it exactly like the template? I still want all the columns to pop up in every page
but otherwise we need to match the template exactly."""

from __future__ import annotations

import os
import re
from pathlib import Path

import pytest

from tests.conftest import PLAN, SEED_WEIGHTS
from tests.test_standards import _stub, standards_dir

WINWORD = Path(os.environ.get("ProgramFiles", r"C:\Program Files")) / "Microsoft Office/root/Office16/WINWORD.EXE"
GREEN = "A0D8A0"  # asi_gutenberg 0.7.2 THEME["mark_bg"], the owner's "nice green, not neon, not forest"
PINK = "F4C2C2"  # asi_gutenberg 0.10.0 THEME["warn_bg"], the owner's "pinkish red" for a gap
DARK_RED = "B71C1C"  # asi_gutenberg 0.10.0 THEME["warn_fg"], the text of a link to the placeholder



def _stamped(path, folder, stem, suffix):
    """0.27.0: an exported file is ``<stem>_<YYYYMMDD-HHMMSS>.<suffix>`` in ``folder`` (any folder when None)."""
    path = Path(path)
    return (folder is None or path.parent == Path(folder)) and bool(
        re.fullmatch(stem + r"_\d{8}-\d{6}\." + suffix, path.name)
    )

def _fill_of(run) -> str | None:
    from docx.oxml.ns import qn

    rpr = run._r.find(qn("w:rPr"))
    shd = rpr.find(qn("w:shd")) if rpr is not None else None
    return shd.get(qn("w:fill")) if shd is not None else None
PARAMETERS = ({"name": "vehicle_type", "value": "Wheel Loader"}, {"name": "fusa_engineer", "value": "the contractor"})


@pytest.fixture
def standards(tmp_path, monkeypatch):
    return standards_dir(tmp_path, monkeypatch)


def _original(path: Path) -> Path:
    """The Word original the test standard's Markdown was imported from: the same units, its blanks unnamed
    (the Markdown's first heading is its title, so the Word file has the title only)."""
    import docx

    d = docx.Document()
    d.add_paragraph("__________ Safety Plan (test)", style="Title")
    p = d.add_paragraph()
    p.add_run("Author: ").bold = True
    p.add_run("__________")
    t = d.add_table(rows=3, cols=3)
    t.style = "Table Grid"
    for c, h in enumerate(("Deliverable", "Expected Duration/Deadline", "Status")):
        t.cell(0, c).text = h
    t.cell(1, 0).text, t.cell(1, 1).text, t.cell(1, 2).text = "Safety Plan", "__ points\n__________", "___\n_____"
    t.cell(2, 0).text, t.cell(2, 1).text, t.cell(2, 2).text = "HARA", "__ points\n__________", "___\n_____"
    d.sections[0].footer.paragraphs[0].text = "Proprietary & Confidential"
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path


def test_export_fills_the_word_original_when_the_standard_has_one(seeded, standards, tmp_path, monkeypatch):
    import docx
    from docx.oxml.ns import qn

    original = _original(standards / "qf-17-test" / "source" / "QF-17 test.docx")
    assert seeded.execute("plan", action="weights", what="load", file=str(SEED_WEIGHTS)).ok
    assert seeded.execute("create", entity="sprint", data={
        "number": 2, "label": "SPRINT 2", "start": "2026-10-27", "end": "2026-11-09", "status": "Active",
    }).ok
    for data in PARAMETERS:
        assert seeded.execute("create", entity="parameter", data=data).ok
    assert seeded.execute("update", entity=PLAN, id=1, set=["status=Ongoing", "jira_key=hwl-2554", "points=5"]).ok
    _stub(seeded, monkeypatch, {"HWL-2554": "In Progress"})
    listed = seeded.execute("standards", action="list")
    assert listed.ok and listed.data["standards"][0]["template"] == str(original)

    out = tmp_path / "out"
    result = seeded.execute("plan", action="export", row=1, out=str(out))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    assert built["template"] == str(original) and built["exact"]["docx"] is True
    assert _stamped(built["files"]["docx"], out, "qf-17-test", "docx") and Path(built["files"]["md"]).is_file()
    word = docx.Document(built["files"]["docx"])
    assert word.paragraphs[0].text == "Wheel Loader Safety Plan (test)" and word.paragraphs[0].style.name == "Title"
    assert word.core_properties.title == "Wheel Loader Safety Plan (test)"
    author = word.paragraphs[1]
    assert author.text == "Author: the contractor" and author.runs[0].bold and author.runs[0].text == "Author: "
    table = word.tables[0]
    assert table.style.name == "Table Grid" and table.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    assert [p.text for p in table.cell(1, 1).paragraphs] == ["5 points", "2026-11-09"]
    assert [p.text for p in table.cell(1, 2).paragraphs] == ["HWL-2554", "In Progress"]  # the key over the status
    assert [p.text for p in table.cell(2, 1).paragraphs] == ["3 points", "2026-11-09"]
    unfilled_cell = [p.text for p in table.cell(2, 2).paragraphs]
    assert unfilled_cell == ["no ticket attached", ""]  # no ticket: the cell says so, never a placeholder (0.25.1)
    assert word.sections[0].footer.paragraphs[0].text == "Proprietary & Confidential"
    assert "==" not in word.paragraphs[0].text  # a mark is shading on the run, never text in the Word file
    green = GREEN
    shaded = [r.text for p in word.paragraphs for r in p.runs if _fill_of(r) == green]
    assert shaded == ["Wheel Loader", "the contractor"]  # 0.17.0: the filled values shaded the theme's green
    status_runs = [r.text for p in table.cell(1, 2).paragraphs for r in p.runs if _fill_of(r) == green]
    assert status_runs == ["HWL-2554", "In Progress"]
    no_ticket = [(r.text, _fill_of(r)) for p in table.cell(2, 2).paragraphs for r in p.runs if r.text]
    assert no_ticket == [("no ticket attached", PINK)]  # 0.26.0: a gap, pinkish red, never green
    assert built["warn"] == ["row_2_jira_key"] and built["warn_urls"] == []
    assert not any(_fill_of(r) == green for r in author.runs[:1])  # the template's own text is not
    assert "==Wheel Loader==" in Path(built["files"]["md"]).read_text(encoding="utf-8")  # the source keeps them
    if WINWORD.is_file():
        assert built["exact"]["pdf"] is True and _stamped(built["files"]["pdf"], out, "qf-17-test", "pdf")
        assert Path(built["files"]["pdf"]).stat().st_size > 1000
        assert not any("pdf through Word" in w for w in result.warnings)
    else:
        assert built["exact"]["pdf"] is False and Path(built["files"]["pdf"]).is_file()
        assert any("the PDF is Gutenberg's rendering" in w for w in result.warnings)


def test_a_failed_print_keeps_gutenbergs_pdf(seeded, standards, tmp_path, monkeypatch):
    """0.16.1: when Word cannot print, the PDF is Gutenberg's rendering, as 0.16.0 promised. asi_gutenberg 0.7.0
    removed that rendering before asking Word, so a failed print (no Word on the machine, as on the CI runner)
    left the export naming a PDF that was gone; 0.7.1 prints to a working name and leaves the target alone."""
    import subprocess

    _original(standards / "qf-17-test" / "source" / "QF-17 test.docx")
    for data in PARAMETERS:
        assert seeded.execute("create", entity="parameter", data=data).ok
    _stub(seeded, monkeypatch, {})

    def no_word(*args, **kwargs):
        raise OSError("no Word here")

    monkeypatch.setattr(subprocess, "run", no_word)  # what asi_gutenberg's pdf_with_word calls
    result = seeded.execute("plan", action="export", row=1, out=str(tmp_path / "out"))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    assert built["exact"] == {"docx": True, "pdf": False}
    assert _stamped(built["files"]["pdf"], tmp_path / "out", "qf-17-test", "pdf") and Path(built["primary"]).is_file()
    assert Path(built["primary"]).stat().st_size > 1000
    assert any("pdf through Word failed: no Word here; the PDF is Gutenberg's rendering" in w for w in result.warnings)


def test_a_folder_without_an_original_exports_as_before(seeded, standards):
    listed = seeded.execute("standards", action="list")
    assert listed.ok and listed.data["standards"][0]["template"] is None


def test_two_originals_name_none(seeded, standards):
    _original(standards / "qf-17-test" / "source" / "one.docx")
    _original(standards / "qf-17-test" / "source" / "two.docx")
    listed = seeded.execute("standards", action="list")
    assert listed.ok and listed.data["standards"][0]["template"] is None


@pytest.mark.skipif(not WINWORD.is_file(), reason="Word is not installed here")
def test_the_pdf_is_words_print(seeded, standards, tmp_path, monkeypatch):
    _original(standards / "qf-17-test" / "source" / "QF-17 test.docx")
    for data in PARAMETERS:
        assert seeded.execute("create", entity="parameter", data=data).ok
    _stub(seeded, monkeypatch, {})
    result = seeded.execute("plan", action="export", row=1, out=str(tmp_path / "out"))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    assert built["exact"] == {"docx": True, "pdf": True} and _stamped(built["primary"], None, "qf-17-test", "pdf")


# -- 0.21.0: the deliverable's link and the abridged headers ----------------------------------------------
MS = "https://www.microsoft.com/"
BLUE = "CFE2F3"  # asi_gutenberg 0.9.0 THEME["link_bg"], the pale blue under a link
ABRIDGED = {"Deliverable": "Del.", "Expected Duration/Deadline": "Duration", "Status": "Status"}


def _linked_standard(standards: Path) -> Path:
    """The test standard with its deliverable names written as links to the row's url (asi_gutenberg 0.9.0,
    the owner: "turn the name of the safety deliverable into a clickable URL") and the abridged headings
    beside the source; the original under source/."""
    import json

    source = standards / "qf-17-test" / "qf-17-test.md"
    text = source.read_text(encoding="utf-8")
    text = text.replace("| Safety Plan |", "| [Safety Plan]({{ row_1_url }}) |")
    text = text.replace("| HARA |", "| [HARA]({{ row_2_url }}) |")
    source.write_text(text, encoding="utf-8")
    mapping = standards / "qf-17-test" / "abridged-headers.json"
    mapping.write_text(json.dumps(ABRIDGED), encoding="utf-8")
    _original(standards / "qf-17-test" / "source" / "QF-17 test.docx")
    return mapping


def _links_of(cell) -> list[tuple[str, str, str | None, str | None]]:
    """Every hyperlink of a Word cell: (target, text, character style, shading fill)."""
    return [(target, text, style, fill) for target, text, style, _colour, fill in _coloured_links_of(cell)]


def _coloured_links_of(cell) -> list[tuple[str, str, str | None, str | None, str | None]]:
    """Every hyperlink of a Word cell: (target, text, character style, text colour, shading fill)."""
    from docx.oxml.ns import qn

    out = []
    for h in cell._tc.iter(qn("w:hyperlink")):
        rpr = h.find(qn("w:r") + "/" + qn("w:rPr"))
        style = rpr.find(qn("w:rStyle")) if rpr is not None else None
        colour = rpr.find(qn("w:color")) if rpr is not None else None
        shd = rpr.find(qn("w:shd")) if rpr is not None else None
        text = "".join(t.text or "" for t in h.iter(qn("w:t")))
        out.append((cell.part.rels[h.get(qn("r:id"))].target_ref, text,
                    style.get(qn("w:val")) if style is not None else None,
                    colour.get(qn("w:val")) if colour is not None else None,
                    shd.get(qn("w:fill")) if shd is not None else None))
    return out


def _cell_fills(cell) -> list[tuple[str, str | None]]:
    """Every run of a cell with text, inside and outside its hyperlinks: (text, shading fill)."""
    from docx.oxml.ns import qn

    out = []
    for r in cell._tc.iter(qn("w:r")):
        text = "".join(t.text or "" for t in r.findall(qn("w:t")))
        if text:
            shd = r.find(qn("w:rPr") + "/" + qn("w:shd"))
            out.append((text, shd.get(qn("w:fill")) if shd is not None else None))
    return out


CONFLUENCE = "https://asirobots.atlassian.net/wiki/pages/viewpage.action?pageId=399343803"


def test_export_links_the_deliverable_and_passes_the_abridged_headers(seeded, standards, tmp_path, monkeypatch):
    """0.21.0: the deliverable cell of a row with a URL is a hyperlink to its first URL, blue when it is a real
    one; 0.26.0: pinkish red, dark red text, still a link, when it is the placeholder (the owner: "when the
    link is microsoft.com because of default ... highlight in pinkish red instead of blue"); the no-ticket
    key pinkish red, the points green; ``abridged-headers.json`` beside the source goes to the fill as the
    mapping and the export carries what Word did (``abridged``, ``pages``), here with the Word session
    replaced."""
    from importlib import import_module

    import docx
    from asi_gutenberg.fill_docx import WordSession
    from docx.oxml.ns import qn

    fill_module = import_module("asi_gutenberg.fill_docx")  # the package exports the function of that name

    mapping = _linked_standard(standards)
    for data in PARAMETERS:
        assert seeded.execute("create", entity="parameter", data=data).ok
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS, "https://dms.example.com/doc/17"]}).ok
    assert seeded.execute("update", entity=PLAN, id=2, set=["ref_qf_17=1", "jira_key=HWL-7", "points=2"]).ok
    assert seeded.execute("update", entity=PLAN, id=2, data={"urls": [CONFLUENCE]}).ok
    _stub(seeded, monkeypatch, {"HWL-7": "Open"})
    sessions: list[dict] = []

    def fake_session(docx_path, pdf_path, plan=None, timeout=180, read_only_recommended=False):
        sessions.append({"docx": docx_path, "pdf": pdf_path, "plan": plan})
        Path(pdf_path).write_bytes(b"%PDF-1.4 fake\n" + b"0" * 2000)
        return WordSession(printed=True, abridged=["Deliverable"], pages=3)

    monkeypatch.setattr(fill_module, "word_session", fake_session)
    listed = seeded.execute("standards", action="list")
    assert listed.ok and listed.data["standards"][0]["abridged_headers"] == str(mapping)

    out = tmp_path / "out"
    result = seeded.execute("plan", action="export", row=1, out=str(out))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    assert built["abridged"] == ["Deliverable"] and built["pages"] == 3 and built["abridged_headers"] == str(mapping)
    assert built["exact"] == {"docx": True, "pdf": True} and _stamped(built["primary"], out, "qf-17-test", "pdf")
    [session] = sessions
    assert session["plan"] == [{"index": 1, "first": "Deliverable", "short": ["Del.", "Duration", "Status"]}]
    assert built["warn"] == ["row_1_jira_key", "row_1_url"]
    assert built["warn_urls"] == [MS]
    word = docx.Document(built["files"]["docx"])
    table = word.tables[0]
    assert table.cell(1, 0).text == "Safety Plan"
    # the first URL is the placeholder: still a hyperlink to it, dark red over pinkish red, not blue (0.26.0)
    assert _coloured_links_of(table.cell(1, 0)) == [(MS, "Safety Plan", "Hyperlink", DARK_RED, PINK)]
    assert not any(s.get(qn("w:fill")) in (GREEN, BLUE) for s in table.cell(1, 0)._tc.iter(qn("w:shd")))
    assert table.cell(2, 0).text == "HARA"  # a real URL: blue, Word's own hyperlink colour
    assert _coloured_links_of(table.cell(2, 0)) == [(CONFLUENCE, "HARA", "Hyperlink", "0563C1", BLUE)]
    assert _cell_fills(table.cell(1, 2)) == [("no ticket attached", PINK)]  # the key: a gap
    assert _cell_fills(table.cell(2, 2)) == [("HWL-7", GREEN), ("Open", GREEN)]  # a ticket: green
    assert ("2", GREEN) in _cell_fills(table.cell(2, 1))  # the points: green
    assert "Hyperlink" in [s.name for s in word.styles]
    md = Path(built["files"]["md"]).read_text(encoding="utf-8")
    assert f"[!!Safety Plan!!]({MS})" in md and f"[HARA]({CONFLUENCE})" in md  # the url never marked
    assert "!!no ticket attached!!<br>" in md and "==HWL-7==<br>==Open==" in md


def test_without_word_the_abridged_headers_are_a_warning(seeded, standards, tmp_path, monkeypatch):
    """Without Word the fill cannot split the table: the export warns once for the pass, the Word file keeps
    the verbatim header repeating, ``abridged`` is empty and ``pages`` is the rendering's."""
    import subprocess

    _linked_standard(standards)
    for data in PARAMETERS:
        assert seeded.execute("create", entity="parameter", data=data).ok
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS]}).ok
    _stub(seeded, monkeypatch, {})

    def no_word(*args, **kwargs):
        raise OSError("no Word here")

    monkeypatch.setattr(subprocess, "run", no_word)
    result = seeded.execute("plan", action="export", row=1, out=str(tmp_path / "out"))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    assert built["exact"] == {"docx": True, "pdf": False} and built["abridged"] == []
    assert isinstance(built["pages"], int) and built["pages"] >= 1
    assert any("abridged headers need Word" in w and "keeps the verbatim header on every page" in w
               for w in result.warnings)
    assert any("pdf through Word failed: no Word here; the PDF is Gutenberg's rendering" in w for w in result.warnings)
    assert any("unfilled parameter row_2_url" in w for w in result.warnings)
    import docx
    from docx.oxml.ns import qn

    word = docx.Document(built["files"]["docx"])
    assert len(word.tables) == 1 and word.tables[0].rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    assert _links_of(word.tables[0].cell(1, 0)) == [(MS, "Safety Plan", "Hyperlink", PINK)]  # the placeholder
    # no URL at all: the plain name, no link, shaded pinkish red (the owner: "or blank")
    assert word.tables[0].cell(2, 0).text == "HARA" and _links_of(word.tables[0].cell(2, 0)) == []
    assert _cell_fills(word.tables[0].cell(2, 0)) == [("HARA", PINK)]
    md = Path(built["files"]["md"]).read_text(encoding="utf-8")
    assert "[HARA]({{ row_2_url }})" in md  # the Markdown keeps the blank; the renderings paint it as a gap


@pytest.mark.skipif(not WINWORD.is_file(), reason="Word is not installed here")
def test_word_prints_the_linked_deliverable_with_the_mapping(seeded, standards, tmp_path, monkeypatch):
    """Word itself: the mapping is checked against the table, which ends on its first page, so nothing is
    split; the PDF is printed with the link in it and the page count is Word's."""
    _linked_standard(standards)
    for data in PARAMETERS:
        assert seeded.execute("create", entity="parameter", data=data).ok
    assert seeded.execute("update", entity=PLAN, id=1, data={"urls": [MS]}).ok
    _stub(seeded, monkeypatch, {})
    result = seeded.execute("plan", action="export", row=1, out=str(tmp_path / "out"))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    assert built["exact"] == {"docx": True, "pdf": True}, result.warnings
    assert built["abridged"] == [] and built["pages"] == 1, (built, result.warnings)
    assert Path(built["primary"]).stat().st_size > 1000
    assert not any("abridged" in w or "pdf through Word" in w for w in result.warnings)


def test_a_refused_gutenberg_rendering_still_fills_the_word_original(seeded, standards, tmp_path, monkeypatch):
    """0.27.1, the owner's "plan failed" dialog: reportlab refused a plan table whose rows had grown taller than
    the page and the whole export failed. Gutenberg's own rendering is a fallback, not the gate: the Markdown is
    written, the Word original is filled (and printed where Word runs), and a warning names the refusal."""
    import asi_gutenberg

    _original(standards / "qf-17-test" / "source" / "QF-17 test.docx")

    def refuse(*args, **kwargs):
        raise RuntimeError("LayoutError: Flowable <Table 4 rows x 10 cols> too large on page 20 in frame 'turned'")

    monkeypatch.setattr(asi_gutenberg, "build", refuse)
    out = tmp_path / "out"
    result = seeded.execute("plan", action="export", row=1, out=str(out))
    assert result.ok, result.errors
    [built] = result.data["exports"]
    md = Path(built["files"]["md"]).read_text(encoding="utf-8")
    assert _stamped(built["files"]["md"], out, "qf-17-test", "md") and "!!no ticket attached!!" in md  # filled, marked
    assert "pdf" not in built["files"] or built["exact"]["pdf"] is True  # Gutenberg's own PDF is not there to fall back on
    assert _stamped(built["files"]["docx"], out, "qf-17-test", "docx") and built["exact"]["docx"] is True
    assert isinstance(built["unfilled"], list) and "vehicle_type" in built["unfilled"]  # nothing set it here
    assert any("could not render its own PDF and Word file (RuntimeError: LayoutError" in w for w in result.warnings)
    assert any("the Word original is filled and printed as before" in w for w in result.warnings)
